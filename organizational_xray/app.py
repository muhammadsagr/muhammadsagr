"""Organizational X-Ray — an "X-ray" of an organisation's structure, in one file.

Loads employee / org-structure data, builds the reporting network and analyses
layers, span of control, departments, reporting relationships, vacancies,
duplicate / similar job titles and data quality. Every finding is explainable
(rule, value, threshold, benchmark, source records), and a What-If Simulator
tests structural changes on a copy of the data.

    Data -> Rules -> Calculations -> Evidence -> Findings -> Visualization

Run:
    streamlit run app.py          (or simply: python app.py)
Other commands:
    python app.py --generate [--employees 1500] [--seed 42]   write the sample dataset + report
    pytest                                                     run the tests

The file is organised in sections: configuration, data (loader, validator,
dummy generator), analytics (hierarchy, span, departments, findings, data
quality, simulations, AI-ready engine), charts, export, and the Streamlit pages.
Business logic never imports Streamlit; only the UI sections do.
"""
from __future__ import annotations

from dataclasses import asdict, dataclass, field, fields
from datetime import date, timedelta
from difflib import SequenceMatcher
from functools import cached_property
from itertools import combinations
from pathlib import Path
from typing import Callable, IO
import argparse
import hashlib
import html
import io
import itertools
import json
import logging
import re

import importlib.util
import subprocess
import sys

# Install missing libraries automatically on first run.
_REQUIRED = {"streamlit": "streamlit>=1.50", "pandas": "pandas>=2.2", "numpy": "numpy>=1.26",
             "plotly": "plotly>=5.20", "networkx": "networkx>=3.2", "openpyxl": "openpyxl>=3.1",
             "faker": "Faker>=24.0"}
_missing = [pkg for mod, pkg in _REQUIRED.items() if importlib.util.find_spec(mod) is None]
if _missing:
    print("Installing required packages:", ", ".join(_missing))
    subprocess.check_call([sys.executable, "-m", "pip", "install", *_missing])

from faker import Faker
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import networkx as nx
import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st

if not st.runtime.exists():  # command line / tests: silence "no runtime" cache warnings
    for _name in ("cache_data_api", "cache_resource_api"):
        logging.getLogger(f"streamlit.runtime.caching.{_name}").setLevel(logging.ERROR)



# ================================================================================================
# Configuration: schema, labels and adjustable thresholds  (utils/config.py)
# ================================================================================================
# --------------------------------------------------------------------------- schema
# Canonical (internal) column names. Everything downstream uses these.
EMPLOYEE_ID = "employee_id"
EMPLOYEE_NAME = "employee_name"
MANAGER_ID = "manager_id"
MANAGER_NAME = "manager_name"

CANONICAL_COLUMNS: list[str] = [
    "employee_id",
    "employee_name",
    "gender",
    "age",
    "hire_date",
    "employment_status",
    "job_title",
    "job_level",
    "grade",
    "department",
    "division",
    "business_unit",
    "location",
    "manager_id",
    "manager_name",
    "salary",
    "position_id",
    "position_status",
    "employment_type",
    "work_arrangement",
    "performance_rating",
    "years_of_service",
]

# Columns required by the import contract (Section 29 of the spec).
IMPORT_REQUIRED_COLUMNS: list[str] = [
    "employee_id",
    "employee_name",
    "manager_id",
    "job_title",
    "department",
    "division",
    "business_unit",
    "job_level",
    "salary",
    "employment_status",
    "position_id",
    "position_status",
]
# Without these two the hierarchy cannot be built at all.
STRUCTURALLY_REQUIRED: list[str] = ["employee_id", "manager_id"]

STRING_COLUMNS = [
    "employee_id", "employee_name", "gender", "employment_status", "job_title", "grade",
    "department", "division", "business_unit", "location", "manager_id", "manager_name",
    "position_id", "position_status", "employment_type", "work_arrangement",
]
NUMERIC_COLUMNS = ["age", "job_level", "salary", "performance_rating", "years_of_service"]
DATE_COLUMNS = ["hire_date"]

# Display names for tables / exports.
DISPLAY_NAMES: dict[str, str] = {
    "employee_id": "Employee ID",
    "employee_name": "Name",
    "gender": "Gender",
    "age": "Age",
    "hire_date": "Hire Date",
    "employment_status": "Employment Status",
    "job_title": "Job Title",
    "job_level": "Job Level",
    "grade": "Grade",
    "department": "Department",
    "division": "Division",
    "business_unit": "Business Unit",
    "location": "Location",
    "manager_id": "Manager ID",
    "manager_name": "Manager Name",
    "salary": "Salary",
    "position_id": "Position ID",
    "position_status": "Position Status",
    "employment_type": "Employment Type",
    "work_arrangement": "Work Arrangement",
    "performance_rating": "Performance Rating",
    "years_of_service": "Years of Service",
    # derived
    "node_id": "Node ID",
    "depth": "Depth",
    "layer": "Layer",
    "reporting_status": "Reporting Status",
    "direct_reports": "Direct Reports",
    "vacant_direct_positions": "Vacant Direct Positions",
    "total_reports": "Total Reports",
    "is_manager": "Is Manager",
    "is_active": "Is Active",
    "is_vacant": "Is Vacant",
    "span_category": "Span Category",
}

# Allowed values (validation + normalisation)
EMPLOYMENT_STATUSES = ["Active", "On Leave", "Terminated"]
ACTIVE_STATUSES = {"Active", "On Leave"}
POSITION_STATUSES = ["Filled", "Vacant"]
EMPLOYMENT_TYPES = ["Full-Time", "Part-Time", "Contractor"]
WORK_ARRANGEMENTS = ["On-site", "Hybrid", "Remote"]

# Nominal names of the job levels (declared hierarchy, not measured depth).
LEVEL_NAMES: dict[int, str] = {
    0: "CEO",
    1: "C-Suite",
    2: "Directors",
    3: "Senior Managers",
    4: "Managers",
    5: "Supervisors",
    6: "Individual Contributors",
}

# Finding vocabulary --------------------------------------------------------------
SEVERITIES = ["Critical", "Significant", "Attention", "Informational"]
SEVERITY_RANK = {s: i for i, s in enumerate(SEVERITIES)}  # 0 = most severe

FINDING_TYPES = [
    "CIRCULAR_REPORTING",
    "ORPHAN_EMPLOYEE",
    "INVALID_MANAGER",
    "TERMINATED_MANAGER",
    "HIGH_SPAN_OF_CONTROL",
    "LOW_SPAN_OF_CONTROL",
    "DEEP_HIERARCHY",
    "HIGH_MANAGEMENT_RATIO",
    "SMALL_DEPARTMENT",
    "DUPLICATE_POSITION",
    "VACANT_POSITION",
]

# Reporting-status vocabulary (hierarchy engine)
RS_CEO = "Top of Hierarchy"
RS_VALID = "Valid"
RS_MISSING = "Missing Manager"
RS_INVALID = "Invalid Manager"
RS_TERMINATED = "Terminated Manager"
RS_CIRCULAR = "Circular Reporting"
RS_DISCONNECTED = "Disconnected (upstream break)"
RS_INACTIVE = "Not Active"
RS_VACANT = "Vacant Position"
RS_DUPLICATE = "Duplicate ID (excluded)"
INVALID_REPORTING_STATUSES = {RS_MISSING, RS_INVALID, RS_TERMINATED, RS_CIRCULAR}


# --------------------------------------------------------------------------- settings
@dataclass
class Settings:
    """User-adjustable analysis thresholds (Settings page)."""

    low_span_threshold: int = 3  # direct reports <= this -> Low Span of Control
    high_span_threshold: int = 12  # direct reports > this -> High Span of Control
    deep_hierarchy_threshold: int = 6  # depth (edges to CEO) > this -> Deep Hierarchy
    small_department_threshold: int = 3  # headcount <= this -> Small Organizational Unit
    high_management_ratio: float = 0.20  # managers / headcount > this -> High Mgmt Ratio
    title_similarity_threshold: float = 0.80  # 0..1, for Similar Job Titles
    exclude_executives_from_span: bool = True  # skip job levels 0-1 in span findings
    max_network_nodes: int = 600  # safety cap for the network chart

    def as_dict(self) -> dict:
        return asdict(self)

    def cache_key(self) -> tuple:
        return tuple(getattr(self, f.name) for f in fields(self))

    @classmethod
    def from_dict(cls, data: dict) -> "Settings":
        known = {f.name for f in fields(cls)}
        return cls(**{k: v for k, v in data.items() if k in known})


SETTING_DESCRIPTIONS: dict[str, str] = {
    "low_span_threshold": "Low Span Threshold — مدير لديه عدد موظفين مباشرين أقل من أو يساوي هذا الحد",
    "high_span_threshold": "High Span Threshold — مدير لديه عدد موظفين مباشرين أكبر من هذا الحد",
    "deep_hierarchy_threshold": "Deep Hierarchy Threshold — عمق (عدد العلاقات حتى الـCEO) أكبر من هذا الحد",
    "small_department_threshold": "Small Department Threshold — إدارة عدد موظفيها أقل من أو يساوي هذا الحد",
    "high_management_ratio": "High Management Ratio — نسبة المديرين إلى إجمالي موظفي الإدارة أكبر من هذا الحد",
    "title_similarity_threshold": "Title Similarity — حد التشابه النصي بين المسميات الوظيفية (0-1)",
    "exclude_executives_from_span": "استثناء CEO والـC-Suite من ملاحظات Span of Control",
    "max_network_nodes": "الحد الأقصى لعدد العقد في رسم الشبكة التنظيمية",
}


# ================================================================================================
# Helpers  (utils/helpers.py)
# ================================================================================================
UNASSIGNED = "(Unassigned)"


def is_text(s: pd.Series) -> bool:
    return pd.api.types.is_object_dtype(s) or pd.api.types.is_string_dtype(s)


def blank_to_label(s: pd.Series, label: str = UNASSIGNED) -> pd.Series:
    if not is_text(s):
        return s
    return s.fillna("").astype(str).replace("", label)


def safe_div(a: float, b: float) -> float:
    return float(a) / float(b) if b else np.nan


def pct(x: float, digits: int = 1) -> str:
    return "—" if x is None or pd.isna(x) else f"{x * 100:.{digits}f}%"


def fmt_num(x, digits: int = 1) -> str:
    if x is None or pd.isna(x):
        return "—"
    if float(x).is_integer():
        return f"{int(x):,}"
    return f"{x:,.{digits}f}"


def df_fingerprint(df: pd.DataFrame) -> str:
    """Stable hash of a DataFrame's content (used as a cache token)."""
    h = hashlib.md5(pd.util.hash_pandas_object(df.astype(str), index=True).values.tobytes())
    return h.hexdigest()[:16]


def display_df(df: pd.DataFrame) -> pd.DataFrame:
    """Rename canonical/derived columns to human-friendly headers."""
    return df.rename(columns={c: DISPLAY_NAMES.get(c, c.replace("_", " ").title())
                              for c in df.columns})


# ================================================================================================
# Data loading & normalisation  (data/loader.py)
# ================================================================================================
# Header aliases -> canonical name (compared after lower-casing and stripping
# spaces, underscores and dashes).
_ALIASES: dict[str, list[str]] = {
    "employee_id": ["employeeid", "empid", "employeenumber", "empno", "id", "رقمالموظف"],
    "employee_name": ["employeename", "name", "fullname", "الاسم", "اسمالموظف"],
    "gender": ["gender", "sex", "الجنس"],
    "age": ["age", "العمر"],
    "hire_date": ["hiredate", "dateofhire", "joiningdate", "startdate", "تاريخالتوظيف"],
    "employment_status": ["employmentstatus", "status", "employeestatus", "الحالةالوظيفية"],
    "job_title": ["jobtitle", "title", "position", "positiontitle", "المسمىالوظيفي"],
    "job_level": ["joblevel", "level", "layer", "المستوىالوظيفي"],
    "grade": ["grade", "jobgrade", "payGrade".lower(), "الدرجة"],
    "department": ["department", "dept", "الإدارة", "الادارة"],
    "division": ["division", "sector", "القطاع"],
    "business_unit": ["businessunit", "bu", "وحدةالأعمال"],
    "location": ["location", "site", "city", "الموقع"],
    "manager_id": ["managerid", "supervisorid", "reportsto", "linemanagerid", "رقمالمدير"],
    "manager_name": ["managername", "supervisorname", "linemanager", "اسمالمدير"],
    "salary": ["salary", "basesalary", "annualsalary", "الراتب"],
    "position_id": ["positionid", "positionnumber", "posid", "رقمالوظيفة"],
    "position_status": ["positionstatus", "حالةالوظيفة"],
    "employment_type": ["employmenttype", "contracttype", "نوعالتوظيف"],
    "work_arrangement": ["workarrangement", "workmode", "workmodel"],
    "performance_rating": ["performancerating", "rating", "performance", "تقييمالأداء"],
    "years_of_service": ["yearsofservice", "tenure", "yos", "سنواتالخدمة"],
}
_ALIAS_LOOKUP = {a: canon for canon, aliases in _ALIASES.items() for a in aliases}
for _c in CANONICAL_COLUMNS:
    _ALIAS_LOOKUP.setdefault(_c.replace("_", ""), _c)

_STATUS_MAP = {
    "active": "Active", "employed": "Active", "onleave": "On Leave", "leave": "On Leave",
    "loa": "On Leave", "terminated": "Terminated", "inactive": "Terminated",
    "resigned": "Terminated", "exited": "Terminated", "separated": "Terminated",
}
_POSITION_MAP = {"filled": "Filled", "occupied": "Filled", "vacant": "Vacant", "open": "Vacant"}
_TYPE_MAP = {"fulltime": "Full-Time", "parttime": "Part-Time", "contractor": "Contractor",
             "contract": "Contractor", "temporary": "Contractor"}
_WORK_MAP = {"onsite": "On-site", "office": "On-site", "hybrid": "Hybrid", "remote": "Remote"}


def _key(text: str) -> str:
    return re.sub(r"[\s_\-./]+", "", str(text).strip().lower())


def map_columns(columns) -> dict[str, str]:
    """Return {original_header: canonical_name} for recognised headers."""
    mapping: dict[str, str] = {}
    taken: set[str] = set()
    for col in columns:
        canon = _ALIAS_LOOKUP.get(_key(col))
        if canon and canon not in taken:
            mapping[col] = canon
            taken.add(canon)
    return mapping


def read_uploaded(file: IO[bytes] | bytes, filename: str) -> pd.DataFrame:
    """Read a CSV or Excel upload into a raw DataFrame (all columns as text)."""
    data = file if isinstance(file, bytes) else file.read()
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm", ".xls")):
        return pd.read_excel(io.BytesIO(data), dtype=str, engine="openpyxl")
    for enc in ("utf-8-sig", "cp1256", "latin-1"):
        try:
            return pd.read_csv(io.BytesIO(data), dtype=str, encoding=enc,
                               keep_default_na=False, sep=None, engine="python")
        except UnicodeDecodeError:
            continue
    raise ValueError("Unable to decode the CSV file")


def _clean_str(s: pd.Series) -> pd.Series:
    out = s.astype("object").where(s.notna(), "")
    out = out.map(lambda v: re.sub(r"\s+", " ", str(v)).strip())
    return out.replace({"nan": "", "None": "", "NaN": "", "<NA>": "", "NaT": ""})


def _clean_id(s: pd.Series) -> pd.Series:
    out = _clean_str(s)
    # Excel often turns numeric IDs into floats ("1042.0")
    return out.str.replace(r"^(\d+)\.0$", r"\1", regex=True)


def _map_values(s: pd.Series, mapping: dict[str, str]) -> pd.Series:
    return s.map(lambda v: mapping.get(_key(v), v) if v else "")


def _parse_level(v) -> float:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return np.nan
    m = re.search(r"-?\d+", str(v))
    return float(m.group()) if m else np.nan


def normalize(raw: pd.DataFrame) -> pd.DataFrame:
    """Return a copy of ``raw`` in the canonical schema.

    Adds helper columns:
      * ``node_id``   - unique key used by the hierarchy graph (Employee ID, or
                        ``VAC-<PositionID>`` for vacant positions without one)
      * ``is_vacant`` - Position Status == Vacant
      * ``is_active`` - filled position held by an Active / On Leave employee
      * ``is_duplicate_id`` - second+ occurrence of an Employee ID
      * ``row_number`` - 1-based row number in the source file (for evidence)
    """
    df = raw.copy()
    df = df.rename(columns=map_columns(df.columns))
    df = df.loc[:, ~df.columns.duplicated()]
    for col in CANONICAL_COLUMNS:
        if col not in df.columns:
            df[col] = np.nan
    extra = [c for c in df.columns if c not in CANONICAL_COLUMNS]
    df = df[CANONICAL_COLUMNS + extra].reset_index(drop=True)
    df.insert(0, "row_number", np.arange(1, len(df) + 1))

    for col in STRING_COLUMNS:
        df[col] = _clean_id(df[col]) if col in ("employee_id", "manager_id", "position_id") \
            else _clean_str(df[col])
    df["employment_status"] = _map_values(df["employment_status"], _STATUS_MAP)
    df["position_status"] = _map_values(df["position_status"], _POSITION_MAP)
    df["employment_type"] = _map_values(df["employment_type"], _TYPE_MAP)
    df["work_arrangement"] = _map_values(df["work_arrangement"], _WORK_MAP)

    for col in NUMERIC_COLUMNS:
        if col == "job_level":
            df[col] = df[col].map(_parse_level).astype(float)
        else:
            cleaned = df[col].astype("object").map(
                lambda v: re.sub(r"[,\s]", "", str(v)) if pd.notna(v) else v)
            df[col] = pd.to_numeric(cleaned, errors="coerce")
    for col in DATE_COLUMNS:
        df[col] = pd.to_datetime(df[col], errors="coerce")

    # Position status default: rows without an employee are vacant positions.
    no_status = df["position_status"].eq("")
    df.loc[no_status, "position_status"] = np.where(
        df.loc[no_status, "employee_id"].eq(""), "Vacant", "Filled")
    df["is_vacant"] = df["position_status"].eq("Vacant")
    # Employment status default for filled positions
    df.loc[~df["is_vacant"] & df["employment_status"].eq(""), "employment_status"] = "Active"

    # Node IDs (graph keys)
    node = df["employee_id"].copy()
    vac_no_id = df["is_vacant"] & node.eq("")
    pos = df["position_id"].where(df["position_id"].ne(""), "ROW" + df["row_number"].astype(str))
    node.loc[vac_no_id] = "VAC-" + pos.loc[vac_no_id]
    missing_id = ~df["is_vacant"] & node.eq("")
    node.loc[missing_id] = "NOID-" + df.loc[missing_id, "row_number"].astype(str)
    df.insert(1, "node_id", node)
    df["is_duplicate_id"] = df["node_id"].duplicated(keep="first")
    df["is_active"] = (~df["is_vacant"] & df["employment_status"].isin(["Active", "On Leave"])
                       & ~df["is_duplicate_id"])
    # Years of service: derive from hire date when absent
    if df["years_of_service"].isna().all() and df["hire_date"].notna().any():
        ref = pd.Timestamp.today().normalize()
        df["years_of_service"] = ((ref - df["hire_date"]).dt.days / 365.25).round(1)
    return df


# ================================================================================================
# Import validation  (data/validator.py)
# ================================================================================================
@dataclass
class ValidationIssue:
    level: str  # Error | Warning | Info
    field: str
    message: str
    count: int = 0
    rows: list[int] = field(default_factory=list)  # 1-based source row numbers


@dataclass
class ValidationReport:
    issues: list[ValidationIssue]
    column_mapping: dict[str, str]
    missing_columns: list[str]
    n_rows: int

    @property
    def ok(self) -> bool:
        return not any(i.level == "Error" for i in self.issues)

    def to_frame(self) -> pd.DataFrame:
        return pd.DataFrame([{
            "Level": i.level, "Field": i.field, "Message": i.message, "Count": i.count,
            "Rows (first 20)": ", ".join(map(str, i.rows[:20])),
        } for i in self.issues], columns=["Level", "Field", "Message", "Count", "Rows (first 20)"])


def _has_value(s: pd.Series) -> pd.Series:
    txt = s.astype("object").map(lambda v: "" if pd.isna(v) else str(v).strip())
    return ~txt.isin(["", "nan", "None", "NaN", "NaT", "<NA>"])


def validate(raw: pd.DataFrame) -> ValidationReport:
    issues: list[ValidationIssue] = []
    mapping = map_columns(raw.columns)
    present = set(mapping.values())
    missing = [c for c in IMPORT_REQUIRED_COLUMNS if c not in present]

    if raw.empty:
        issues.append(ValidationIssue("Error", "*", "الملف لا يحتوي على أي سجلات."))
        return ValidationReport(issues, mapping, missing, 0)
    for col in missing:
        level = "Error" if col in STRUCTURALLY_REQUIRED else "Warning"
        issues.append(ValidationIssue(level, col, f"العمود المطلوب '{col}' غير موجود في الملف."
                                      + ("" if level == "Error" else " سيتم إنشاؤه فارغًا.")))
    unmapped = [c for c in raw.columns if c not in mapping]
    if unmapped:
        issues.append(ValidationIssue("Info", "*", "أعمدة إضافية لم يتم التعرف عليها وسيتم الاحتفاظ بها: "
                                      + ", ".join(map(str, unmapped[:15])), len(unmapped)))
    if any(i.level == "Error" for i in issues):
        return ValidationReport(issues, mapping, missing, len(raw))

    df = normalize(raw)

    def add(level, fld, msg, mask):
        n = int(mask.sum())
        if n:
            issues.append(ValidationIssue(level, fld, msg, n, df.loc[mask, "row_number"].tolist()))

    add("Warning", "employee_id", "Employee ID مكرر (سيتم استبعاد التكرارات من التحليل الهيكلي).",
        df["employee_id"].ne("") & df["employee_id"].duplicated(keep=False))
    add("Warning", "employee_id", "سجل موظف (وظيفة مشغولة) بدون Employee ID.",
        ~df["is_vacant"] & df["employee_id"].eq(""))
    ids = set(df["node_id"])
    add("Warning", "manager_id", "Manager ID غير موجود ضمن أرقام الموظفين.",
        df["manager_id"].ne("") & ~df["manager_id"].isin(ids))
    add("Warning", "manager_id", "الموظف مسجل كمدير لنفسه.", df["manager_id"].eq(df["node_id"]))
    roots = df["is_active"] & df["manager_id"].eq("")
    if roots.sum() == 0:
        issues.append(ValidationIssue("Warning", "manager_id",
                                      "لا يوجد موظف نشط بدون مدير؛ لا يمكن تحديد الـCEO وسيتعذر حساب العمق."))
    elif roots.sum() > 1:
        add("Info", "manager_id", "أكثر من موظف نشط بدون مدير؛ سيتم اختيار الـCEO تلقائيًا "
            "واعتبار الباقين Orphan Employees.", roots)

    raw_norm = raw.rename(columns=mapping)
    for col in ("salary", "job_level", "age"):
        if col in raw_norm.columns:
            had_value = _has_value(raw_norm[col])
            add("Warning", col, f"قيم غير رقمية في {col} تم تحويلها إلى فارغ.",
                had_value.values & df[col].isna().values)
    if "hire_date" in raw_norm.columns:
        add("Warning", "hire_date", "تواريخ توظيف غير قابلة للقراءة.",
            _has_value(raw_norm["hire_date"]).values & df["hire_date"].isna().values)
    add("Warning", "employment_status",
        f"حالة وظيفية غير معروفة (المسموح: {', '.join(EMPLOYMENT_STATUSES)}).",
        ~df["is_vacant"] & ~df["employment_status"].isin(EMPLOYMENT_STATUSES))
    add("Warning", "position_status",
        f"Position Status غير معروف (المسموح: {', '.join(POSITION_STATUSES)}).",
        ~df["position_status"].isin(POSITION_STATUSES))
    if not issues or all(i.level == "Info" for i in issues):
        issues.append(ValidationIssue("Info", "*", "لم يتم العثور على مشكلات تمنع التحليل."))
    return ValidationReport(issues, mapping, missing, len(raw))


# ================================================================================================
# Dummy data generator  (data/dummy_generator.py)
# ================================================================================================
SAMPLE_DIR = Path(__file__).resolve().parent / "data" / "sample"


# =========================================================================== blueprints
@dataclass
class Unit:
    """A position template. ``children`` is a list of (count, Unit)."""

    title: str
    level: int
    children: list = field(default_factory=list)
    ics: object = 0  # int | (lo, hi) | [per-instance counts]
    ic_titles: list[str] | None = None


@dataclass
class Dept:
    name: str
    division: str
    business_unit: str
    head: Unit
    ic_titles: list[str]
    vacancy_rate: float
    salary_factor: float
    locations: dict[str, float]
    remote_bias: float = 0.1  # share of Remote among IC work arrangements
    onsite_bias: float = 0.6


def U(title, level, children=(), ics=0, ic_titles=None) -> Unit:
    return Unit(title, level, list(children), ics, ic_titles)


HQ = {"Riyadh": 1.0}
DEPARTMENTS: list[Dept] = [
    # ---------------------------------------------------------------- Human Resources
    Dept("HR Operations", "Human Resources", "Corporate Services",
         U("Director, HR Operations", 2, [(3, U("HR Operations Manager", 4, ics=(15, 18)))]),
         ["HR Officer", "HR Specialist", "HR Analyst", "HR Analyst", "Senior HR Analyst",
          "HR Coordinator"], 0.03, 1.0, {"Riyadh": 0.7, "Jeddah": 0.2, "Dammam": 0.1}),
    # HRIS: deliberately manager-heavy (high management ratio, low spans).
    Dept("HRIS", "Human Resources", "Corporate Services",
         U("Director, HRIS", 2, [(2, U("HRIS Senior Manager", 3,
            [([2, 1], U("HRIS Manager", 4, [(1, U("HRIS Supervisor", 5, ics=[3, 2, 1]))]))]))]),
         ["HRIS Analyst", "HRIS Analyst", "HR Data Analyst", "HRIS Specialist", "Hris Analyst"],
         0.0, 1.1, HQ, remote_bias=0.2, onsite_bias=0.3),
    Dept("Talent Acquisition", "Human Resources", "Corporate Services",
         U("Director, Talent Acquisition", 2, [(2, U("Talent Acquisition Manager", 4, ics=(14, 17)))]),
         ["Recruiter", "Senior Recruiter", "Talent Acquisition Specialist",
          "Talent Acquisition Specialist", "Recruitment Specialist"], 0.08, 0.95,
         {"Riyadh": 0.6, "Jeddah": 0.25, "Dubai": 0.15}),
    Dept("Learning & Development", "Human Resources", "Corporate Services",
         U("Director, Learning & Development", 2, [(2, U("L&D Manager", 4, ics=(8, 10)))]),
         ["Learning & Development Specialist", "Training Specialist", "L&D Specialist",
          "Instructional Designer"], 0.05, 0.95, HQ),
    Dept("Total Rewards", "Human Resources", "Corporate Services",
         U("Head of Total Rewards", 2, ics=2), ["Compensation Analyst", "Benefits Analyst"],
         0.0, 1.05, HQ),
    # ---------------------------------------------------------------- Finance
    Dept("Accounting", "Finance", "Corporate Services",
         U("Director, Accounting", 2, [(2, U("Senior Accounting Manager", 3,
            [(2, U("Accounting Manager", 4, [(3, U("Accounting Supervisor", 5, ics=(5, 8)))]))]))]),
         ["Accountant", "Senior Accountant", "Staff Accountant", "Accounts Payable Accountant",
          "Accounts Receivable Accountant"], 0.04, 1.0, {"Riyadh": 0.85, "Jeddah": 0.15}),
    Dept("Financial Planning", "Finance", "Corporate Services",
         U("Director, Financial Planning", 2, [(2, U("FP&A Manager", 4, ics=(11, 13)))]),
         ["Financial Analyst", "Senior Financial Analyst", "FP&A Analyst"], 0.05, 1.1, HQ),
    Dept("Treasury", "Finance", "Corporate Services",
         U("Treasury Director", 2, ics=1), ["Treasury Analyst"], 0.0, 1.15, HQ),
    Dept("Procurement", "Finance", "Corporate Services",
         U("Director, Procurement", 2, [(3, U("Procurement Manager", 4, ics=[12, 1, 13]))]),
         ["Procurement Specialist", "Buyer", "Senior Buyer", "Category Specialist"], 0.06, 0.95,
         {"Riyadh": 0.6, "Dammam": 0.4}),
    # ---------------------------------------------------------------- Information Technology
    Dept("Infrastructure", "Information Technology", "Technology",
         U("Director, Infrastructure", 2, [(2, U("Senior Infrastructure Manager", 3,
            [(3, U("Infrastructure Manager", 4, ics=(9, 12)))]))]),
         ["Systems Engineer", "Senior Systems Engineer", "Network Engineer", "Cloud Engineer",
          "IT Support Specialist"], 0.07, 1.15, {"Riyadh": 0.8, "Dammam": 0.2},
         remote_bias=0.15, onsite_bias=0.4),
    Dept("Cybersecurity", "Information Technology", "Technology",
         U("Director, Cybersecurity", 2, [(2, U("Senior Security Manager", 3,
            [([2, 1], U("Security Manager", 4, ics=(7, 9)))]))]),
         ["Security Analyst", "Senior Security Analyst", "SOC Analyst", "Security Engineer"],
         0.15, 1.25, HQ, remote_bias=0.1, onsite_bias=0.5),
    Dept("Applications", "Information Technology", "Technology",
         U("Director, Applications", 2, [(3, U("Senior Applications Manager", 3,
            [(3, U("Applications Manager", 4, [(2, U("Team Lead", 5, ics=(5, 8)))]))]))]),
         ["Software Engineer", "Senior Software Engineer", "Software Developer", "QA Engineer",
          "Business Analyst"], 0.06, 1.2, {"Riyadh": 0.7, "Cairo": 0.3},
         remote_bias=0.3, onsite_bias=0.25),
    Dept("Data & Analytics", "Information Technology", "Technology",
         U("Director, Data & Analytics", 2, [(3, U("Analytics Manager", 4, ics=[10, 11, 3]))]),
         ["Data Analyst", "Senior Data Analyst", "BI Analyst", "Data Engineer", "Data Scientist"],
         0.08, 1.2, HQ, remote_bias=0.25, onsite_bias=0.3),
    # ---------------------------------------------------------------- Operations
    # Operations: one deliberately deep branch (VP -> Sr Director -> Director -> ...)
    # next to a flatter plant branch with very wide supervisor spans.
    Dept("Operations", "Operations", "Operations",
         U("VP Operations", 2, [
             (1, U("Senior Director, Field Operations", 2, [
                 (2, U("Director, Regional Operations", 2, [
                     (2, U("Senior Operations Manager", 3, [
                         (3, U("Operations Manager", 4, [
                             (3, U("Operations Supervisor", 5, ics=(8, 11)))]))]))]))])),
             (1, U("Director, Plant Operations", 2, [
                 (3, U("Plant Manager", 4, [
                     (2, U("Shift Supervisor", 5, ics=[14, 11, 20, 12, 22, 10]))]))])),
         ]),
         ["Operations Technician", "Senior Operations Technician", "Machine Operator",
          "Maintenance Technician", "Operations Associate"], 0.02, 0.8,
         {"Dammam": 0.55, "Jubail": 0.45}, remote_bias=0.0, onsite_bias=0.97),
    # Quality: flat - the director manages most ICs directly (high span).
    Dept("Quality", "Operations", "Operations",
         U("Director, Quality", 2, [(1, U("Quality Manager", 4, ics=7))], ics=17),
         ["Quality Inspector", "Quality Engineer", "Quality Analyst", "QA/QC Technician"],
         0.04, 0.9, {"Dammam": 0.6, "Jubail": 0.4}, remote_bias=0.0, onsite_bias=0.9),
    Dept("Supply Chain", "Operations", "Operations",
         U("Director, Supply Chain", 2, [(2, U("Senior Supply Chain Manager", 3,
            [(2, U("Supply Chain Manager", 4, [(3, U("Warehouse Supervisor", 5,
                                                   ics=[9, 11, 20, 10, 12, 8, 15, 10, 11, 9, 12, 10]))]))]))]),
         ["Warehouse Associate", "Logistics Coordinator", "Inventory Specialist",
          "Forklift Operator", "Supply Chain Analyst"], 0.03, 0.8,
         {"Dammam": 0.5, "Jeddah": 0.3, "Riyadh": 0.2}, remote_bias=0.0, onsite_bias=0.9),
    Dept("Customer Operations", "Operations", "Operations",
         U("Director, Customer Operations", 2, [(3, U("Customer Operations Manager", 4,
            [(3, U("Customer Service Supervisor", 5, ics=[12, 10, 34, 11, 9, 13, 12, 11, 10]))]))]),
         ["Customer Service Agent", "Senior Customer Service Agent",
          "Customer Support Specialist", "Call Center Agent"], 0.05, 0.7,
         {"Riyadh": 0.5, "Cairo": 0.5}, remote_bias=0.15, onsite_bias=0.6),
    # ---------------------------------------------------------------- Commercial
    Dept("Sales", "Commercial", "Commercial",
         U("Director, Sales", 2, [(3, U("Regional Sales Director", 3,
            [(3, U("Sales Manager", 4, ics=(8, 11)))]))]),
         ["Sales Representative", "Senior Sales Representative", "Account Executive",
          "Key Account Manager", "Sales Executive"], 0.05, 1.0,
         {"Riyadh": 0.4, "Jeddah": 0.35, "Dubai": 0.25}, remote_bias=0.05, onsite_bias=0.5),
    Dept("Marketing", "Commercial", "Commercial",
         U("Director, Marketing", 2, [(3, U("Marketing Manager", 4, ics=(9, 11)))]),
         ["Marketing Specialist", "Digital Marketing Specialist", "Content Specialist",
          "Marketing Analyst", "Brand Specialist"], 0.06, 1.0, {"Riyadh": 0.7, "Dubai": 0.3},
         remote_bias=0.2, onsite_bias=0.3),
    Dept("Business Development", "Commercial", "Commercial",
         U("Director, Business Development", 2, ics=11),
         ["Business Development Executive", "Partnership Specialist",
          "Business Development Analyst"], 0.0, 1.1, {"Riyadh": 0.6, "Dubai": 0.4}),
    # ---------------------------------------------------------------- Governance (CEO direct)
    Dept("Internal Audit", "Governance", "Corporate Services",
         U("Head of Internal Audit", 2, ics=2), ["Internal Auditor"], 0.0, 1.1, HQ),
    Dept("Strategy Office", "Governance", "Corporate Services",
         U("Strategy Director", 2), [], 0.0, 1.2, HQ),
]

C_SUITE = {
    "Human Resources": "Chief Human Resources Officer",
    "Finance": "Chief Financial Officer",
    "Information Technology": "Chief Information Officer",
    "Operations": "Chief Operating Officer",
    "Commercial": "Chief Commercial Officer",
}
EXEC_DEPT = ("Executive Office", "Executive Office", "Corporate Services")
ESG_DEPT = ("ESG & Sustainability", "Governance", "Corporate Services")

BASE_SALARY = {0: 1_800_000, 1: 950_000, 2: 480_000, 3: 330_000, 4: 245_000, 5: 170_000, 6: 110_000}
GRADE_RANGE = {0: (20, 20), 1: (18, 19), 2: (16, 17), 3: (14, 15), 4: (12, 13), 5: (10, 11), 6: (5, 9)}
AGE_BY_LEVEL = {0: 56, 1: 51, 2: 46, 3: 42, 4: 39, 5: 36, 6: 31}


# =========================================================================== builder
class _Builder:
    def __init__(self, seed: int, as_of: date, scale: float, locale: str):
        self.rng = np.random.default_rng(seed)
        self.fake = Faker(locale)
        self.fake.seed_instance(seed)
        self.as_of = as_of
        self.scale = scale
        self.rows: list[dict] = []
        self._pos_seq = itertools.count(10001)
        self._instance_counter: dict[tuple, int] = {}

    # ---------------------------------------------------------------- helpers
    def _count(self, spec, key: tuple | None = None, scale: float = 1.0) -> int:
        if isinstance(spec, list):
            i = self._instance_counter.get(key, 0)
            self._instance_counter[key] = i + 1
            n = spec[i % len(spec)]
        elif isinstance(spec, tuple):
            n = int(self.rng.integers(spec[0], spec[1] + 1))
        else:
            n = int(spec)
        return max(0, int(round(n * scale))) if n else 0

    def _person(self, level: int) -> dict:
        gender = "Male" if self.rng.random() < 0.64 else "Female"
        first = self.fake.first_name_male() if gender == "Male" else self.fake.first_name_female()
        age = int(np.clip(self.rng.normal(AGE_BY_LEVEL[level], 5.5), 21, 64))
        max_yos = max(0, min(age - 21, 28))
        yos = float(np.clip(self.rng.gamma(2.2, 2.6 + (6 - level) * 0.9), 0.1, max_yos or 0.1))
        hire = self.as_of - timedelta(days=int(yos * 365.25))
        return {
            "employee_name": f"{first} {self.fake.last_name()}",
            "gender": gender,
            "age": age,
            "hire_date": pd.Timestamp(hire),
            "years_of_service": round(yos, 1),
        }

    def _add(self, *, title, level, dept_name, division, bu, manager_id, location,
             salary_factor=1.0, vacant=False, work=None, emp_type=None) -> dict:
        pos_id = f"P{next(self._pos_seq)}"
        row = {c: None for c in CANONICAL_COLUMNS}
        row.update(job_title=title, job_level=level, department=dept_name, division=division,
                   business_unit=bu, manager_id=manager_id, location=location, position_id=pos_id)
        lo, hi = GRADE_RANGE[level]
        grade = int(self.rng.integers(lo, hi + 1))
        if level == 6 and title.lower().startswith("senior"):
            grade = min(hi + 1, grade + 1)
        row["grade"] = f"G{grade}"
        if vacant:
            row.update(position_status="Vacant", employee_id="", employee_name="",
                       employment_status="", employment_type="", work_arrangement="")
        else:
            row.update(self._person(level))
            row["position_status"] = "Filled"
            row["employment_status"] = "On Leave" if self.rng.random() < 0.035 else "Active"
            row["employment_type"] = emp_type or (
                self.rng.choice(["Full-Time", "Contractor", "Part-Time"], p=[0.86, 0.10, 0.04])
                if level == 6 else "Full-Time")
            row["work_arrangement"] = work or "On-site"
            sal = (BASE_SALARY[level] * salary_factor * self.rng.lognormal(0, 0.10)
                   * (1 + 0.012 * row["years_of_service"]))
            if level == 6 and title.lower().startswith("senior"):
                sal *= 1.2
            row["salary"] = float(round(sal, -2))
            row["performance_rating"] = int(self.rng.choice([1, 2, 3, 4, 5],
                                                             p=[0.04, 0.14, 0.50, 0.24, 0.08]))
        self.rows.append(row)
        return row

    def _loc(self, dept: Dept, parent_loc: str | None) -> str:
        if parent_loc and self.rng.random() < 0.8 and parent_loc in dept.locations:
            return parent_loc
        names = list(dept.locations)
        p = np.array(list(dept.locations.values()), dtype=float)
        return str(self.rng.choice(names, p=p / p.sum()))

    def _work(self, dept: Dept, level: int) -> str:
        if level <= 3:
            return "On-site" if self.rng.random() < 0.7 else "Hybrid"
        r = self.rng.random()
        if r < dept.remote_bias:
            return "Remote"
        return "On-site" if r < dept.remote_bias + dept.onsite_bias else "Hybrid"

    def _node_id(self, row: dict) -> str:
        return row["employee_id"] or f"VAC-{row['position_id']}"

    # ---------------------------------------------------------------- recursion
    def _build_unit(self, unit: Unit, dept: Dept, manager_id: str, parent_loc: str | None,
                    is_head: bool) -> None:
        loc = self._loc(dept, parent_loc)
        row = self._add(title=unit.title, level=unit.level, dept_name=dept.name,
                        division=dept.division, bu=dept.business_unit, manager_id=manager_id,
                        location=loc, salary_factor=dept.salary_factor,
                        work=self._work(dept, unit.level))
        row["employee_id"] = self._next_emp_id()
        my_id = row["employee_id"]
        child_scale = self.scale if is_head else 1.0
        for count_spec, child in unit.children:
            n = self._count(count_spec, ("c", id(count_spec)), child_scale)
            for _ in range(n):
                self._build_unit(child, dept, my_id, loc, False)
        # individual contributors
        ic_scale = self.scale if (is_head and not unit.children) else 1.0
        n_ic = self._count(unit.ics, ("i", id(unit)), ic_scale)
        titles = unit.ic_titles or dept.ic_titles
        for _ in range(n_ic):
            vacant = self.rng.random() < dept.vacancy_rate
            ic = self._add(title=str(self.rng.choice(titles)), level=6, dept_name=dept.name,
                           division=dept.division, bu=dept.business_unit, manager_id=my_id,
                           location=self._loc(dept, loc), salary_factor=dept.salary_factor,
                           vacant=vacant, work=self._work(dept, 6))
            if not vacant:
                ic["employee_id"] = self._next_emp_id()

    def _next_emp_id(self) -> str:
        return self._emp_ids.pop()

    # ---------------------------------------------------------------- top level
    def build(self) -> pd.DataFrame:
        # Pre-shuffle a pool of employee IDs so IDs do not reveal the structure.
        pool = np.arange(10001, 10001 + int(3000 * max(self.scale, 1.0)) + 2000)
        self.rng.shuffle(pool)
        self._emp_ids = [f"E{n}" for n in pool]

        ex_name, ex_div, ex_bu = EXEC_DEPT
        ceo = self._add(title="Chief Executive Officer", level=0, dept_name=ex_name,
                        division=ex_div, bu=ex_bu, manager_id="", location="Riyadh",
                        work="On-site", emp_type="Full-Time")
        ceo["employee_id"] = self._next_emp_id()
        ceo_id = ceo["employee_id"]
        for _ in range(2):
            ea = self._add(title="Executive Assistant", level=6, dept_name=ex_name,
                           division=ex_div, bu=ex_bu, manager_id=ceo_id, location="Riyadh",
                           work="On-site", emp_type="Full-Time")
            ea["employee_id"] = self._next_emp_id()

        for division, title in C_SUITE.items():
            cx = self._add(title=title, level=1, dept_name=ex_name, division=ex_div, bu=ex_bu,
                           manager_id=ceo_id, location="Riyadh", work="On-site",
                           emp_type="Full-Time")
            cx["employee_id"] = self._next_emp_id()
            ea = self._add(title="Executive Assistant", level=6, dept_name=ex_name,
                           division=ex_div, bu=ex_bu, manager_id=cx["employee_id"],
                           location="Riyadh", work="On-site", emp_type="Full-Time")
            ea["employee_id"] = self._next_emp_id()
            for dept in (d for d in DEPARTMENTS if d.division == division):
                self._build_unit(dept.head, dept, cx["employee_id"], "Riyadh", True)

        for dept in (d for d in DEPARTMENTS if d.division == "Governance"):
            self._build_unit(dept.head, dept, ceo_id, "Riyadh", True)

        # ESG & Sustainability: an "empty" department - only vacant positions.
        esg_name, esg_div, esg_bu = ESG_DEPT
        head = self._add(title="Head of ESG & Sustainability", level=2, dept_name=esg_name,
                         division=esg_div, bu=esg_bu, manager_id=ceo_id, location="Riyadh",
                         vacant=True)
        self._add(title="ESG Analyst", level=6, dept_name=esg_name, division=esg_div,
                  bu=esg_bu, manager_id=self._node_id(head), location="Riyadh", vacant=True)

        df = pd.DataFrame(self.rows, columns=CANONICAL_COLUMNS)
        return df


# =========================================================================== injections
def _gen_set_manager(df: pd.DataFrame, idx, manager_id: str, name_lookup: dict) -> None:
    df.loc[idx, "manager_id"] = manager_id
    df.loc[idx, "manager_name"] = name_lookup.get(manager_id, "")


def _inject_issues(df: pd.DataFrame, rng: np.random.Generator, fake: Faker) -> list[dict]:
    """Inject deliberate anomalies. Returns the manifest of what was injected."""
    manifest: list[dict] = []
    names = dict(zip(df["employee_id"], df["employee_name"]))
    used: set = set()

    def pick_ics(dept: str, n: int, exclude_mgrs: set | None = None) -> list:
        filled = df["position_status"].eq("Filled")
        cand = df.index[filled & df["department"].eq(dept) & df["job_level"].eq(6)
                        & ~df.index.isin(list(used))
                        & ~df["employee_id"].isin(set(df["manager_id"]))]
        if exclude_mgrs:
            cand = cand[~df.loc[cand, "manager_id"].isin(exclude_mgrs)]
        chosen = list(rng.choice(cand, size=n, replace=False))
        used.update(chosen)
        return chosen

    # 1. Orphans: no Manager ID -------------------------------------------------------
    for dept in ["Sales", "Infrastructure", "Supply Chain"]:
        (i,) = pick_ics(dept, 1)
        df.loc[i, ["manager_id", "manager_name"]] = ["", ""]
        manifest.append({"issue": "Missing Manager ID", "expected_finding": "ORPHAN_EMPLOYEE",
                         "ids": [df.at[i, "employee_id"]], "department": dept})

    # 2. Manager ID that does not exist ----------------------------------------------
    for k, dept in enumerate(["Talent Acquisition", "Applications", "Customer Operations"]):
        (i,) = pick_ics(dept, 1)
        ghost = f"E{90001 + k}"
        df.loc[i, ["manager_id", "manager_name"]] = [ghost, fake.name()]
        manifest.append({"issue": "Manager ID not found", "expected_finding": "INVALID_MANAGER",
                         "ids": [df.at[i, "employee_id"]], "department": dept,
                         "detail": ghost})

    # 3. Managers who are terminated but still have reports ---------------------------
    for dept, title in [("Customer Operations", "Customer Service Supervisor"),
                        ("Marketing", "Marketing Manager")]:
        cand = df.index[df["department"].eq(dept) & df["job_title"].eq(title)]
        # choose a supervisor with a normal team (keep the 34-span one intact)
        sizes = df["manager_id"].value_counts()
        cand = [c for c in cand if 3 <= sizes.get(df.at[c, "employee_id"], 0) <= 13]
        i = cand[0]
        df.loc[i, "employment_status"] = "Terminated"  # position left as "Filled" (stale)
        reports = df.index[df["manager_id"].eq(df.at[i, "employee_id"])
                           & df["position_status"].eq("Filled")].tolist()
        used.update(reports)
        manifest.append({"issue": "Reports point to a terminated manager",
                         "expected_finding": "TERMINATED_MANAGER",
                         "ids": df.loc[reports, "employee_id"].tolist(), "department": dept,
                         "detail": df.at[i, "employee_id"]})

    # 4. Circular reporting A -> B -> C -> A and a self-reporting record ---------------
    a, b, c = pick_ics("Marketing", 3)
    ida, idb, idc = (df.at[x, "employee_id"] for x in (a, b, c))
    _gen_set_manager(df, a, idb, names)
    _gen_set_manager(df, b, idc, names)
    _gen_set_manager(df, c, ida, names)
    manifest.append({"issue": "Circular reporting (3-cycle)",
                     "expected_finding": "CIRCULAR_REPORTING", "ids": [ida, idb, idc],
                     "department": "Marketing"})
    (s,) = pick_ics("Data & Analytics", 1)
    _gen_set_manager(df, s, df.at[s, "employee_id"], names)
    manifest.append({"issue": "Self-reporting (1-cycle)", "expected_finding": "CIRCULAR_REPORTING",
                     "ids": [df.at[s, "employee_id"]], "department": "Data & Analytics"})

    # 5. Duplicate Position IDs -------------------------------------------------------
    for dept in ["Accounting", "Sales"]:
        x, y = pick_ics(dept, 2)
        df.loc[y, "position_id"] = df.at[x, "position_id"]
        manifest.append({"issue": "Duplicate Position ID", "expected_finding": "DUPLICATE_POSITION",
                         "ids": [df.at[x, "employee_id"], df.at[y, "employee_id"]],
                         "department": dept, "detail": df.at[x, "position_id"]})

    # 6. A team lead whose whole team is vacant (manager-level, zero active reports) --
    tls = df.index[df["job_title"].eq("Team Lead")]
    tl = tls[0]
    team = df.index[df["manager_id"].eq(df.at[tl, "employee_id"])]
    df.loc[team, ["position_status", "employee_id", "employee_name", "employment_status",
                  "employment_type", "work_arrangement"]] = ["Vacant", "", "", "", "", ""]
    df.loc[team, ["salary", "age", "performance_rating", "years_of_service"]] = np.nan
    df.loc[team, "hire_date"] = pd.NaT
    manifest.append({"issue": "Manager-level position with no active reports (team vacant)",
                     "expected_finding": "VACANT_POSITION", "ids": [df.at[tl, "employee_id"]],
                     "department": "Applications", "detail": f"{len(team)} vacant positions"})

    # 7. Data-quality defects ---------------------------------------------------------
    (i,) = pick_ics("HR Operations", 1)
    df.loc[i, "job_title"] = "hr analyst"  # case variant of an existing title
    manifest.append({"issue": "Job title case variant ('hr analyst')",
                     "expected_finding": "DUPLICATE_POSITION", "ids": [df.at[i, "employee_id"]],
                     "department": "HR Operations"})
    for dept in ["Procurement", "Quality"]:
        (i,) = pick_ics(dept, 1)
        df.loc[i, "job_title"] = ""
        manifest.append({"issue": "Missing job title", "expected_finding": "DATA_QUALITY",
                         "ids": [df.at[i, "employee_id"]], "department": dept})
    for dept in ["Supply Chain", "Customer Operations", "Sales"]:
        (i,) = pick_ics(dept, 1)
        df.loc[i, "salary"] = np.nan
        manifest.append({"issue": "Missing salary", "expected_finding": "DATA_QUALITY",
                         "ids": [df.at[i, "employee_id"]], "department": dept})
    (i,) = pick_ics("Operations", 1)
    df.loc[i, "salary"] = 0.0
    manifest.append({"issue": "Salary = 0", "expected_finding": "DATA_QUALITY",
                     "ids": [df.at[i, "employee_id"]], "department": "Operations"})
    (i,) = pick_ics("Applications", 1)
    df.loc[i, "department"] = ""
    manifest.append({"issue": "Missing department", "expected_finding": "DATA_QUALITY",
                     "ids": [df.at[i, "employee_id"]], "department": "Applications"})

    # 8. Duplicate employee ID (a second record re-using an existing ID) --------------
    (i,) = pick_ics("Accounting", 1)
    dup = df.loc[[i]].copy()
    dup["employee_name"] = fake.name()
    dup["position_id"] = "P99001"
    df_dup_idx = len(df)
    df.loc[df_dup_idx] = dup.iloc[0].values
    manifest.append({"issue": "Duplicate Employee ID", "expected_finding": "DATA_QUALITY",
                     "ids": [df.at[i, "employee_id"]], "department": "Accounting"})
    return manifest


# =========================================================================== public API
def generate_org(n_employees: int = 1500, seed: int = 42, as_of: date | None = None,
                 locale: str = "en_US", inject_issues: bool = True
                 ) -> tuple[pd.DataFrame, list[dict]]:
    """Build a synthetic organisation of roughly ``n_employees`` active employees.

    Returns ``(dataframe, manifest)`` - the manifest lists every injected issue.
    """
    as_of = as_of or date.today()
    base = 1500
    scale = max(0.3, n_employees / base)
    builder = _Builder(seed, as_of, scale, locale)
    df = builder.build()
    manifest = _inject_issues(df, builder.rng, builder.fake) if inject_issues else []
    # Manager Name is looked up from the (possibly stale) manager record.
    names = dict(zip(df["employee_id"], df["employee_name"]))
    keep = df["manager_name"].notna() & df["manager_name"].ne("")
    df.loc[~keep, "manager_name"] = df.loc[~keep, "manager_id"].map(names).fillna("")
    vac_names = {f"VAC-{p}": "(Vacant)" for p in df.loc[df["position_status"].eq("Vacant"),
                                                         "position_id"]}
    df.loc[df["manager_id"].isin(vac_names), "manager_name"] = "(Vacant)"
    df = df.reset_index(drop=True)
    df["job_level"] = df["job_level"].astype("Int64")
    for col in ("age", "performance_rating"):
        df[col] = pd.to_numeric(df[col]).astype("Int64")
    df["years_of_service"] = pd.to_numeric(df["years_of_service"])
    df["salary"] = pd.to_numeric(df["salary"])
    return df, manifest


def save_sample(df: pd.DataFrame, out_dir: Path = SAMPLE_DIR) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    csv_path = out_dir / "org_dummy_data.csv"
    xlsx_path = out_dir / "org_dummy_data.xlsx"
    df.to_csv(csv_path, index=False, encoding="utf-8-sig")
    df.to_excel(xlsx_path, index=False, engine="openpyxl")
    return {"csv": csv_path, "xlsx": xlsx_path}


# ================================================================================================
# Hierarchy engine  (analytics/hierarchy.py)
# ================================================================================================
@dataclass
class OrgModel:
    """Result of :func:`build_org_model`. ``df`` is the enriched dataset."""

    df: pd.DataFrame
    graph: nx.DiGraph
    ceo_id: str | None
    cycles: list[list[str]] = field(default_factory=list)
    root_candidates: list[str] = field(default_factory=list)

    # -------------------------------------------------------------- convenience
    @property
    def active(self) -> pd.DataFrame:
        return self.df[self.df["is_active"]]

    @property
    def managers(self) -> pd.DataFrame:
        return self.df[self.df["is_manager"]]

    def node(self, node_id: str) -> pd.Series | None:
        idx = self._index().get(node_id)
        return None if idx is None else self.df.loc[idx]

    def _index(self) -> dict:
        if not hasattr(self, "_idx_cache"):
            first = self.df[~self.df["is_duplicate_id"]]
            self._idx_cache = dict(zip(first["node_id"], first.index))
        return self._idx_cache

    def chain_up(self, node_id: str, max_steps: int = 200) -> list[str]:
        """Node IDs from ``node_id`` up the declared manager chain.

        Follows ``manager_id`` (not the graph) so broken chains are visible, and
        stops at a repeated node (cycle), a missing manager, or ``max_steps``.
        """
        chain, seen = [], set()
        idx = self._index()
        cur = node_id
        while cur and cur not in seen and len(chain) < max_steps:
            chain.append(cur)
            seen.add(cur)
            row_idx = idx.get(cur)
            if row_idx is None:
                break
            cur = self.df.at[row_idx, "manager_id"]
        if cur and cur in seen:
            chain.append(cur)  # show where the loop closes
        return chain

    def direct_report_ids(self, node_id: str, include_vacant: bool = False) -> list[str]:
        if node_id not in self.graph:
            return []
        out = []
        for child in self.graph.successors(node_id):
            if include_vacant or self.graph.nodes[child].get("active"):
                out.append(child)
        return out

    def subtree_ids(self, node_id: str, include_root: bool = True,
                    include_vacant: bool = False) -> list[str]:
        if node_id not in self.graph:
            return [node_id] if include_root else []
        ids = nx.descendants(self.graph, node_id)
        if not include_vacant:
            ids = {i for i in ids if self.graph.nodes[i].get("active")}
        if include_root:
            ids = ids | {node_id}
        return list(ids)


# ============================================================================ build
def _find_ceo(df: pd.DataFrame, graph: nx.DiGraph, in_cycle: set[str]) -> tuple[str | None, list]:
    roots = df[df["is_active"] & df["manager_id"].eq("") & ~df["node_id"].isin(in_cycle)]
    candidates = roots["node_id"].tolist()
    if not candidates:
        return None, []
    if len(candidates) == 1:
        return candidates[0], candidates
    # Prefer the lowest job level (0 = CEO), then the largest organisation.
    sizes = {c: len(nx.descendants(graph, c)) for c in candidates}
    title_ceo = roots["job_title"].str.contains(r"\b(?:CEO|Chief Executive)", case=False, regex=True)
    ranked = roots.assign(
        _lvl=roots["job_level"].fillna(99),
        _ceo=~title_ceo,
        _size=[-sizes[c] for c in roots["node_id"]],
    ).sort_values(["_lvl", "_ceo", "_size"])
    return ranked["node_id"].iloc[0], candidates


def build_org_model(df_norm: pd.DataFrame) -> OrgModel:
    """Build the reporting graph and enrich every record with structural metrics.

    Input must come from :func:`data.loader.normalize`.
    """
    df = df_norm.copy()
    base = df[~df["is_duplicate_id"]]
    status_of = dict(zip(base["node_id"], base["employment_status"]))
    active_of = dict(zip(base["node_id"], base["is_active"]))
    vacant_of = dict(zip(base["node_id"], base["is_vacant"]))
    in_graph = base[base["is_active"] | base["is_vacant"]]

    g = nx.DiGraph()
    for nid, act in zip(in_graph["node_id"], in_graph["is_active"]):
        g.add_node(nid, active=bool(act))
    for nid, mid in zip(in_graph["node_id"], in_graph["manager_id"]):
        if mid and mid in g:  # manager exists and is active / vacant position
            g.add_edge(mid, nid)

    # ---------------------------------------------------------------- cycles
    cycles: list[list[str]] = []
    in_cycle: set[str] = set()
    for comp in nx.strongly_connected_components(g):
        if len(comp) > 1:
            cyc = next(nx.simple_cycles(g.subgraph(comp)))
            cycles.append(_rotate_cycle(cyc))
            in_cycle |= set(comp)
    for n in list(nx.nodes_with_selfloops(g)):
        cycles.append([n])
        in_cycle.add(n)

    ceo_id, root_candidates = _find_ceo(base, g, in_cycle)

    # ---------------------------------------------------------------- depth (BFS from CEO)
    depth: dict[str, int] = {}
    if ceo_id is not None:
        tree = g.copy()
        tree.remove_edges_from(nx.selfloop_edges(tree))
        depth = nx.single_source_shortest_path_length(tree, ceo_id)

    # ---------------------------------------------------------------- reporting status
    def status(row) -> str:
        nid, mid = row.node_id, row.manager_id
        if row.is_duplicate_id:
            return RS_DUPLICATE
        if row.is_vacant:
            return RS_VACANT
        if not row.is_active:
            return RS_INACTIVE
        if nid == ceo_id:
            return RS_CEO
        if nid in in_cycle:
            return RS_CIRCULAR
        if not mid:
            return RS_MISSING
        if mid not in status_of:
            return RS_INVALID
        if not active_of.get(mid) and not vacant_of.get(mid):
            return RS_TERMINATED
        return RS_VALID if nid in depth else RS_DISCONNECTED

    df["reporting_status"] = [status(r) for r in df.itertuples(index=False)]
    df["depth"] = df["node_id"].map(depth).astype("float")
    df.loc[df["is_duplicate_id"], "depth"] = np.nan
    df["manager_status"] = df["manager_id"].map(status_of).fillna("")

    # ---------------------------------------------------------------- spans & subtree sizes
    active_nodes = {n for n, a in g.nodes(data="active") if a}
    direct_active: dict[str, int] = {}
    direct_vacant: dict[str, int] = {}
    for u, v in g.edges():
        if u == v:
            continue
        if v in active_nodes:
            direct_active[u] = direct_active.get(u, 0) + 1
        else:
            direct_vacant[u] = direct_vacant.get(u, 0) + 1
    total = _subtree_active_counts(g, active_nodes, in_cycle)

    df["direct_reports"] = df["node_id"].map(direct_active).fillna(0).astype(int)
    df["vacant_direct_positions"] = df["node_id"].map(direct_vacant).fillna(0).astype(int)
    df["total_reports"] = df["node_id"].map(total).fillna(0).astype(int)
    not_node = ~df["is_active"]
    df.loc[not_node & ~df["is_vacant"], ["direct_reports", "vacant_direct_positions",
                                         "total_reports"]] = 0
    df.loc[df["is_duplicate_id"], ["direct_reports", "vacant_direct_positions",
                                   "total_reports"]] = 0
    df["is_manager"] = df["is_active"] & df["direct_reports"].gt(0)
    df["layer"] = df["depth"].map(lambda d: f"Layer {int(d)}" if pd.notna(d) else "Unreachable")

    return OrgModel(df=df, graph=g, ceo_id=ceo_id, cycles=cycles,
                    root_candidates=root_candidates)


def _rotate_cycle(cyc: list[str]) -> list[str]:
    """Express the loop in reporting order (employee -> manager -> ...), starting at
    its smallest id so results are deterministic. Graph edges run manager -> employee."""
    cyc = list(reversed(cyc))
    i = cyc.index(min(cyc))
    return cyc[i:] + cyc[:i]


def _subtree_active_counts(g: nx.DiGraph, active_nodes: set[str],
                           in_cycle: set[str]) -> dict[str, int]:
    """Number of active descendants per node (excluding itself).

    Cycles are condensed first so the computation always terminates.
    """
    cond = nx.condensation(g)
    members = cond.graph["mapping"]  # node -> component id
    comp_active = {c: sum(1 for n in data["members"] if n in active_nodes)
                   for c, data in cond.nodes(data=True)}
    below: dict[int, int] = {}
    for c in reversed(list(nx.topological_sort(cond))):
        below[c] = sum(below[s] + comp_active[s] for s in cond.successors(c))
    out = {}
    for n in g.nodes:
        c = members[n]
        own = comp_active[c] - (1 if n in active_nodes else 0)
        out[n] = below[c] + own  # members of the same cycle count as reports
    return out


# ============================================================================ queries
def depth_stats(model: OrgModel, df: pd.DataFrame | None = None) -> dict:
    """Min / max / mean / median depth of active employees (CEO excluded)."""
    d = (df if df is not None else model.df)
    d = d[d["is_active"] & d["depth"].notna() & d["node_id"].ne(model.ceo_id)]["depth"]
    if d.empty:
        return {"min": np.nan, "max": np.nan, "mean": np.nan, "median": np.nan, "count": 0}
    return {"min": int(d.min()), "max": int(d.max()), "mean": float(d.mean()),
            "median": float(d.median()), "count": int(d.size)}


def layer_count(df: pd.DataFrame) -> int:
    """Number of structural layers = distinct depths among reachable active employees."""
    d = df.loc[df["is_active"], "depth"].dropna()
    return int(d.nunique()) if not d.empty else 0


def longest_chains(model: OrgModel, top: int = 10, df: pd.DataFrame | None = None) -> pd.DataFrame:
    """The deepest reporting chains (one per leaf employee), with the path."""
    d = df if df is not None else model.df
    d = d[d["is_active"] & d["depth"].notna()].sort_values(
        ["depth", "node_id"], ascending=[False, True]).head(top)
    names = dict(zip(model.df["node_id"], model.df["job_title"]))
    rows = []
    for r in d.itertuples():
        chain = model.chain_up(r.node_id)
        rows.append({
            "employee_id": r.node_id,
            "employee_name": r.employee_name,
            "job_title": r.job_title,
            "department": r.department,
            "depth": int(r.depth),
            "chain": " → ".join(f"{names.get(n, n)}" for n in reversed(chain)),
            "chain_ids": " → ".join(reversed(chain)),
        })
    return pd.DataFrame(rows)


def department_heads(model: OrgModel) -> dict[str, str]:
    """Department -> node_id of its head (top-most active member, largest team)."""
    df = model.df
    act = df[df["is_active"] & df["department"].ne("")]
    dept_of = dict(zip(df["node_id"], df["department"]))
    heads: dict[str, str] = {}
    for dept, grp in act.groupby("department"):
        tops = grp[grp["manager_id"].map(dept_of).ne(dept)]
        if tops.empty:
            tops = grp
        tops = tops.sort_values(["depth", "total_reports"], ascending=[True, False],
                                na_position="last")
        heads[dept] = tops["node_id"].iloc[0]
    return heads


def level_depth_matrix(df: pd.DataFrame) -> pd.DataFrame:
    """Cross-tab of declared Job Level vs measured depth (active employees)."""
    a = df[df["is_active"] & df["depth"].notna() & df["job_level"].notna()]
    if a.empty:
        return pd.DataFrame()
    return pd.crosstab(a["job_level"].astype(int), a["depth"].astype(int))


# ================================================================================================
# Span of Control analytics  (analytics/span_of_control.py)
# ================================================================================================
SPAN_LOW = "Low"
SPAN_HIGH = "High"
SPAN_OK = "Within Range"


def span_category(n: int, settings: Settings) -> str:
    if n <= settings.low_span_threshold:
        return SPAN_LOW
    if n > settings.high_span_threshold:
        return SPAN_HIGH
    return SPAN_OK


def manager_table(df: pd.DataFrame, settings: Settings) -> pd.DataFrame:
    """One row per manager with span metrics and category."""
    m = df[df["is_manager"]].copy()
    cols = ["node_id", "employee_name", "job_title", "job_level", "department", "division",
            "business_unit", "location", "depth", "direct_reports", "vacant_direct_positions",
            "total_reports", "salary"]
    m = m[cols]
    m["span_category"] = m["direct_reports"].map(lambda n: span_category(int(n), settings))
    return m.sort_values("direct_reports", ascending=False).reset_index(drop=True)


def span_stats(spans: pd.Series) -> dict:
    s = pd.Series(spans, dtype=float).dropna()
    if s.empty:
        return {k: np.nan for k in ("min", "max", "mean", "median", "p25", "p75")} | {"count": 0}
    return {
        "count": int(s.size),
        "min": float(s.min()),
        "max": float(s.max()),
        "mean": float(s.mean()),
        "median": float(s.median()),
        "p25": float(s.quantile(0.25)),
        "p75": float(s.quantile(0.75)),
    }


def span_distribution(df: pd.DataFrame) -> pd.DataFrame:
    s = df.loc[df["is_manager"], "direct_reports"]
    return s.value_counts().sort_index().rename_axis("direct_reports").reset_index(name="managers")


def span_by(df: pd.DataFrame, by: str, settings: Settings) -> pd.DataFrame:
    """Span statistics grouped by a column (department, job_level, ...)."""
    m = df[df["is_manager"]]
    if m.empty:
        return pd.DataFrame(columns=[by, "managers", "avg_span", "median_span", "min_span",
                                     "max_span", "low_span_managers", "high_span_managers"])
    key = blank_to_label(m[by])
    g = m.groupby(key)["direct_reports"]
    out = pd.DataFrame({
        "managers": g.size(),
        "avg_span": g.mean().round(2),
        "median_span": g.median(),
        "min_span": g.min(),
        "max_span": g.max(),
        "low_span_managers": g.apply(lambda s: int((s <= settings.low_span_threshold).sum())),
        "high_span_managers": g.apply(lambda s: int((s > settings.high_span_threshold).sum())),
    })
    return out.rename_axis(by).reset_index()


# ================================================================================================
# Department / layer / job-title analytics  (analytics/departments.py)
# ================================================================================================
UNIT_COLUMNS = [
    "headcount", "managers", "management_ratio", "avg_span", "avg_depth", "max_depth",
    "layers", "vacancies", "positions", "vacancy_rate", "avg_salary", "median_salary",
    "invalid_reporting",
]


def unit_metrics(df: pd.DataFrame, by: str = "department", settings: Settings | None = None
                 ) -> pd.DataFrame:
    """Metrics per organisational unit (``by`` = department / division / business_unit).

    * headcount         active employees (Active + On Leave)
    * managers          active employees with >= 1 active direct report
    * management_ratio  managers / headcount
    * avg_span          mean direct reports of the unit's managers
    * avg_depth         mean distance to the CEO of the unit's employees
    * layers            number of distinct depth levels present in the unit
    * vacancies         positions with Position Status = Vacant
    * vacancy_rate      vacancies / (vacancies + filled positions of active employees)
    """
    settings = settings or Settings()
    d = df.assign(_unit=blank_to_label(df[by]))
    act = d[d["is_active"]]
    vac = d[d["is_vacant"] & ~d["is_duplicate_id"]]
    units = sorted(set(act["_unit"]) | set(vac["_unit"]))
    if not units:
        return pd.DataFrame(columns=[by] + UNIT_COLUMNS)

    ga = act.groupby("_unit")
    mgr = act[act["is_manager"]].groupby("_unit")["direct_reports"]
    out = pd.DataFrame(index=pd.Index(units, name=by))
    out["headcount"] = ga.size()
    out["managers"] = act.groupby("_unit")["is_manager"].sum()
    out = out.fillna({"headcount": 0, "managers": 0}).astype({"headcount": int, "managers": int})
    out["management_ratio"] = [safe_div(m, h) for m, h in zip(out["managers"], out["headcount"])]
    out["avg_span"] = mgr.mean().round(2)
    out["median_span"] = mgr.median()
    out["avg_depth"] = ga["depth"].mean().round(2)
    out["max_depth"] = ga["depth"].max()
    out["layers"] = ga["depth"].nunique()
    out["vacancies"] = vac.groupby("_unit").size()
    out["vacancies"] = out["vacancies"].fillna(0).astype(int)
    out["positions"] = out["headcount"] + out["vacancies"]
    out["vacancy_rate"] = [safe_div(v, p) for v, p in zip(out["vacancies"], out["positions"])]
    out["avg_salary"] = ga["salary"].mean().round(0)
    out["median_salary"] = ga["salary"].median()
    out["invalid_reporting"] = act.groupby("_unit")["reporting_status"].apply(
        lambda s: int(s.isin(INVALID_REPORTING_STATUSES).sum()))
    out["invalid_reporting"] = out["invalid_reporting"].fillna(0).astype(int)
    out["layers"] = out["layers"].fillna(0).astype(int)
    out["low_span_managers"] = mgr.apply(lambda s: int((s <= settings.low_span_threshold).sum()))
    out["high_span_managers"] = mgr.apply(lambda s: int((s > settings.high_span_threshold).sum()))
    out[["low_span_managers", "high_span_managers"]] = (
        out[["low_span_managers", "high_span_managers"]].fillna(0).astype(int))
    if by == "department":
        parents = d.groupby("_unit")[["division", "business_unit"]].agg(
            lambda s: s.mode().iloc[0] if not s.mode().empty else "")
        out = out.join(parents)
    return out.reset_index().sort_values("headcount", ascending=False).reset_index(drop=True)


def layer_metrics(df: pd.DataFrame, basis: str = "job_level") -> pd.DataFrame:
    """Per-layer summary. ``basis`` = 'job_level' (declared) or 'depth' (measured)."""
    act = df[df["is_active"]]
    total = len(act)
    key = act[basis]
    rows = []
    for layer, grp in act.groupby(key.fillna(-1).astype(int)):
        mg = grp[grp["is_manager"]]
        if basis == "job_level":
            name = LEVEL_NAMES.get(layer, "Unknown") if layer >= 0 else "Unknown level"
        else:
            name = f"Layer {layer}" if layer >= 0 else "Unreachable"
        rows.append({
            "layer": layer if layer >= 0 else None,
            "layer_name": name,
            "employees": len(grp),
            "pct_of_total": safe_div(len(grp), total),
            "managers": len(mg),
            "avg_salary": grp["salary"].mean(),
            "avg_span": mg["direct_reports"].mean() if len(mg) else np.nan,
        })
    out = pd.DataFrame(rows)
    if out.empty:
        return pd.DataFrame(columns=["layer", "layer_name", "employees", "pct_of_total",
                                     "managers", "avg_salary", "avg_span"])
    return out.sort_values("layer", na_position="last").reset_index(drop=True)


# ============================================================================ job titles
_SENIORITY = {"senior", "sr", "junior", "jr", "staff", "lead", "principal", "chief", "head",
              "assistant", "associate", "i", "ii", "iii"}
_STOP = {"&", "and", "of", "the", "-", ","}


def normalize_title(title: str) -> str:
    t = re.sub(r"[^\w&/ ]+", " ", str(title).lower())
    return re.sub(r"\s+", " ", t).strip()


def _core(title_norm: str) -> str:
    return " ".join(w for w in title_norm.split() if w not in _SENIORITY)


def _acronym_form(title_norm: str) -> str:
    """'learning & development specialist' -> 'ld specialist'; 'l&d specialist' -> same."""
    words = [w for w in title_norm.split() if w not in _STOP]
    if len(words) < 2:
        return title_norm
    head = "".join(w.replace("&", "") if "&" in w else w[0] for w in words[:-1])
    return f"{head} {words[-1]}"


def title_similarity(a: str, b: str) -> float:
    """Similarity (0-1) of two job titles.

    Character similarity alone over-matches ("Brand Specialist" ~ "HR Specialist"),
    so a pair only scores when the titles are structurally related:
      * acronym variants ("L&D Specialist" ~ "Learning & Development Specialist") -> 0.95
      * one title's words are contained in the other's ("HR Analyst" ~ "HR Data Analyst")
      * exactly one word differs and those words are themselves alike ("HR" ~ "HRIS")
    Otherwise the score is 0.
    """
    na, nb = normalize_title(a), normalize_title(b)
    if na == nb:
        return 1.0
    if _acronym_form(na) == _acronym_form(nb) and len(na.split()) != len(nb.split()):
        return 0.95
    ta = [w for w in na.split() if w not in _STOP]
    tb = [w for w in nb.split() if w not in _STOP]
    sa, sb = set(ta), set(tb)
    ratio = SequenceMatcher(None, na, nb).ratio()
    if sa < sb or sb < sa:
        return ratio
    da, db = sa - sb, sb - sa
    if len(da) == 1 and len(db) == 1 and len(ta) == len(tb):
        wa, wb = next(iter(da)), next(iter(db))
        if SequenceMatcher(None, wa, wb).ratio() >= 0.5:
            return ratio
    return 0.0


def job_title_analysis(df: pd.DataFrame, settings: Settings | None = None) -> dict:
    """Duplicate titles, raw variants of the same title, and similar-title pairs.

    Similar titles are compared within the same division (overlapping roles
    inside one function are the ones worth reviewing).
    """
    settings = settings or Settings()
    rows = df[(df["is_active"] | df["is_vacant"]) & df["job_title"].ne("")]
    counts = (rows.groupby("job_title")
              .agg(positions=("node_id", "size"),
                   departments=("department", lambda s: ", ".join(sorted(set(s) - {""}))),
                   n_departments=("department", lambda s: len(set(s) - {""})))
              .reset_index().sort_values("positions", ascending=False))
    counts["normalized"] = counts["job_title"].map(normalize_title)

    # Titles that are identical after normalisation but typed differently
    variants = []
    for norm, grp in counts.groupby("normalized"):
        if len(grp) > 1:
            variants.append({"normalized_title": norm,
                             "variants": " | ".join(grp["job_title"]),
                             "positions": int(grp["positions"].sum()),
                             "departments": ", ".join(sorted(set(", ".join(grp["departments"])
                                                                 .split(", ")) - {""}))})
    variants_df = pd.DataFrame(variants, columns=["normalized_title", "variants", "positions",
                                                  "departments"])

    # Similar (not identical) titles, within each division
    norms = dict(zip(counts["job_title"], counts["normalized"]))
    pos = dict(zip(counts["job_title"], counts["positions"]))
    depts = dict(zip(counts["job_title"], counts["departments"]))
    thr = settings.title_similarity_threshold
    seen: set[tuple[str, str]] = set()
    pairs = []
    for division, grp in rows.groupby(blank_to_label(rows["division"])):
        titles = list(dict.fromkeys(t for t in grp["job_title"]))
        titles = sorted({norms[t]: t for t in titles}.values())[:1500]
        for a, b in combinations(titles, 2):
            key = tuple(sorted((a, b)))
            if key in seen:
                continue
            na, nb = norms[a], norms[b]
            if _core(na) == _core(nb):
                continue  # same job family, different seniority: a career ladder
            score = title_similarity(a, b)
            if score >= thr:
                seen.add(key)
                pairs.append({"title_a": a, "title_b": b, "similarity": round(score, 3),
                              "division": division,
                              "positions_a": pos[a], "positions_b": pos[b],
                              "departments_a": depts[a], "departments_b": depts[b]})
    similar_df = pd.DataFrame(pairs, columns=["title_a", "title_b", "similarity", "division",
                                              "positions_a", "positions_b", "departments_a",
                                              "departments_b"])
    similar_df = similar_df.sort_values("similarity", ascending=False).reset_index(drop=True)
    return {"titles": counts.reset_index(drop=True), "variants": variants_df,
            "similar": similar_df}


def duplicate_position_ids(df: pd.DataFrame) -> pd.DataFrame:
    d = df[df["position_id"].ne("") & ~df["is_duplicate_id"]]
    dup = d[d["position_id"].duplicated(keep=False)]
    return dup.sort_values("position_id")


# ================================================================================================
# X-Ray Findings engine  (analytics/anomalies.py)
# ================================================================================================
FINDING_COLUMNS = [
    "finding_id", "type", "severity", "entity_type", "entity_id", "entity_name", "department",
    "description", "evidence", "recommended_investigation", "rule", "metric", "value",
    "threshold", "benchmark", "related_ids",
]

TYPE_LABELS_AR = {
    "LOW_SPAN_OF_CONTROL": "Span of Control منخفض",
    "HIGH_SPAN_OF_CONTROL": "Span of Control مرتفع",
    "DEEP_HIERARCHY": "تسلسل إداري عميق",
    "ORPHAN_EMPLOYEE": "موظف بدون مدير",
    "INVALID_MANAGER": "مدير غير موجود",
    "TERMINATED_MANAGER": "مدير منتهي الخدمة",
    "CIRCULAR_REPORTING": "علاقة إدارية دائرية",
    "SMALL_DEPARTMENT": "وحدة تنظيمية صغيرة",
    "HIGH_MANAGEMENT_RATIO": "نسبة مديرين مرتفعة",
    "VACANT_POSITION": "وظيفة شاغرة",
    "DUPLICATE_POSITION": "وظائف مكررة / متشابهة",
}


@dataclass
class Finding:
    type: str
    severity: str
    entity_type: str
    entity_id: str
    entity_name: str
    department: str
    description: str
    evidence: str
    recommended_investigation: str
    rule: str
    metric: str
    value: object
    threshold: object = None
    benchmark: str = ""
    related_ids: list[str] = field(default_factory=list)
    finding_id: str = ""


def _name(row) -> str:
    return f"{row['employee_name']} ({row['job_title']})" if row["employee_name"] else row["node_id"]


def _lvl(v) -> str:
    return "—" if pd.isna(v) else str(int(v))


# ============================================================================ detectors
def detect_span(model: OrgModel, s: Settings) -> list[Finding]:
    df = model.df
    mgrs = df[df["is_manager"]]
    if mgrs.empty:
        return []
    median = float(mgrs["direct_reports"].median())
    # Loop members are reported as CIRCULAR_REPORTING; their "spans" are artefacts.
    scope = mgrs[mgrs["reporting_status"].ne(RS_CIRCULAR)]
    if s.exclude_executives_from_span:
        scope = scope[~scope["job_level"].isin([0, 1])]
    out: list[Finding] = []
    for _, r in scope.iterrows():
        n = int(r["direct_reports"])
        common = (f"الوسيط التنظيمي لـSpan of Control = {median:g}. "
                  f"إجمالي التابعين (مباشر وغير مباشر) = {int(r['total_reports'])}. "
                  f"وظائف شاغرة تابعة مباشرة = {int(r['vacant_direct_positions'])}. "
                  f"Job Level = {_lvl(r['job_level'])}.")
        if n <= s.low_span_threshold:
            out.append(Finding(
                type="LOW_SPAN_OF_CONTROL",
                severity="Attention" if n == 1 else "Informational",
                entity_type="Manager", entity_id=r["node_id"], entity_name=_name(r),
                department=r["department"],
                description=(f"تم اكتشاف Span of Control منخفض ويحتاج إلى مراجعة: لدى هذا المدير "
                             f"{n} موظف/موظفين مباشرين، مقارنة بالحد المحدد حاليًا وهو "
                             f"{s.low_span_threshold}."),
                evidence=f"عدد الموظفين المباشرين النشطين = {n}. " + common,
                recommended_investigation=(
                    "مراجعة سبب محدودية نطاق الإشراف: هل الدور تخصصي أو قيادي فني؟ هل توجد وظائف "
                    "شاغرة ستُشغل قريبًا؟ هل يمكن دمج الفريق مع فريق مجاور؟ النتيجة لا تعني أن الدور "
                    "غير ضروري."),
                rule=f"direct_reports <= Low Span Threshold ({s.low_span_threshold})",
                metric="Direct Reports", value=n, threshold=s.low_span_threshold,
                benchmark=f"Organizational median = {median:g}",
                related_ids=[r["node_id"]] + model.direct_report_ids(r["node_id"])))
        elif n > s.high_span_threshold:
            ratio = n / s.high_span_threshold
            sev = "Attention" if ratio <= 1.5 else ("Significant" if ratio <= 2.5 else "Critical")
            out.append(Finding(
                type="HIGH_SPAN_OF_CONTROL", severity=sev,
                entity_type="Manager", entity_id=r["node_id"], entity_name=_name(r),
                department=r["department"],
                description=(f"تم اكتشاف Span of Control مرتفع: لدى هذا المدير {n} موظفًا مباشرًا "
                             f"مقارنة بالحد المحدد حاليًا وهو {s.high_span_threshold} "
                             f"وبوسيط تنظيمي يبلغ {median:g}."),
                evidence=(f"هذا المدير لديه {n} موظفًا مباشرًا مقارنة بوسيط تنظيمي يبلغ {median:g} "
                          f"(أي {n / median:.1f}× الوسيط). " + common) if median else common,
                recommended_investigation=(
                    "مراجعة طبيعة العمل (عمل متجانس/تشغيلي قد يتحمل نطاقًا أوسع)، وجود قادة فرق "
                    "غير رسميين، وعبء الإشراف والتقييم على المدير."),
                rule=f"direct_reports > High Span Threshold ({s.high_span_threshold})",
                metric="Direct Reports", value=n, threshold=s.high_span_threshold,
                benchmark=f"Organizational median = {median:g}",
                related_ids=[r["node_id"]] + model.direct_report_ids(r["node_id"])))
    return out


def detect_deep_hierarchy(model: OrgModel, s: Settings) -> list[Finding]:
    df = model.df
    act = df[df["is_active"] & df["depth"].notna()]
    if act.empty:
        return []
    org_median = float(act["depth"].median())
    org_max = int(act["depth"].max())
    titles = dict(zip(df["node_id"], df["job_title"]))
    out = []
    for dept, grp in act.groupby("department"):
        deep = grp[grp["depth"] > s.deep_hierarchy_threshold]
        if deep.empty:
            continue
        max_d = int(grp["depth"].max())
        deepest = grp.sort_values(["depth", "node_id"], ascending=[False, True]).iloc[0]
        chain = list(reversed(model.chain_up(deepest["node_id"])))
        chain_txt = " → ".join(titles.get(n, n) or n for n in chain)
        excess = max_d - s.deep_hierarchy_threshold
        out.append(Finding(
            type="DEEP_HIERARCHY", severity="Attention" if excess == 1 else "Significant",
            entity_type="Department", entity_id=dept or "(Unassigned)",
            entity_name=dept or "(Unassigned)", department=dept,
            description=(f"تم اكتشاف تسلسل إداري عميق (Deep Reporting Chain): أقصى عمق في الإدارة "
                         f"{max_d} علاقات إدارية حتى الـCEO، مقارنة بالحد المحدد {s.deep_hierarchy_threshold}."),
            evidence=(f"{len(deep)} موظف/موظفين على عمق أكبر من {s.deep_hierarchy_threshold}. "
                      f"عدد الطبقات في الإدارة = {grp['depth'].nunique()}. "
                      f"الوسيط التنظيمي للعمق = {org_median:g}، والحد الأقصى على مستوى المنظمة = {org_max}. "
                      f"أطول سلسلة: {chain_txt}."),
            recommended_investigation=(
                "مراجعة ما إذا كانت كل طبقة تضيف نطاق مسؤولية مختلفًا، وهل توجد مسميات متتالية "
                "بنفس المستوى الوظيفي (مثل Senior Director ثم Director)، وأثر العمق على سرعة القرار."),
            rule=f"employee depth > Deep Hierarchy Threshold ({s.deep_hierarchy_threshold})",
            metric="Max Depth", value=max_d, threshold=s.deep_hierarchy_threshold,
            benchmark=f"Organizational median depth = {org_median:g}",
            related_ids=deep["node_id"].tolist()))
    return out


def detect_reporting(model: OrgModel, s: Settings) -> list[Finding]:
    df = model.df
    names = dict(zip(df["node_id"], df["employee_name"]))
    out = []
    for _, r in df[df["reporting_status"].eq(RS_MISSING)].iterrows():
        out.append(Finding(
            type="ORPHAN_EMPLOYEE", severity="Significant", entity_type="Employee",
            entity_id=r["node_id"], entity_name=_name(r), department=r["department"],
            description="موظف بدون Manager ID وليس هو الـCEO (Invalid Reporting Relationship).",
            evidence=(f"Manager ID فارغ. الـCEO المحدد في البيانات = {model.ceo_id}. "
                      f"عدد التابعين لهذا الموظف = {int(r['total_reports'])}. "
                      f"رقم السجل في الملف = {r['row_number']}."),
            recommended_investigation="التحقق من المدير المباشر الفعلي وتحديث Manager ID في نظام الموارد البشرية.",
            rule="manager_id is blank AND employee is not the CEO",
            metric="Manager ID", value="(blank)", related_ids=[r["node_id"]]))
    for _, r in df[df["reporting_status"].eq(RS_INVALID)].iterrows():
        out.append(Finding(
            type="INVALID_MANAGER", severity="Significant", entity_type="Employee",
            entity_id=r["node_id"], entity_name=_name(r), department=r["department"],
            description=f"Manager ID ({r['manager_id']}) غير موجود في البيانات (Invalid Reporting Relationship).",
            evidence=(f"Manager ID = {r['manager_id']}، Manager Name المسجل = "
                      f"'{r['manager_name'] or '—'}'. لا يوجد أي سجل بهذا الرقم. "
                      f"رقم السجل في الملف = {r['row_number']}."),
            recommended_investigation="التحقق من صحة رقم المدير (خطأ إدخال، مدير لم يُنقل للنظام، أو رقم قديم).",
            rule="manager_id not found among Employee IDs",
            metric="Manager ID", value=r["manager_id"], related_ids=[r["node_id"]]))
    term = df[df["reporting_status"].eq(RS_TERMINATED)]
    for mid, grp in term.groupby("manager_id"):
        mrow = df[df["node_id"].eq(mid)].iloc[0]
        out.append(Finding(
            type="TERMINATED_MANAGER", severity="Significant", entity_type="Manager",
            entity_id=mid, entity_name=_name(mrow), department=mrow["department"],
            description=(f"{len(grp)} موظف/موظفين ما زالوا يتبعون مديرًا حالته الوظيفية "
                         f"'{mrow['employment_status']}' (Invalid Reporting Relationship)."),
            evidence=(f"المدير {mid} – {names.get(mid, '')}: Employment Status = "
                      f"{mrow['employment_status']}، Position Status = {mrow['position_status']}. "
                      f"الموظفون المتأثرون: {', '.join(grp['node_id'].head(15))}"
                      f"{' …' if len(grp) > 15 else ''}."),
            recommended_investigation=("تحديد المدير البديل أو المكلف، وتحديث Manager ID للموظفين، "
                                       "ومراجعة حالة الوظيفة (قد تكون شاغرة فعليًا)."),
            rule="manager record exists but Employment Status is not Active / On Leave",
            metric="Affected employees", value=len(grp),
            related_ids=[mid] + grp["node_id"].tolist()))
    return out


def detect_cycles(model: OrgModel, s: Settings) -> list[Finding]:
    df = model.df
    names = dict(zip(df["node_id"], df["employee_name"]))
    depts = dict(zip(df["node_id"], df["department"]))
    out = []
    for cyc in model.cycles:
        path = " → ".join(f"{n} ({names.get(n, '')})" for n in cyc + [cyc[0]])
        downstream = set()
        for n in cyc:
            downstream |= set(model.subtree_ids(n, include_root=False))
        downstream -= set(cyc)
        out.append(Finding(
            type="CIRCULAR_REPORTING", severity="Critical", entity_type="Reporting Loop",
            entity_id=" → ".join(cyc), entity_name=f"{len(cyc)}-employee loop",
            department=depts.get(cyc[0], ""),
            description=("تم اكتشاف علاقة إدارية دائرية: سلسلة المديرين تعود إلى نفس الموظف ولا تصل "
                         "إلى الـCEO." if len(cyc) > 1 else
                         "موظف مسجل كمدير لنفسه (Manager ID = Employee ID)."),
            evidence=(f"المسار: {path}. عدد الموظفين خارج الحلقة المتأثرين (لا يصلون إلى الـCEO) = "
                      f"{len(downstream)}."),
            recommended_investigation="تصحيح Manager ID لأحد أعضاء الحلقة على الأقل بحيث تصل السلسلة إلى الـCEO.",
            rule="strongly connected component (size > 1) or self-loop in the reporting graph",
            metric="Loop length", value=len(cyc), related_ids=list(cyc) + sorted(downstream)))
    return out


def detect_units(model: OrgModel, s: Settings) -> list[Finding]:
    df = model.df
    units = unit_metrics(df, "department", s)
    act = df[df["is_active"]]
    org_ratio = act["is_manager"].sum() / len(act) if len(act) else np.nan
    median_hc = float(units["headcount"].median()) if not units.empty else np.nan
    out = []
    for _, u in units.iterrows():
        dept = u["department"]
        members = df[(df["department"].eq(dept) | ((dept == "(Unassigned)") & df["department"].eq("")))
                     & (df["is_active"] | df["is_vacant"])]["node_id"].tolist()
        hc = int(u["headcount"])
        if hc <= s.small_department_threshold:
            out.append(Finding(
                type="SMALL_DEPARTMENT", severity="Attention" if hc == 0 else "Informational",
                entity_type="Department", entity_id=dept, entity_name=dept, department=dept,
                description=(f"Small Organizational Unit: عدد الموظفين النشطين في الإدارة = {hc}، "
                             f"مقارنة بالحد المحدد {s.small_department_threshold}."),
                evidence=(f"Headcount = {hc}، الوظائف الشاغرة = {int(u['vacancies'])}، "
                          f"عدد المديرين = {int(u['managers'])}. وسيط حجم الإدارات = {median_hc:g}."),
                recommended_investigation=("مراجعة ما إذا كانت الوحدة وظيفة متخصصة مستقلة بطبيعتها، أو "
                                           "قيد التأسيس، أو يمكن ضمها لوحدة ذات صلة."),
                rule=f"department headcount <= Small Department Threshold ({s.small_department_threshold})",
                metric="Headcount", value=hc, threshold=s.small_department_threshold,
                benchmark=f"Median department headcount = {median_hc:g}", related_ids=members))
        ratio = u["management_ratio"]
        if hc > s.small_department_threshold and pd.notna(ratio) and ratio > s.high_management_ratio:
            out.append(Finding(
                type="HIGH_MANAGEMENT_RATIO",
                severity="Attention" if ratio <= 2 * s.high_management_ratio else "Significant",
                entity_type="Department", entity_id=dept, entity_name=dept, department=dept,
                description=(f"نسبة المديرين في الإدارة {pct(ratio)} مقارنة بالحد المحدد "
                             f"{pct(s.high_management_ratio)}."),
                evidence=(f"Management Ratio = {int(u['managers'])} مدير ÷ {hc} موظف = {pct(ratio)}. "
                          f"النسبة على مستوى المنظمة = {pct(org_ratio)}. متوسط Span of Control في "
                          f"الإدارة = {u['avg_span']:.2f}. عدد الطبقات = {int(u['layers'])}."),
                recommended_investigation=("مراجعة عدد المستويات الإشرافية مقارنة بحجم الفريق، ووجود "
                                           "مديرين بفرق صغيرة جدًا، وإمكانية تبسيط الهيكل."),
                rule=f"managers / headcount > High Management Ratio ({pct(s.high_management_ratio, 0)})",
                metric="Management Ratio", value=round(float(ratio), 4),
                threshold=s.high_management_ratio,
                benchmark=f"Organization ratio = {pct(org_ratio)}", related_ids=members))
    return out


def detect_vacancies(model: OrgModel, s: Settings) -> list[Finding]:
    df = model.df
    out = []
    for _, r in df[df["is_vacant"] & ~df["is_duplicate_id"]].iterrows():
        lvl = r["job_level"]
        has_reports = int(r["direct_reports"]) + int(r["vacant_direct_positions"])
        managerial = pd.notna(lvl) and lvl <= 4
        sev = "Attention" if (managerial or has_reports) else "Informational"
        out.append(Finding(
            type="VACANT_POSITION", severity=sev, entity_type="Position",
            entity_id=r["position_id"] or r["node_id"], entity_name=r["job_title"] or "(no title)",
            department=r["department"],
            description=f"وظيفة شاغرة (Position Status = Vacant): {r['job_title']}.",
            evidence=(f"Position ID = {r['position_id']}، Job Level = {_lvl(lvl)}، تتبع المدير "
                      f"{r['manager_id'] or '—'}. وظائف تابعة لهذه الوظيفة = {has_reports}."),
            recommended_investigation=("التحقق من حالة التوظيف لهذه الوظيفة ومدة شغورها، وأثرها على "
                                       "الفريق." + (" الوظيفة ذات مستوى إشرافي." if managerial else "")),
            rule="position_status == 'Vacant'", metric="Position Status", value="Vacant",
            related_ids=[r["node_id"]]))
    # Manager-level roles whose entire team is vacant (0 active direct reports)
    empty = df[df["is_active"] & df["direct_reports"].eq(0) & df["vacant_direct_positions"].gt(0)]
    for _, r in empty.iterrows():
        out.append(Finding(
            type="VACANT_POSITION", severity="Attention", entity_type="Manager",
            entity_id=r["node_id"], entity_name=_name(r), department=r["department"],
            description=(f"مدير بدون موظفين نشطين: جميع الوظائف التابعة له ({int(r['vacant_direct_positions'])}) شاغرة."),
            evidence=(f"Direct Reports النشطون = 0، الوظائف الشاغرة التابعة = "
                      f"{int(r['vacant_direct_positions'])}، Job Level = {_lvl(r['job_level'])}."),
            recommended_investigation="مراجعة خطة شغل الفريق أو ما إذا كان الفريق قد نُقل/أُعيد تنظيمه.",
            rule="active employee with 0 active and >= 1 vacant direct positions",
            metric="Vacant direct positions", value=int(r["vacant_direct_positions"]),
            related_ids=[r["node_id"]] + model.direct_report_ids(r["node_id"], include_vacant=True)))
    return out


def detect_duplicates(model: OrgModel, s: Settings) -> list[Finding]:
    df = model.df
    out = []
    dups = duplicate_position_ids(df)
    for pid, grp in dups.groupby("position_id"):
        out.append(Finding(
            type="DUPLICATE_POSITION", severity="Significant", entity_type="Position",
            entity_id=pid, entity_name=", ".join(grp["job_title"].unique()),
            department=grp["department"].iloc[0],
            description=f"Position ID {pid} مسجل لأكثر من سجل ({len(grp)} سجلات).",
            evidence="السجلات: " + "; ".join(f"{r.node_id} – {r.employee_name} (row {r.row_number})"
                                             for r in grp.itertuples()),
            recommended_investigation="التحقق من أن كل موظف مرتبط بوظيفة مستقلة في هيكل الوظائف.",
            rule="position_id appears on more than one record", metric="Records", value=len(grp),
            related_ids=grp["node_id"].tolist()))
    jt = job_title_analysis(df, s)
    for v in jt["variants"].itertuples():
        titles = v.variants.split(" | ")
        ids = df[df["job_title"].isin(titles) & (df["is_active"] | df["is_vacant"])]["node_id"].tolist()
        out.append(Finding(
            type="DUPLICATE_POSITION", severity="Attention", entity_type="Job Title",
            entity_id=v.normalized_title, entity_name=v.variants, department=v.departments,
            description="مسمى وظيفي مكرر بصيغ كتابة مختلفة (Duplicate Job Title).",
            evidence=f"الصيغ: {v.variants}. عدد الوظائف = {v.positions}. الإدارات: {v.departments}.",
            recommended_investigation="توحيد المسمى في دليل الوظائف (Job Catalogue).",
            rule="job titles identical after normalising case / spacing / punctuation",
            metric="Variants", value=len(titles), related_ids=ids))
    for p in jt["similar"].itertuples():
        ids = df[df["job_title"].isin([p.title_a, p.title_b])
                 & (df["is_active"] | df["is_vacant"])]["node_id"].tolist()
        out.append(Finding(
            type="DUPLICATE_POSITION", severity="Informational", entity_type="Job Title",
            entity_id=f"{p.title_a} ~ {p.title_b}", entity_name=f"{p.title_a} ~ {p.title_b}",
            department=", ".join(sorted({*p.departments_a.split(", "), *p.departments_b.split(", ")} - {""})),
            description="مسميات وظيفية متشابهة نصيًا (Similar Job Titles) قد تشير إلى تداخل أدوار.",
            evidence=(f"'{p.title_a}' ({p.positions_a} وظيفة؛ {p.departments_a}) و'{p.title_b}' "
                      f"({p.positions_b} وظيفة؛ {p.departments_b}). درجة التشابه = {p.similarity:.2f}."),
            recommended_investigation="مقارنة الوصف الوظيفي للمسميين لتحديد ما إذا كانا دورًا واحدًا أم دورين مختلفين.",
            rule=f"text similarity >= {s.title_similarity_threshold}", metric="Similarity",
            value=p.similarity, threshold=s.title_similarity_threshold, related_ids=ids))
    return out


DETECTORS = [detect_cycles, detect_reporting, detect_span, detect_deep_hierarchy, detect_units,
             detect_duplicates, detect_vacancies]


# ============================================================================ public API
def detect_findings(model: OrgModel, settings: Settings | None = None) -> pd.DataFrame:
    """Run every detector. A failing detector never stops the others."""
    settings = settings or Settings()
    findings: list[Finding] = []
    errors = []
    for det in DETECTORS:
        try:
            findings.extend(det(model, settings))
        except Exception as exc:  # pragma: no cover - defensive
            errors.append(f"{det.__name__}: {exc}")
    type_rank = {t: i for i, t in enumerate(FINDING_TYPES)}
    findings.sort(key=lambda f: (SEVERITY_RANK[f.severity], type_rank.get(f.type, 99),
                                 str(f.department), str(f.entity_id)))
    for i, f in enumerate(findings, 1):
        f.finding_id = f"F-{i:04d}"
    out = pd.DataFrame([asdict(f) for f in findings], columns=FINDING_COLUMNS)
    out.attrs["errors"] = errors
    return out


def explain_finding(findings: pd.DataFrame, finding_id: str, model: OrgModel) -> dict:
    """Everything needed to justify a finding (rule, numbers, source records)."""
    row = findings[findings["finding_id"].eq(finding_id)]
    if row.empty:
        raise KeyError(finding_id)
    f = row.iloc[0].to_dict()
    ids = list(f["related_ids"])
    records = model.df[model.df["node_id"].isin(ids)]
    f["records"] = records
    return f


def filter_findings(findings: pd.DataFrame, node_ids: set[str], departments: set[str]
                    ) -> pd.DataFrame:
    """Keep findings touching the filtered population."""
    if findings.empty:
        return findings
    keep = findings.apply(
        lambda r: bool(set(r["related_ids"]) & node_ids)
        or (r["entity_type"] == "Department" and r["department"] in departments), axis=1)
    return findings[keep]


def summarize_findings(findings: pd.DataFrame) -> pd.DataFrame:
    if findings.empty:
        return pd.DataFrame(columns=["type", "severity", "count"])
    return findings.groupby(["type", "severity"]).size().reset_index(name="count")


# ================================================================================================
# Data Quality engine  (analytics/data_quality.py)
# ================================================================================================
@dataclass
class Check:
    check_id: str
    dimension: str
    name: str
    field: str
    description: str
    scope: Callable[[pd.DataFrame, OrgModel], pd.Series]  # rows the check applies to
    fails: Callable[[pd.DataFrame, OrgModel], pd.Series]  # rows that fail (within scope)


def _filled(df, m):
    return ~df["is_vacant"]


def _all(df, m):
    return pd.Series(True, index=df.index)


def _positions(df, m):
    return ~df["is_duplicate_id"]


def _active(df, m):
    return df["is_active"]


def _non_ceo_filled(df, m):
    return ~df["is_vacant"] & df["node_id"].ne(m.ceo_id) & df["employment_status"].ne("Terminated")


def _blank(col):
    return lambda df, m: df[col].astype(str).str.strip().isin(["", "nan", "NaT", "<NA>", "None"]) \
        | df[col].isna()


def _manager_level_inverted(df, m):
    lvl = dict(zip(df["node_id"], df["job_level"]))
    mgr_lvl = df["manager_id"].map(lvl)
    return mgr_lvl.notna() & df["job_level"].notna() & (mgr_lvl > df["job_level"])


def _dept_multi_division(df, m):
    d = df[df["department"].ne("") & df["division"].ne("")]
    n = d.groupby("department")["division"].nunique()
    bad = set(n[n > 1].index)
    return df["department"].isin(bad)


def _manager_name_mismatch(df, m):
    names = dict(zip(df["node_id"], df["employee_name"]))
    actual = df["manager_id"].map(names)
    recorded = df["manager_name"]
    return actual.notna() & recorded.ne("") & actual.ne("") & recorded.ne("(Vacant)") \
        & actual.str.lower().ne(recorded.str.lower())


def _salary_band_outlier(df, m):
    """Salary outside the 1.5*IQR fences of its job level."""
    s = df["salary"]
    out = pd.Series(False, index=df.index)
    for lvl, grp in df[s.notna() & df["job_level"].notna()].groupby("job_level"):
        q1, q3 = grp["salary"].quantile([0.25, 0.75])
        iqr = q3 - q1
        lo, hi = q1 - 1.5 * iqr, q3 + 1.5 * iqr
        out.loc[grp.index] = (grp["salary"] < lo) | (grp["salary"] > hi)
    return out


CHECKS: list[Check] = [
    # ------------------------------------------------------------------ completeness
    Check("C01", "Completeness", "Employee ID Completeness", "employee_id",
          "Employee ID populated for filled positions", _filled, _blank("employee_id")),
    Check("C02", "Completeness", "Employee Name Completeness", "employee_name",
          "Name populated for filled positions", _filled, _blank("employee_name")),
    Check("C03", "Completeness", "Manager ID Completeness", "manager_id",
          "Manager ID populated for everyone except the CEO", _non_ceo_filled, _blank("manager_id")),
    Check("C04", "Completeness", "Job Title Completeness", "job_title",
          "Job title populated", _positions, _blank("job_title")),
    Check("C05", "Completeness", "Department Completeness", "department",
          "Department populated", _positions, _blank("department")),
    Check("C06", "Completeness", "Division Completeness", "division",
          "Division populated", _positions, _blank("division")),
    Check("C07", "Completeness", "Business Unit Completeness", "business_unit",
          "Business unit populated", _positions, _blank("business_unit")),
    Check("C08", "Completeness", "Job Level Completeness", "job_level",
          "Job level populated", _positions, lambda df, m: df["job_level"].isna()),
    Check("C09", "Completeness", "Salary Completeness", "salary",
          "Salary populated for filled positions", _filled, lambda df, m: df["salary"].isna()),
    Check("C10", "Completeness", "Position ID Completeness", "position_id",
          "Position ID populated", _all, _blank("position_id")),
    Check("C11", "Completeness", "Hire Date Completeness", "hire_date",
          "Hire date populated for filled positions", _filled, lambda df, m: df["hire_date"].isna()),
    # ------------------------------------------------------------------ uniqueness
    Check("U01", "Uniqueness", "Employee ID Uniqueness", "employee_id",
          "Each Employee ID appears once",
          lambda df, m: ~df["is_vacant"] & df["employee_id"].ne(""),
          lambda df, m: df["employee_id"].ne("") & df["employee_id"].duplicated(keep=False)),
    Check("U02", "Uniqueness", "Position ID Uniqueness", "position_id",
          "Each Position ID appears once", lambda df, m: df["position_id"].ne(""),
          lambda df, m: df["position_id"].ne("") & df["position_id"].duplicated(keep=False)),
    # ------------------------------------------------------------------ validity
    Check("V01", "Validity", "Salary Validity", "salary", "Salary is a positive number",
          lambda df, m: ~df["is_vacant"] & df["salary"].notna(),
          lambda df, m: df["salary"] <= 0),
    Check("V02", "Validity", "Job Level Validity", "job_level", "Job level is an integer 0-6",
          lambda df, m: df["job_level"].notna(),
          lambda df, m: ~df["job_level"].between(0, 6) | (df["job_level"] % 1 != 0)),
    Check("V03", "Validity", "Employment Status Validity", "employment_status",
          f"One of {', '.join(EMPLOYMENT_STATUSES)}", _filled,
          lambda df, m: ~df["employment_status"].isin(EMPLOYMENT_STATUSES)),
    Check("V04", "Validity", "Position Status Validity", "position_status",
          f"One of {', '.join(POSITION_STATUSES)}", _all,
          lambda df, m: ~df["position_status"].isin(POSITION_STATUSES)),
    Check("V05", "Validity", "Employment Type Validity", "employment_type",
          f"One of {', '.join(EMPLOYMENT_TYPES)}",
          lambda df, m: ~df["is_vacant"] & df["employment_type"].ne(""),
          lambda df, m: ~df["employment_type"].isin(EMPLOYMENT_TYPES)),
    Check("V06", "Validity", "Work Arrangement Validity", "work_arrangement",
          f"One of {', '.join(WORK_ARRANGEMENTS)}",
          lambda df, m: ~df["is_vacant"] & df["work_arrangement"].ne(""),
          lambda df, m: ~df["work_arrangement"].isin(WORK_ARRANGEMENTS)),
    Check("V07", "Validity", "Age Validity", "age", "Age between 16 and 80",
          lambda df, m: df["age"].notna(), lambda df, m: ~df["age"].between(16, 80)),
    Check("V08", "Validity", "Hire Date Validity", "hire_date", "Hire date is not in the future",
          lambda df, m: df["hire_date"].notna(),
          lambda df, m: df["hire_date"] > pd.Timestamp.today().normalize()),
    Check("V09", "Validity", "Performance Rating Validity", "performance_rating",
          "Rating between 1 and 5", lambda df, m: df["performance_rating"].notna(),
          lambda df, m: ~df["performance_rating"].between(1, 5)),
    # ------------------------------------------------------------------ consistency
    Check("S01", "Consistency", "Terminated Employee / Position Status", "position_status",
          "Terminated employees should not hold a 'Filled' position",
          lambda df, m: df["employment_status"].eq("Terminated"),
          lambda df, m: df["position_status"].eq("Filled")),
    Check("S02", "Consistency", "Manager Job Level Seniority", "job_level",
          "Manager's job level is equal or more senior than the employee's",
          lambda df, m: df["is_active"] & df["manager_id"].ne(""), _manager_level_inverted),
    Check("S03", "Consistency", "Manager-level Role Without Reports", "job_level",
          "Job levels 0-5 are expected to have at least one active direct report",
          lambda df, m: df["is_active"] & df["job_level"].le(5),
          lambda df, m: df["direct_reports"].eq(0)),
    Check("S04", "Consistency", "Department → Division Mapping", "division",
          "Each department belongs to exactly one division", _positions, _dept_multi_division),
    Check("S05", "Consistency", "Manager Name Matches Manager ID", "manager_name",
          "Recorded manager name matches the manager's record",
          lambda df, m: df["manager_id"].ne("") & df["manager_name"].ne(""),
          _manager_name_mismatch),
    Check("S06", "Consistency", "Vacant Position Without Employee Data", "employee_name",
          "Vacant positions carry no employee name",
          lambda df, m: df["is_vacant"], lambda df, m: df["employee_name"].ne("")),
    Check("S07", "Consistency", "Salary Within Job Level Band", "salary",
          "Salary within 1.5×IQR of its job level",
          lambda df, m: df["is_active"] & df["salary"].notna(), _salary_band_outlier),
    # ------------------------------------------------------------------ referential integrity
    Check("R01", "Referential Integrity", "Manager Reference Validity", "manager_id",
          "Manager ID refers to an existing record",
          lambda df, m: df["manager_id"].ne("") & (df["is_active"] | df["is_vacant"]),
          lambda df, m: df["reporting_status"].eq(RS_INVALID)),
    Check("R02", "Referential Integrity", "Manager Is Active", "manager_id",
          "Manager record is Active / On Leave (or a vacant position)",
          lambda df, m: df["is_active"] & df["manager_id"].ne(""),
          lambda df, m: df["reporting_status"].eq(RS_TERMINATED)),
    Check("R03", "Referential Integrity", "No Circular Reporting", "manager_id",
          "Reporting chain does not loop", _active,
          lambda df, m: df["reporting_status"].eq(RS_CIRCULAR)),
    Check("R04", "Referential Integrity", "Reaches the CEO", "manager_id",
          "Reporting chain reaches the top of the hierarchy", _active,
          lambda df, m: df["depth"].isna()),
    Check("R05", "Referential Integrity", "Single Top of Hierarchy", "manager_id",
          "Only one active employee without a manager", lambda df, m: df["is_active"] & df["manager_id"].eq(""),
          lambda df, m: df["node_id"].ne(m.ceo_id)),
]


def run_checks(model: OrgModel, df: pd.DataFrame | None = None) -> pd.DataFrame:
    """Run all checks. Returns one row per check with score and failing node IDs."""
    data = df if df is not None else model.df
    rows = []
    for c in CHECKS:
        try:
            scope = c.scope(data, model).fillna(False).astype(bool)
            fails = c.fails(data, model).fillna(False).astype(bool) & scope
            n_scope, n_fail = int(scope.sum()), int(fails.sum())
            error = ""
        except Exception as exc:  # pragma: no cover - defensive
            n_scope, n_fail, fails, error = 0, 0, pd.Series(False, index=data.index), str(exc)
        score = 1.0 if n_scope == 0 else (n_scope - n_fail) / n_scope
        rows.append({
            "check_id": c.check_id, "dimension": c.dimension, "check": c.name,
            "field": c.field, "rule": c.description, "records_checked": n_scope,
            "failed": n_fail, "score": score,
            "status": "Pass" if n_fail == 0 else ("Warning" if score >= 0.98 else "Fail"),
            "failing_rows": data.index[fails].tolist(), "error": error,
        })
    return pd.DataFrame(rows)


def dimension_scores(checks: pd.DataFrame) -> pd.DataFrame:
    """Average score per dimension (weighted by records checked)."""
    if checks.empty:
        return pd.DataFrame(columns=["dimension", "score", "failed"])
    g = checks.assign(passed=checks["records_checked"] - checks["failed"]).groupby("dimension")
    out = pd.DataFrame({"records": g["records_checked"].sum(), "passed": g["passed"].sum(),
                        "failed": g["failed"].sum()})
    out["score"] = out["passed"] / out["records"].where(out["records"] > 0)
    return out.reset_index()[["dimension", "score", "failed", "records"]]


def overall_score(checks: pd.DataFrame) -> float:
    tot = checks["records_checked"].sum()
    return float((tot - checks["failed"].sum()) / tot) if tot else 1.0


def affected_records(model: OrgModel, checks: pd.DataFrame, check_id: str,
                     df: pd.DataFrame | None = None) -> pd.DataFrame:
    data = df if df is not None else model.df
    row = checks[checks["check_id"].eq(check_id)]
    if row.empty:
        return data.iloc[0:0]
    return data.loc[row.iloc[0]["failing_rows"]]


# ================================================================================================
# What-If Simulator engine  (analytics/simulations.py)
# ================================================================================================
class SimulationError(ValueError):
    """Raised when an action cannot be applied (the message is user-facing)."""


@dataclass
class ScenarioResult:
    df: pd.DataFrame
    model: OrgModel
    log: list[str] = field(default_factory=list)


# ============================================================================ helpers
def _row_ix(df: pd.DataFrame, node_id: str) -> int:
    ix = df.index[df["node_id"].eq(node_id) & ~df["is_duplicate_id"]]
    if len(ix) == 0:
        raise SimulationError(f"الرقم {node_id} غير موجود في البيانات.")
    return ix[0]


def _set_manager(df: pd.DataFrame, rows, new_manager: str) -> None:
    name = df.loc[df["node_id"].eq(new_manager), "employee_name"]
    df.loc[rows, "manager_id"] = new_manager
    df.loc[rows, "manager_name"] = name.iloc[0] if len(name) else ""


def _children_rows(df: pd.DataFrame, node_id: str) -> pd.Index:
    """All records (active or vacant) whose manager is ``node_id``."""
    return df.index[df["manager_id"].eq(node_id) & df["node_id"].ne(node_id)
                    & (df["is_active"] | df["is_vacant"])]


# ============================================================================ actions
def merge_departments(df: pd.DataFrame, model: OrgModel, departments: list[str], new_name: str,
                      consolidate_heads: bool = False, primary_head: str | None = None
                      ) -> tuple[pd.DataFrame, list[str]]:
    departments = [d for d in departments if d]
    if len(departments) < 2:
        raise SimulationError("اختر إدارتين على الأقل للدمج.")
    missing = [d for d in departments if d not in set(df["department"])]
    if missing:
        raise SimulationError(f"إدارات غير موجودة: {', '.join(missing)}")
    new_name = (new_name or " & ".join(departments)).strip()
    df = df.copy()
    log = []
    heads = department_heads(model)
    primary_dept = departments[0]
    mask = df["department"].isin(departments)
    div = df.loc[df["department"].eq(primary_dept), "division"].mode()
    bu = df.loc[df["department"].eq(primary_dept), "business_unit"].mode()
    df.loc[mask, "department"] = new_name
    if len(div):
        df.loc[mask, "division"] = div.iloc[0]
    if len(bu):
        df.loc[mask, "business_unit"] = bu.iloc[0]
    log.append(f"دمج {' + '.join(departments)} → {new_name} ({int(mask.sum())} سجل).")

    if consolidate_heads:
        primary = primary_head or heads.get(primary_dept)
        if primary is None:
            raise SimulationError("تعذر تحديد رئيس الإدارة الأساسية.")
        for d in departments:
            h = heads.get(d)
            if h is None or h == primary:
                continue
            rows = _children_rows(df, h)
            _set_manager(df, rows, primary)
            df = df.drop(index=_row_ix(df, h))
            log.append(f"توحيد القيادة: {len(rows)} سجل كان يتبع {h} أصبح يتبع {primary}، "
                       f"وتم حذف وظيفة رئيس الإدارة {h}.")
    return df, log


def remove_layer(df: pd.DataFrame, model: OrgModel, job_level: int,
                 department: str | None = None) -> tuple[pd.DataFrame, list[str]]:
    """Remove every *manager* at ``job_level`` (optionally within one department).

    Their reports move up to the nearest remaining manager in their chain.
    """
    m = model.df
    target_mask = m["is_manager"] & m["job_level"].eq(job_level) & m["node_id"].ne(model.ceo_id)
    if department:
        target_mask &= m["department"].eq(department)
    targets = set(m.loc[target_mask, "node_id"])
    if not targets:
        raise SimulationError("لا يوجد مديرون في الطبقة/الإدارة المحددة.")
    df = df.copy()
    moved = 0
    new_parent = {}
    for t in targets:
        chain = model.chain_up(t)[1:]
        parent = next((c for c in chain if c not in targets), None)
        new_parent[t] = parent if parent is not None else df.at[_row_ix(df, t), "manager_id"]
    for t in targets:
        rows = _children_rows(df, t)
        rows = [r for r in rows if df.at[r, "node_id"] not in targets]
        _set_manager(df, rows, new_parent[t])
        moved += len(rows)
    df = df[~(df["node_id"].isin(targets) & ~df["is_duplicate_id"])]
    scope = f" في إدارة {department}" if department else ""
    return df, [f"إزالة الطبقة الإدارية Job Level {job_level}{scope}: حذف {len(targets)} وظيفة "
                f"إشرافية ونقل {moved} سجل إلى المدير الأعلى التالي."]


def move_employee(df: pd.DataFrame, model: OrgModel, employee_id: str, new_manager_id: str,
                  move_department: bool = True) -> tuple[pd.DataFrame, list[str]]:
    if employee_id == new_manager_id:
        raise SimulationError("لا يمكن أن يكون الموظف مديرًا لنفسه.")
    e_ix = _row_ix(df, employee_id)
    m_ix = _row_ix(df, new_manager_id)
    if not (df.at[m_ix, "is_active"] or df.at[m_ix, "is_vacant"]):
        raise SimulationError("المدير الجديد ليس موظفًا نشطًا.")
    if new_manager_id in set(model.subtree_ids(employee_id, include_root=False, include_vacant=True)):
        raise SimulationError("المدير الجديد يتبع هذا الموظف (مباشرة أو غير مباشرة)؛ "
                              "النقل سينشئ علاقة دائرية.")
    df = df.copy()
    old = df.at[e_ix, "manager_id"]
    _set_manager(df, [e_ix], new_manager_id)
    msg = f"نقل {employee_id} من المدير {old or '—'} إلى المدير {new_manager_id}"
    if move_department:
        for col in ("department", "division", "business_unit"):
            df.at[e_ix, col] = df.at[m_ix, col]
        msg += f" (والإدارة إلى {df.at[m_ix, 'department']})"
    return df, [msg + "."]


def remove_manager(df: pd.DataFrame, model: OrgModel, manager_id: str,
                   reassign_to: str = "manager_of_removed") -> tuple[pd.DataFrame, list[str]]:
    t_ix = _row_ix(df, manager_id)
    if manager_id == model.ceo_id:
        raise SimulationError("لا يمكن حذف الـCEO في هذا السيناريو.")
    df = df.copy()
    rows = list(_children_rows(df, manager_id))
    parent = df.at[t_ix, "manager_id"]
    directs = set(df.loc[rows, "node_id"])
    log = []
    if reassign_to == "manager_of_removed":
        target = parent
        if not target:
            raise SimulationError("المدير المحذوف ليس له مدير أعلى؛ اختر مستلمًا آخر.")
    elif reassign_to == "department_head":
        dept = df.at[t_ix, "department"]
        target = department_heads(model).get(dept)
        if target in (None, manager_id):
            target = parent
            log.append("المدير المحذوف هو رئيس الإدارة؛ تم تحويل الفريق إلى مديره الأعلى.")
    else:
        target = reassign_to
        _row_ix(df, target)
        subtree = set(model.subtree_ids(manager_id, include_root=False, include_vacant=True))
        if target in subtree and target not in directs:
            raise SimulationError("المستلم يقع في مستوى أدنى داخل فريق المدير المحذوف؛ اختر مديرًا "
                                  "من خارج الفريق أو أحد الموظفين المباشرين (ترقية).")
    if target in directs:  # promote a direct report
        p_ix = _row_ix(df, target)
        _set_manager(df, [p_ix], parent)
        rows = [r for r in rows if r != p_ix]
        log.append(f"ترقية {target} ليحل محل {manager_id} ويتبع {parent or '—'}.")
    _set_manager(df, rows, target)
    df = df.drop(index=t_ix)
    log.append(f"حذف المدير {manager_id} ونقل {len(rows)} سجل إلى {target}.")
    return df, log


ACTIONS = {
    "merge_departments": lambda df, m, a: merge_departments(
        df, m, a["departments"], a.get("new_name", ""), a.get("consolidate_heads", False),
        a.get("primary_head")),
    "remove_layer": lambda df, m, a: remove_layer(df, m, int(a["job_level"]), a.get("department")),
    "move_employee": lambda df, m, a: move_employee(
        df, m, a["employee_id"], a["new_manager_id"], a.get("move_department", True)),
    "remove_manager": lambda df, m, a: remove_manager(
        df, m, a["manager_id"], a.get("reassign_to", "manager_of_removed")),
}


def apply_scenario(df_norm: pd.DataFrame, actions: list[dict],
                   base_model: OrgModel | None = None) -> ScenarioResult:
    """Apply actions sequentially to a copy of ``df_norm``."""
    df = df_norm.copy()
    model = base_model or build_org_model(df)
    log: list[str] = []
    for i, action in enumerate(actions, 1):
        fn = ACTIONS.get(action.get("type"))
        if fn is None:
            raise SimulationError(f"نوع إجراء غير معروف: {action.get('type')}")
        df, msgs = fn(df, model, action)
        log.extend(f"[{i}] {m}" for m in msgs)
        model = build_org_model(df)
    return ScenarioResult(df=df, model=model, log=log)


# ============================================================================ comparison
def org_summary(model: OrgModel, settings: Settings) -> dict:
    df = model.df
    act = df[df["is_active"]]
    mgr = act[act["is_manager"]]["direct_reports"]
    depth = act.loc[act["node_id"].ne(model.ceo_id), "depth"].dropna()
    return {
        "Employees": len(act),
        "Managers": len(mgr),
        "Management Ratio": safe_div(len(mgr), len(act)),
        "Layers": layer_count(df),
        "Max Depth": int(depth.max()) if len(depth) else np.nan,
        "Avg Depth": float(depth.mean()) if len(depth) else np.nan,
        "Avg Span": float(mgr.mean()) if len(mgr) else np.nan,
        "Median Span": float(mgr.median()) if len(mgr) else np.nan,
        "Low Span Managers": int((mgr <= settings.low_span_threshold).sum()),
        "High Span Managers": int((mgr > settings.high_span_threshold).sum()),
        "Invalid Reporting": int(act["reporting_status"].isin(INVALID_REPORTING_STATUSES).sum()),
        "Vacancies": int((df["is_vacant"] & ~df["is_duplicate_id"]).sum()),
        "Departments": int(act["department"].replace("", np.nan).nunique()),
    }


def _fmt(metric: str, v) -> str:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return "—"
    if metric == "Management Ratio":
        return f"{v * 100:.1f}%"
    if isinstance(v, float) and not float(v).is_integer():
        return f"{v:.2f}"
    return f"{int(v):,}"


def compare(before: OrgModel, after: OrgModel, settings: Settings) -> dict:
    """Before/After comparison with an itemised list of what changed."""
    sb, sa = org_summary(before, settings), org_summary(after, settings)
    summary = pd.DataFrame({"Metric": list(sb), "Before": [sb[k] for k in sb],
                            "After": [sa[k] for k in sb]})
    summary["Change"] = [(a - b) if isinstance(a, (int, float)) and isinstance(b, (int, float))
                         else np.nan for a, b in zip(summary["After"], summary["Before"])]
    summary["Before (fmt)"] = [_fmt(m, v) for m, v in zip(summary["Metric"], summary["Before"])]
    summary["After (fmt)"] = [_fmt(m, v) for m, v in zip(summary["Metric"], summary["After"])]

    # Managers whose span changed (including removed / new managers)
    b = before.df[before.df["is_active"]].set_index("node_id")
    a = after.df[after.df["is_active"]].set_index("node_id")
    b = b[~b.index.duplicated()]
    a = a[~a.index.duplicated()]
    ids = sorted(set(b.index[b["is_manager"]]) | set(a.index[a["is_manager"]]))
    spans = pd.DataFrame(index=pd.Index(ids, name="node_id"))
    spans["employee_name"] = b["employee_name"].reindex(ids).fillna(a["employee_name"].reindex(ids))
    spans["job_title"] = b["job_title"].reindex(ids).fillna(a["job_title"].reindex(ids))
    spans["department_before"] = b["department"].reindex(ids)
    spans["span_before"] = b["direct_reports"].reindex(ids)
    spans["span_after"] = a["direct_reports"].reindex(ids)
    spans["status"] = np.where(spans["span_after"].isna(), "Removed",
                               np.where(spans["span_before"].isna(), "New", "Changed"))
    spans = spans[spans["span_before"].fillna(-1) != spans["span_after"].fillna(-1)].reset_index()

    # Employees whose depth changed
    common = b.index.intersection(a.index)
    db, da = b.loc[common, "depth"], a.loc[common, "depth"]
    changed = (db.fillna(-1) != da.fillna(-1))
    depth_changes = pd.DataFrame({
        "node_id": common[changed], "employee_name": b.loc[common[changed], "employee_name"].values,
        "department": a.loc[common[changed], "department"].values,
        "depth_before": db[changed].values, "depth_after": da[changed].values})

    # Department metrics
    ub = unit_metrics(before.df, "department", settings).set_index("department")
    ua = unit_metrics(after.df, "department", settings).set_index("department")
    cols = ["headcount", "managers", "management_ratio", "avg_span", "layers", "vacancies"]
    depts = ub.index.union(ua.index)
    dcomp = pd.concat({"before": ub[cols].reindex(depts), "after": ua[cols].reindex(depts)}, axis=1)
    diff_mask = pd.Series(False, index=depts)
    for c in cols:
        diff_mask |= dcomp[("before", c)].fillna(-1).round(4) != dcomp[("after", c)].fillna(-1).round(4)
    dept_changes = dcomp[diff_mask]
    dept_changes.columns = [f"{c}_{w}" for w, c in dept_changes.columns]
    dept_changes = dept_changes[[f"{c}_{w}" for c in cols for w in ("before", "after")]].reset_index()

    # Narrative
    lines = []
    for k in sb:
        if _fmt(k, sb[k]) != _fmt(k, sa[k]):
            lines.append(f"{k}: {_fmt(k, sb[k])} → {_fmt(k, sa[k])}")
    for r in spans.head(25).itertuples():
        bef = "—" if pd.isna(r.span_before) else int(r.span_before)
        aft = "removed" if pd.isna(r.span_after) else int(r.span_after)
        lines.append(f"Manager {r.node_id} ({r.job_title}): Direct Reports {bef} → {aft}")
    for r in dept_changes.head(25).itertuples():
        parts = []
        for c, label in [("headcount", "Headcount"), ("managers", "Managers"),
                         ("layers", "Layers"), ("management_ratio", "Management Ratio")]:
            bv, av = getattr(r, f"{c}_before"), getattr(r, f"{c}_after")
            if _fmt("Management Ratio" if c == "management_ratio" else c, bv) != \
                    _fmt("Management Ratio" if c == "management_ratio" else c, av):
                fmt = (lambda v: _fmt("Management Ratio", v)) if c == "management_ratio" \
                    else (lambda v: _fmt(c, v))
                parts.append(f"{label} {fmt(bv)} → {fmt(av)}")
        if parts:
            lines.append(f"Department {r.department}: " + "; ".join(parts))
    if len(depth_changes):
        lines.append(f"Organizational Depth تغير لـ {len(depth_changes)} موظف/موظفين.")
    return {"summary": summary, "span_changes": spans, "depth_changes": depth_changes,
            "department_changes": dept_changes, "narrative": lines,
            "before": sb, "after": sa}


# ================================================================================================
# X-Ray answers  (analytics/insights.py)
# ================================================================================================
def key_questions(model: OrgModel, df: pd.DataFrame, findings: pd.DataFrame,
                  settings: Settings) -> list[dict]:
    """``df`` is the (possibly filtered) enriched dataset."""
    act = df[df["is_active"]]
    mgr = act[act["is_manager"]]
    units = unit_metrics(df, "department", settings)
    units = units[units["headcount"] > 0]
    out: list[dict] = []

    def q(question, answer, evidence="", page=""):
        out.append({"question": question, "answer": answer, "evidence": evidence, "page": page})

    d = act["depth"].dropna()
    q("كم عدد الطبقات الإدارية؟",
      f"{layer_count(df)} طبقات هيكلية" if len(d) else "لا يمكن الحساب (لا يوجد CEO)",
      f"العمق من {int(d.min())} إلى {int(d.max())} علاقة إدارية حتى الـCEO." if len(d) else "",
      "Layers")
    q("كم عدد المديرين؟", f"{len(mgr):,} مدير",
      f"Management Ratio = {pct(len(mgr) / len(act) if len(act) else float('nan'))} "
      f"من {len(act):,} موظف نشط.", "Span of Control")
    if len(mgr):
        q("كم عدد الموظفين تحت كل مدير؟",
          f"الوسيط {mgr['direct_reports'].median():g} والمتوسط {mgr['direct_reports'].mean():.1f}",
          f"النطاق من {int(mgr['direct_reports'].min())} إلى {int(mgr['direct_reports'].max())} موظف مباشر.",
          "Span of Control")
    low = units.sort_values("low_span_managers", ascending=False)
    low = low[low["low_span_managers"] > 0].head(3)
    q("أين توجد الإدارات ذات Span of Control منخفض؟",
      ", ".join(f"{r.department} ({r.low_span_managers})" for r in low.itertuples()) or "لا توجد",
      f"عدد المديرين الذين لديهم ≤ {settings.low_span_threshold} موظفين مباشرين في كل إدارة.",
      "Departments")
    high = units.sort_values("high_span_managers", ascending=False)
    high = high[high["high_span_managers"] > 0].head(3)
    q("أين توجد الإدارات ذات Span of Control مرتفع؟",
      ", ".join(f"{r.department} ({r.high_span_managers})" for r in high.itertuples()) or "لا توجد",
      f"عدد المديرين الذين لديهم > {settings.high_span_threshold} موظفًا مباشرًا.", "Departments")
    if len(d):
        deepest = act.loc[act["depth"].idxmax()]
        q("ما أطول سلسلة إدارية؟", f"{int(deepest['depth'])} علاقات إدارية",
          f"مثال: {deepest['node_id']} – {deepest['job_title']} ({deepest['department']}).", "Layers")
    rs = act["reporting_status"].value_counts()
    q("هل توجد موظفون بدون مدير؟",
      f"{int(rs.get(RS_MISSING, 0))} بدون Manager ID، {int(rs.get(RS_INVALID, 0))} بمدير غير موجود",
      f"بالإضافة إلى {int(rs.get(RS_DISCONNECTED, 0))} موظف لا تصل سلسلتهم إلى الـCEO بسبب انقطاع أعلى.",
      "X-Ray Findings")
    term_mgrs = act.loc[act["reporting_status"].eq(RS_TERMINATED), "manager_id"].nunique()
    q("هل توجد علاقات إدارية تشير إلى مدير منتهي الخدمة؟",
      f"{int(rs.get(RS_TERMINATED, 0))} موظف يتبعون {term_mgrs} مدير/مديرين منتهي الخدمة",
      "Manager ID يشير إلى سجل حالته Terminated.", "X-Ray Findings")
    q("هل توجد حلقات إدارية Circular Reporting؟",
      f"{len(model.cycles)} حلقة ({int(rs.get(RS_CIRCULAR, 0))} موظف)",
      " | ".join(" → ".join(c + [c[0]]) for c in model.cycles[:3]), "X-Ray Findings")
    small = units[units["headcount"] <= settings.small_department_threshold]
    q("هل توجد إدارات صغيرة جدًا؟",
      ", ".join(f"{r.department} ({r.headcount})" for r in small.itertuples()) or "لا توجد",
      f"Headcount ≤ {settings.small_department_threshold}.", "Departments")
    hm = units[(units["headcount"] > settings.small_department_threshold)
               & (units["management_ratio"] > settings.high_management_ratio)]
    q("هل نسبة المديرين مرتفعة في بعض الإدارات؟",
      ", ".join(f"{r.department} ({pct(r.management_ratio)})"
                for r in hm.sort_values("management_ratio", ascending=False).itertuples())
      or "لا توجد",
      f"Management Ratio > {pct(settings.high_management_ratio, 0)}.", "Departments")
    dup = findings[findings["type"].eq("DUPLICATE_POSITION")] if not findings.empty else findings
    q("هل توجد وظائف متكررة أو متشابهة؟",
      f"{len(dup)} ملاحظة" if len(dup) else "لا توجد",
      "Position IDs مكررة، صيغ كتابة مختلفة لنفس المسمى، ومسميات متشابهة نصيًا.", "X-Ray Findings")
    vac = df[df["is_vacant"] & ~df["is_duplicate_id"]]
    top_vac = vac["department"].replace("", "(Unassigned)").value_counts().head(3)
    q("أين توجد شواغر؟", f"{len(vac)} وظيفة شاغرة",
      ", ".join(f"{k} ({v})" for k, v in top_vac.items()), "Departments")
    return out


# ================================================================================================
# Dataset report  (analytics/report.py)
# ================================================================================================
def dataset_facts(df_raw: pd.DataFrame, settings: Settings | None = None) -> dict:
    settings = settings or Settings()
    model = build_org_model(normalize(df_raw))
    d = model.df
    act = d[d["is_active"]]
    findings = detect_findings(model, settings)
    return {
        "records": len(d),
        "active_employees": len(act),
        "terminated_records": int(d["employment_status"].eq("Terminated").sum()),
        "vacant_positions": int(d["is_vacant"].sum()),
        "departments": int(d.loc[d["department"].ne(""), "department"].nunique()),
        "divisions": int(d.loc[d["division"].ne(""), "division"].nunique()),
        "business_units": int(d.loc[d["business_unit"].ne(""), "business_unit"].nunique()),
        "managers": int(act["is_manager"].sum()),
        "layers": layer_count(d),
        "depth": depth_stats(model),
        "span": span_stats(act.loc[act["is_manager"], "direct_reports"]),
        "findings_by_type": findings["type"].value_counts().to_dict(),
        "findings_by_severity": findings["severity"].value_counts().to_dict(),
        "findings_total": len(findings),
        "cycles": model.cycles,
    }


def dataset_report(df_raw: pd.DataFrame, manifest: list[dict] | None = None) -> str:
    f = dataset_facts(df_raw)
    lines = [
        "## Dataset facts (computed)",
        f"- Records: {f['records']:,}",
        f"- Active employees (headcount): {f['active_employees']:,}",
        f"- Vacant positions: {f['vacant_positions']}",
        f"- Terminated records still referenced: {f['terminated_records']}",
        f"- Business units / divisions / departments: {f['business_units']} / {f['divisions']} / {f['departments']}",
        f"- Managers (>= 1 active direct report): {f['managers']}",
        f"- Structural layers: {f['layers']} (depth {f['depth']['min']}–{f['depth']['max']}, "
        f"mean {f['depth']['mean']:.2f}, median {f['depth']['median']:g})",
        f"- Span of control: min {f['span']['min']:g}, P25 {f['span']['p25']:g}, median "
        f"{f['span']['median']:g}, mean {f['span']['mean']:.2f}, P75 {f['span']['p75']:g}, "
        f"max {f['span']['max']:g}",
        f"- Findings (default settings): {f['findings_total']}",
        "",
        "| Finding type | Count |", "|---|---|",
    ]
    lines += [f"| {k} | {v} |" for k, v in sorted(f["findings_by_type"].items())]
    lines += ["", "| Severity | Count |", "|---|---|"]
    lines += [f"| {k} | {v} |" for k, v in f["findings_by_severity"].items()]
    if manifest:
        lines += ["", "## Injected issues (generator manifest)", "",
                  "| Issue | Expected finding | Department | IDs |", "|---|---|---|---|"]
        for m in manifest:
            ids = ", ".join(m["ids"][:4]) + (" …" if len(m["ids"]) > 4 else "")
            lines.append(f"| {m['issue']} | {m['expected_finding']} | {m['department']} | {ids} |")
    return "\n".join(lines)


# ================================================================================================
# Analytics facade for a future AI assistant  (analytics/engine.py)
# ================================================================================================
def _records(df: pd.DataFrame, limit: int | None = 50) -> list[dict]:
    d = df.head(limit) if limit else df
    return json.loads(d.to_json(orient="records", date_format="iso", default_handler=str))


class OrgXRayEngine:
    def __init__(self, data: pd.DataFrame, settings: Settings | None = None,
                 normalized: bool = False, model: OrgModel | None = None):
        self.df = data if normalized else normalize(data)
        self.settings = settings or Settings()
        self._model = model

    # ------------------------------------------------------------------ core objects
    @property
    def model(self) -> OrgModel:
        if self._model is None:
            self._model = build_org_model(self.df)
        return self._model

    @cached_property
    def findings(self) -> pd.DataFrame:
        return detect_findings(self.model, self.settings)

    # ------------------------------------------------------------------ tools
    def summary(self) -> dict:
        s = org_summary(self.model, self.settings)
        s["Findings"] = len(self.findings)
        return {k: (None if isinstance(v, float) and np.isnan(v) else v) for k, v in s.items()}

    def departments_by(self, metric: str = "management_ratio", top: int = 10,
                       ascending: bool = False) -> list[dict]:
        u = unit_metrics(self.model.df, "department", self.settings)
        if metric not in u.columns:
            raise ValueError(f"Unknown metric '{metric}'. Options: {list(u.columns)}")
        return _records(u.sort_values(metric, ascending=ascending), top)

    def managers_by_span(self, min_reports: int | None = None, max_reports: int | None = None,
                         department: str | None = None) -> list[dict]:
        m = manager_table(self.model.df, self.settings)
        if min_reports is not None:
            m = m[m["direct_reports"] >= min_reports]
        if max_reports is not None:
            m = m[m["direct_reports"] <= max_reports]
        if department:
            m = m[m["department"].eq(department)]
        return _records(m, None)

    def span_statistics(self) -> dict:
        return span_stats(self.model.df.loc[self.model.df["is_manager"], "direct_reports"])

    def depth_statistics(self) -> dict:
        return depth_stats(self.model)

    def layers(self, basis: str = "job_level") -> list[dict]:
        return _records(layer_metrics(self.model.df, basis), None)

    def longest_reporting_chains(self, top: int = 10) -> list[dict]:
        return _records(longest_chains(self.model, top), None)

    def list_findings(self, finding_type: str | None = None, severity: str | None = None,
                      department: str | None = None, limit: int = 50) -> list[dict]:
        f = self.findings
        if finding_type:
            f = f[f["type"].eq(finding_type)]
        if severity:
            f = f[f["severity"].eq(severity)]
        if department:
            f = f[f["department"].eq(department)]
        return _records(f.drop(columns=["related_ids"]), limit)

    def explain(self, finding_id: str) -> dict:
        e = explain_finding(self.findings, finding_id, self.model)
        recs = e.pop("records")
        e["records"] = _records(recs[["node_id", "employee_name", "job_title", "department",
                                      "manager_id", "direct_reports", "depth",
                                      "reporting_status"]], 30)
        return e

    def explain_department(self, department: str) -> dict:
        u = unit_metrics(self.model.df, "department", self.settings)
        row = u[u["department"].eq(department)]
        if row.empty:
            raise ValueError(f"Department '{department}' not found")
        f = self.findings
        return {"metrics": _records(row)[0],
                "findings": _records(f[f["department"].eq(department)]
                                     .drop(columns=["related_ids"]), None),
                "thresholds": self.settings.as_dict()}

    def employee(self, employee_id: str) -> dict:
        r = self.model.node(employee_id)
        if r is None:
            raise ValueError(f"Employee '{employee_id}' not found")
        return {"record": _records(r.to_frame().T)[0],
                "chain_up": self.model.chain_up(employee_id),
                "direct_reports": self.model.direct_report_ids(employee_id),
                "total_reports": int(r["total_reports"])}

    def simulate(self, actions: list[dict]) -> dict:
        res = apply_scenario(self.df, actions, self.model)
        cmp = compare(self.model, res.model, self.settings)
        return {"log": res.log, "narrative": cmp["narrative"],
                "summary": _records(cmp["summary"][["Metric", "Before (fmt)", "After (fmt)"]], None),
                "department_changes": _records(cmp["department_changes"], 50),
                "span_changes": _records(cmp["span_changes"], 50)}

    def data_quality(self) -> list[dict]:
        return _records(run_checks(self.model).drop(columns=["failing_rows"]), None)

    def job_titles(self) -> dict:
        jt = job_title_analysis(self.model.df, self.settings)
        return {"variants": _records(jt["variants"], None), "similar": _records(jt["similar"], None)}


# ============================================================================ tool registry
TOOL_SPECS: list[dict] = [
    {"name": "summary", "description": "Organisation-level KPIs (headcount, managers, layers, spans, vacancies, findings).",
     "parameters": {"type": "object", "properties": {}}},
    {"name": "departments_by", "description": "Departments ranked by a metric, e.g. management_ratio, headcount, avg_span, layers, vacancy_rate.",
     "parameters": {"type": "object", "properties": {
         "metric": {"type": "string"}, "top": {"type": "integer"}, "ascending": {"type": "boolean"}}}},
    {"name": "managers_by_span", "description": "Managers filtered by number of direct reports.",
     "parameters": {"type": "object", "properties": {
         "min_reports": {"type": "integer"}, "max_reports": {"type": "integer"},
         "department": {"type": "string"}}}},
    {"name": "span_statistics", "description": "Min/max/mean/median/P25/P75 span of control.",
     "parameters": {"type": "object", "properties": {}}},
    {"name": "depth_statistics", "description": "Min/max/mean/median organisational depth.",
     "parameters": {"type": "object", "properties": {}}},
    {"name": "layers", "description": "Per-layer metrics; basis = job_level or depth.",
     "parameters": {"type": "object", "properties": {"basis": {"type": "string", "enum": ["job_level", "depth"]}}}},
    {"name": "longest_reporting_chains", "description": "The deepest reporting chains with full path.",
     "parameters": {"type": "object", "properties": {"top": {"type": "integer"}}}},
    {"name": "list_findings", "description": "X-Ray findings filtered by type / severity / department.",
     "parameters": {"type": "object", "properties": {
         "finding_type": {"type": "string"}, "severity": {"type": "string"},
         "department": {"type": "string"}, "limit": {"type": "integer"}}}},
    {"name": "explain", "description": "Rule, evidence, thresholds and source records behind a finding.",
     "parameters": {"type": "object", "properties": {"finding_id": {"type": "string"}}, "required": ["finding_id"]}},
    {"name": "explain_department", "description": "Metrics and findings of one department (e.g. why it is a Deep Hierarchy).",
     "parameters": {"type": "object", "properties": {"department": {"type": "string"}}, "required": ["department"]}},
    {"name": "employee", "description": "Profile, reporting chain and reports of an employee.",
     "parameters": {"type": "object", "properties": {"employee_id": {"type": "string"}}, "required": ["employee_id"]}},
    {"name": "simulate", "description": "Run a What-If scenario (list of actions) and return before/after.",
     "parameters": {"type": "object", "properties": {"actions": {"type": "array", "items": {"type": "object"}}},
                    "required": ["actions"]}},
    {"name": "data_quality", "description": "Data quality checks with scores.",
     "parameters": {"type": "object", "properties": {}}},
    {"name": "job_titles", "description": "Duplicate and similar job titles.",
     "parameters": {"type": "object", "properties": {}}},
]


def call_tool(engine: OrgXRayEngine, name: str, arguments: dict | None = None):
    """Dispatch a tool call (as an LLM would issue it) to the engine."""
    allowed = {t["name"] for t in TOOL_SPECS}
    if name not in allowed:
        raise ValueError(f"Unknown tool '{name}'")
    return getattr(engine, name)(**(arguments or {}))


# ================================================================================================
# Chart theme  (visualization/theme.py)
# ================================================================================================
CATEGORICAL = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7",
               "#e34948"]
SEQUENTIAL = ["#cde2fb", "#9ec5f4", "#6da7ec", "#3987e5", "#2a78d6", "#1c5cab", "#104281",
              "#0d366b"]
PRIMARY = "#2a78d6"
SECONDARY = "#eb6834"
MUTED = "#898781"
GRID = "rgba(137,135,129,0.25)"

SEVERITY_COLORS = {
    "Critical": "#d03b3b",
    "Significant": "#ec835a",
    "Attention": "#fab219",
    "Informational": "#898781",
}
SPAN_COLORS = {"Low": "#eb6834", "Within Range": "#2a78d6", "High": "#4a3aa7"}
FONT = "system-ui, -apple-system, 'Segoe UI', sans-serif"


def style(fig: go.Figure, height: int = 360, title: str | None = None,
          legend: bool = True) -> go.Figure:
    fig.update_layout(
        height=height,
        margin=dict(l=8, r=8, t=40 if title else 12, b=8),
        title=dict(text=title, x=0, font=dict(size=14)) if title else None,
        font=dict(family=FONT, size=12),
        showlegend=legend,
        legend=dict(orientation="h", yanchor="bottom", y=1.02, x=0, title=None),
        hoverlabel=dict(font=dict(family=FONT)),
        bargap=0.25,
    )
    fig.update_xaxes(showgrid=False, zeroline=False, linecolor=GRID)
    fig.update_yaxes(gridcolor=GRID, zeroline=False)
    return fig


def threshold_line(fig: go.Figure, value: float, label: str, axis: str = "x",
                   color: str = MUTED) -> go.Figure:
    kw = dict(line_dash="dash", line_color=color, line_width=1.5, annotation_text=label,
              annotation_font_size=11, annotation_font_color=color)
    if axis == "x":
        fig.add_vline(x=value, annotation_position="top", **kw)
    else:
        fig.add_hline(y=value, annotation_position="top left", **kw)
    return fig


# ================================================================================================
# Hierarchy charts  (visualization/hierarchy_charts.py)
# ================================================================================================
def layer_bars(layers: pd.DataFrame, title: str | None = None) -> go.Figure:
    """Employees and managers per layer (two series, grouped)."""
    d = layers.copy()
    d["label"] = [f"{int(l)} · {n}" if pd.notna(l) else n
                  for l, n in zip(d["layer"], d["layer_name"])]
    fig = go.Figure()
    fig.add_bar(y=d["label"], x=d["employees"], name="Employees", orientation="h",
                marker_color=PRIMARY,
                customdata=list(zip(d["pct_of_total"] * 100, d["managers"])),
                hovertemplate="%{y}<br>Employees: %{x:,}<br>Share: %{customdata[0]:.1f}%"
                              "<br>Managers: %{customdata[1]:,}<extra></extra>")
    fig.add_bar(y=d["label"], x=d["managers"], name="Managers", orientation="h",
                marker_color=SECONDARY,
                hovertemplate="%{y}<br>Managers: %{x:,}<extra></extra>")
    fig.update_layout(barmode="group", yaxis=dict(autorange="reversed"))
    fig.update_xaxes(title="Count")
    return style(fig, height=max(260, 42 * len(d)), title=title)


def depth_histogram(df: pd.DataFrame, threshold: int | None = None) -> go.Figure:
    d = df[df["is_active"] & df["depth"].notna()]["depth"].astype(int)
    counts = d.value_counts().sort_index()
    colors = [SECONDARY if threshold is not None and k > threshold else PRIMARY
              for k in counts.index]
    fig = go.Figure(go.Bar(x=counts.index, y=counts.values, marker_color=colors,
                           hovertemplate="Depth %{x}<br>Employees: %{y:,}<extra></extra>"))
    fig.update_xaxes(title="Depth (reporting steps to the CEO)", dtick=1)
    fig.update_yaxes(title="Employees")
    return style(fig, legend=False)


def level_depth_heatmap(matrix: pd.DataFrame) -> go.Figure:
    if matrix.empty:
        return style(go.Figure())
    fig = go.Figure(go.Heatmap(
        z=matrix.values, x=[f"Depth {c}" for c in matrix.columns],
        y=[f"Level {i}" for i in matrix.index],
        colorscale=[[i / (len(SEQUENTIAL) - 1), c] for i, c in enumerate(SEQUENTIAL)],
        text=matrix.values, texttemplate="%{text}", hoverongaps=False,
        hovertemplate="%{y} at %{x}: %{z:,} employees<extra></extra>"))
    fig.update_yaxes(autorange="reversed")
    return style(fig, height=320, legend=False)


def structure_sunburst(df: pd.DataFrame) -> go.Figure:
    a = df[df["is_active"]].copy()
    for c in ("business_unit", "division", "department"):
        a[c] = blank_to_label(a[c])
    g = a.groupby(["business_unit", "division", "department"]).size().reset_index(name="headcount")
    fig = px.sunburst(g, path=["business_unit", "division", "department"], values="headcount",
                      color_discrete_sequence=["#2a78d6", "#eb6834", "#1baf7a", "#4a3aa7",
                                               "#eda100"])
    fig.update_traces(hovertemplate="%{label}<br>Headcount: %{value:,}<extra></extra>",
                      insidetextorientation="radial")
    return style(fig, height=420, legend=False)


# ================================================================================================
# Department charts  (visualization/department_charts.py)
# ================================================================================================
def headcount_bar(units: pd.DataFrame, by: str = "department", top: int = 30) -> go.Figure:
    d = units.sort_values("headcount", ascending=True).tail(top)
    fig = go.Figure(go.Bar(
        y=d[by], x=d["headcount"], orientation="h", marker_color=PRIMARY,
        customdata=list(zip(d["managers"], d["vacancies"])),
        hovertemplate="%{y}<br>Headcount: %{x:,}<br>Managers: %{customdata[0]}"
                      "<br>Vacancies: %{customdata[1]}<extra></extra>"))
    fig.update_xaxes(title="Active employees")
    return style(fig, height=max(280, 22 * len(d)), legend=False)


def management_ratio_bar(units: pd.DataFrame, threshold: float, by: str = "department"
                         ) -> go.Figure:
    d = units.dropna(subset=["management_ratio"]).sort_values("management_ratio")
    colors = [SECONDARY if r > threshold else PRIMARY for r in d["management_ratio"]]
    fig = go.Figure(go.Bar(
        y=d[by], x=d["management_ratio"] * 100, orientation="h", marker_color=colors,
        customdata=list(zip(d["managers"], d["headcount"])),
        hovertemplate="%{y}<br>Management ratio: %{x:.1f}%<br>%{customdata[0]} managers / "
                      "%{customdata[1]} employees<extra></extra>"))
    fig.update_xaxes(title="Management ratio (%)", ticksuffix="%")
    threshold_line(fig, threshold * 100, f"Threshold {threshold * 100:.0f}%")
    return style(fig, height=max(280, 22 * len(d)), legend=False)


def span_vs_headcount(units: pd.DataFrame, by: str = "department") -> go.Figure:
    d = units.dropna(subset=["avg_span"])
    fig = go.Figure(go.Scatter(
        x=d["headcount"], y=d["avg_span"], mode="markers+text", text=d[by],
        textposition="top center", textfont=dict(size=10, color=MUTED),
        marker=dict(size=10, color=PRIMARY, line=dict(width=2, color="white")),
        customdata=list(zip(d["layers"], d["management_ratio"] * 100)),
        hovertemplate="%{text}<br>Headcount: %{x:,}<br>Avg span: %{y:.2f}"
                      "<br>Layers: %{customdata[0]}<br>Mgmt ratio: %{customdata[1]:.1f}%<extra></extra>"))
    fig.update_xaxes(title="Headcount", type="log")
    fig.update_yaxes(title="Average span of control")
    return style(fig, height=420, legend=False)


def vacancy_bar(units: pd.DataFrame, by: str = "department") -> go.Figure:
    d = units[units["vacancies"] > 0].sort_values("vacancies")
    fig = go.Figure(go.Bar(
        y=d[by], x=d["vacancies"], orientation="h", marker_color=PRIMARY,
        customdata=d["vacancy_rate"] * 100,
        hovertemplate="%{y}<br>Vacancies: %{x}<br>Vacancy rate: %{customdata:.1f}%<extra></extra>"))
    fig.update_xaxes(title="Vacant positions")
    return style(fig, height=max(260, 22 * len(d)), legend=False)


# ================================================================================================
# Overview charts  (visualization/dashboards.py)
# ================================================================================================
def span_histogram(spans: pd.Series, low: int, high: int, median: float | None = None
                   ) -> go.Figure:
    counts = spans.astype(int).value_counts().sort_index()
    cats = ["Low" if k <= low else ("High" if k > high else "Within Range") for k in counts.index]
    fig = go.Figure()
    for cat in ["Low", "Within Range", "High"]:
        idx = [i for i, c in zip(counts.index, cats) if c == cat]
        if not idx:
            continue
        fig.add_bar(x=idx, y=counts.loc[idx].values, name=cat, marker_color=SPAN_COLORS[cat],
                    hovertemplate="%{x} direct reports<br>Managers: %{y}<extra>" + cat + "</extra>")
    fig.update_xaxes(title="Direct reports per manager", dtick=1 if counts.index.max() <= 40 else 5)
    fig.update_yaxes(title="Managers")
    threshold_line(fig, low + 0.5, f"Low ≤ {low}")
    threshold_line(fig, high + 0.5, f"High > {high}")
    if median is not None:
        fig.add_vline(x=median, line_color=PRIMARY, line_width=1, line_dash="dot",
                      annotation_text=f"Median {median:g}", annotation_position="bottom right",
                      annotation_font_size=11)
    fig.update_layout(barmode="stack")
    return style(fig)


def findings_by_type(findings: pd.DataFrame) -> go.Figure:
    fig = go.Figure()
    if findings.empty:
        return style(fig)
    order = [t for t in FINDING_TYPES if t in set(findings["type"])]
    for sev in SEVERITIES:
        sub = findings[findings["severity"].eq(sev)]["type"].value_counts()
        fig.add_bar(y=order, x=[int(sub.get(t, 0)) for t in order], name=sev, orientation="h",
                    marker_color=SEVERITY_COLORS[sev],
                    hovertemplate="%{y}<br>" + sev + ": %{x}<extra></extra>")
    fig.update_layout(barmode="stack", yaxis=dict(autorange="reversed"))
    fig.update_xaxes(title="Findings")
    return style(fig, height=max(260, 30 * len(order)))


def department_distribution(df: pd.DataFrame, top: int = 25) -> go.Figure:
    a = df[df["is_active"]]
    counts = a["department"].replace("", "(Unassigned)").value_counts().head(top).sort_values()
    fig = go.Figure(go.Bar(y=counts.index, x=counts.values, orientation="h",
                           marker_color=PRIMARY,
                           hovertemplate="%{y}<br>Employees: %{x:,}<extra></extra>"))
    fig.update_xaxes(title="Active employees")
    return style(fig, height=max(300, 20 * len(counts)), legend=False)


# ================================================================================================
# Organisation network  (visualization/network.py)
# ================================================================================================
SCOPES = ["Entire Organization", "Business Unit", "Division", "Department", "Specific Manager"]


@dataclass
class NetworkSelection:
    nodes: list[str]
    total_in_scope: int
    truncated: bool
    levels_shown: int | None


def select_nodes(model: OrgModel, scope: str, value: str | None = None,
                 managers_only: bool = False, max_levels: int | None = None,
                 max_nodes: int = 600, include_vacant: bool = False) -> NetworkSelection:
    """Pick the nodes to draw for a scope, keeping the chart readable."""
    df = model.df
    base = df[df["is_active"] | (df["is_vacant"] & include_vacant)]
    base = base[~base["is_duplicate_id"]]
    if scope == "Entire Organization":
        ids = set(base["node_id"])
    elif scope == "Business Unit":
        ids = set(base.loc[base["business_unit"].eq(value), "node_id"])
    elif scope == "Division":
        ids = set(base.loc[base["division"].eq(value), "node_id"])
    elif scope == "Department":
        ids = set(base.loc[base["department"].eq(value), "node_id"])
    elif scope == "Specific Manager":
        ids = set(model.subtree_ids(value, include_root=True, include_vacant=include_vacant))
    else:
        raise ValueError(scope)
    if managers_only:
        mgr = set(df.loc[df["is_manager"], "node_id"])
        keep = ids & mgr
        if scope == "Specific Manager" and value:
            keep.add(value)
        ids = keep
    total = len(ids)

    # Relative depth inside the selection (roots = parent outside the selection)
    rel = _relative_depth(model.graph, ids, root_hint=value if scope == "Specific Manager" else None)
    levels = max_levels
    if levels is not None:
        ids = {n for n in ids if rel.get(n, 0) <= levels}
    truncated = False
    if len(ids) > max_nodes:
        truncated = True
        by_level = pd.Series({n: rel.get(n, 0) for n in ids})
        counts = by_level.value_counts().sort_index().cumsum()
        allowed = counts[counts <= max_nodes]
        lvl = int(allowed.index.max()) if len(allowed) else 0
        ids = set(by_level[by_level <= lvl].index)
        if len(ids) > max_nodes:  # a single very wide level
            ids = set(sorted(ids, key=lambda n: (by_level[n], n))[:max_nodes])
        levels = lvl
    return NetworkSelection(sorted(ids), total, truncated, levels)


def _relative_depth(g: nx.DiGraph, ids: set[str], root_hint: str | None = None) -> dict[str, int]:
    roots = [n for n in ids if not any(p in ids for p in g.predecessors(n))] if ids else []
    if root_hint and root_hint in ids and root_hint not in roots:
        roots.append(root_hint)  # e.g. a manager inside a reporting loop
    depth: dict[str, int] = {}
    frontier = sorted(roots)
    for r in frontier:
        depth[r] = 0
    while frontier:
        nxt = []
        for n in frontier:
            for c in g.successors(n) if n in g else []:
                if c in ids and c not in depth:
                    depth[c] = depth[n] + 1
                    nxt.append(c)
        frontier = nxt
    for n in ids:  # anything unreached (pure cycles) sits at level 0
        depth.setdefault(n, 0)
    return depth


def tree_layout(g: nx.DiGraph, ids: list[str]) -> dict[str, tuple[float, float]]:
    idset = set(ids)
    rel = _relative_depth(g, idset)
    children = {n: sorted(c for c in g.successors(n) if c in idset and rel.get(c) == rel[n] + 1)
                for n in ids if n in g}
    roots = sorted(n for n in ids if rel[n] == 0)
    x: dict[str, float] = {}
    counter = 0.0
    for root in roots:
        stack = [(root, False)]
        while stack:
            node, done = stack.pop()
            kids = children.get(node, [])
            if done or not kids:
                if kids:
                    x[node] = float(np.mean([x[k] for k in kids if k in x]))
                else:
                    x[node] = counter
                    counter += 1
                continue
            stack.append((node, True))
            for k in reversed(kids):
                if k not in x:
                    stack.append((k, False))
        counter += 0.5  # gap between separate trees
    return {n: (x.get(n, 0.0), -float(rel[n])) for n in ids}


def network_figure(model: OrgModel, sel: NetworkSelection, color_by: str = "span",
                   low: int = 3, high: int = 12, height: int = 650) -> go.Figure:
    ids = sel.nodes
    if not ids:
        return style(go.Figure(), height=200)
    pos = tree_layout(model.graph, ids)
    df = model.df.set_index("node_id")
    df = df[~df.index.duplicated()]
    idset = set(ids)

    ex, ey = [], []
    for u, v in model.graph.edges():
        if u in idset and v in idset and u != v:
            (x0, y0), (x1, y1) = pos[u], pos[v]
            ym = (y0 + y1) / 2
            ex += [x0, x0, x1, x1, None]
            ey += [y0, ym, ym, y1, None]
    fig = go.Figure(go.Scatter(x=ex, y=ey, mode="lines", line=dict(color=GRID, width=1),
                               hoverinfo="skip", showlegend=False))

    rows = df.loc[ids]
    xs = [pos[n][0] for n in ids]
    ys = [pos[n][1] for n in ids]
    tot = rows["total_reports"].clip(lower=0).values.astype(float)
    size = 8 + 16 * np.sqrt(tot / tot.max()) if tot.max() > 0 else np.full(len(ids), 8.0)
    hover = [
        f"<b>{n}</b> – {r.employee_name or '(Vacant)'}<br>{r.job_title}<br>{r.department}"
        f"<br>Direct reports: {int(r.direct_reports)} · Total: {int(r.total_reports)}"
        f"<br>Depth: {'' if pd.isna(r.depth) else int(r.depth)} · {r.reporting_status}"
        for n, r in zip(ids, rows.itertuples())
    ]
    if color_by == "span":
        def cat(r):
            if r.is_vacant:
                return "Vacant"
            if not r.is_manager:
                return "Individual contributor"
            if r.direct_reports <= low:
                return "Low"
            return "High" if r.direct_reports > high else "Within Range"
        cats = [cat(r) for r in rows.itertuples()]
        palette = dict(SPAN_COLORS, **{"Individual contributor": "#c3c2b7", "Vacant": "#ffffff"})
        for c in ["Low", "Within Range", "High", "Individual contributor", "Vacant"]:
            idx = [i for i, k in enumerate(cats) if k == c]
            if not idx:
                continue
            fig.add_scatter(x=[xs[i] for i in idx], y=[ys[i] for i in idx], mode="markers",
                            name=f"Span: {c}" if c in SPAN_COLORS else c,
                            marker=dict(size=[size[i] for i in idx], color=palette[c],
                                        line=dict(width=1.5, color=MUTED if c == "Vacant" else "white")),
                            text=[hover[i] for i in idx], hovertemplate="%{text}<extra></extra>")
    else:  # depth (sequential)
        d = rows["depth"].fillna(-1).values
        fig.add_scatter(x=xs, y=ys, mode="markers", name="Depth",
                        marker=dict(size=size, color=d, colorscale=[[i / (len(SEQUENTIAL) - 1), c]
                                                                     for i, c in enumerate(SEQUENTIAL)],
                                    showscale=True, colorbar=dict(title="Depth", thickness=10),
                                    line=dict(width=1, color="white")),
                        text=hover, hovertemplate="%{text}<extra></extra>")
    # Direct labels for the top of the tree only (keeps the chart legible)
    top_ids = [n for n in ids if pos[n][1] == 0 or (len(ids) <= 60)]
    if len(top_ids) <= 40:
        fig.add_scatter(x=[pos[n][0] for n in top_ids], y=[pos[n][1] + 0.18 for n in top_ids],
                        mode="text", text=[str(df.at[n, "job_title"])[:28] for n in top_ids],
                        textfont=dict(size=10, color=MUTED), hoverinfo="skip", showlegend=False)
    fig.update_xaxes(visible=False)
    fig.update_yaxes(visible=False)
    fig.update_layout(dragmode="pan")
    return style(fig, height=height)


# ================================================================================================
# Excel / CSV export  (exports/exporter.py)
# ================================================================================================
EMPLOYEE_EXPORT_COLUMNS = ["node_id"] + CANONICAL_COLUMNS + [
    "reporting_status", "depth", "direct_reports", "vacant_direct_positions", "total_reports",
    "is_manager",
]


def build_export_tables(model: OrgModel, findings: pd.DataFrame, settings: Settings,
                        df: pd.DataFrame | None = None) -> dict[str, pd.DataFrame]:
    """All export sheets, computed from ``df`` (filtered) or the full model."""
    data = df if df is not None else model.df
    mgr = manager_table(data, settings)
    stats = span_stats(mgr["direct_reports"])
    span_sheet = pd.concat([
        pd.DataFrame({"Statistic": [k.upper() if k.startswith("p") else k.title() for k in stats],
                      "Value": list(stats.values())}),
        pd.DataFrame({"Statistic": [""], "Value": [None]}),
        span_by(data, "department", settings).rename(columns={"department": "Statistic"}),
    ], ignore_index=True)
    dq = run_checks(model, data).drop(columns=["failing_rows"])
    layers = pd.concat([layer_metrics(data, "job_level").assign(basis="Job Level"),
                        layer_metrics(data, "depth").assign(basis="Structural Depth")],
                       ignore_index=True)
    f = findings.copy()
    if not f.empty:
        f["related_ids"] = f["related_ids"].map(lambda ids: ", ".join(ids[:50]))
    cols = [c for c in EMPLOYEE_EXPORT_COLUMNS if c in data.columns]
    return {
        "Employees": display_df(data[cols]),
        "Managers": display_df(mgr),
        "Departments": display_df(unit_metrics(data, "department", settings)),
        "Findings": display_df(f),
        "Span of Control": span_sheet,
        "Organizational Layers": display_df(layers),
        "Data Quality": display_df(dq),
    }


def _flatten(v):
    return ", ".join(map(str, v)) if isinstance(v, (list, tuple, set)) else v


def to_excel_bytes(sheets: dict[str, pd.DataFrame]) -> bytes:
    buf = io.BytesIO()
    header_fill = PatternFill("solid", fgColor="1C5CAB")
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            out = frame.copy()
            for c in out.columns:  # Excel cannot store tz-aware datetimes / lists
                if isinstance(out[c].dtype, pd.DatetimeTZDtype):
                    out[c] = out[c].dt.tz_localize(None)
                if out[c].dtype == object:
                    out[c] = out[c].map(_flatten)
            sheet = name[:31]
            out.to_excel(writer, sheet_name=sheet, index=False)
            ws = writer.sheets[sheet]
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            for i, col in enumerate(out.columns, 1):
                sample = out[col].head(200).map(lambda v: len(str(v)))
                width = min(60, max(10, len(str(col)) + 2, int(sample.max() if len(sample) else 0) + 2))
                ws.column_dimensions[get_column_letter(i)].width = width
            if len(out):
                ws.auto_filter.ref = ws.dimensions
    return buf.getvalue()


def to_csv_bytes(df: pd.DataFrame) -> bytes:
    out = df.copy()
    for c in out.columns:
        if out[c].dtype == object:
            out[c] = out[c].map(_flatten)
    return out.to_csv(index=False).encode("utf-8-sig")


# ================================================================================================
# UI: session state & caching  (ui/state.py)
# ================================================================================================
FILTERS = [
    ("business_unit", "Business Unit"),
    ("division", "Division"),
    ("department", "Department"),
    ("location", "Location"),
    ("job_level", "Job Level"),
    ("employment_status", "Employment Status"),
    ("employment_type", "Employment Type"),
]


# ============================================================================ cached work
@st.cache_data(show_spinner="Generating dummy organisation…", max_entries=4)
def cached_dummy(n: int, seed: int) -> tuple[pd.DataFrame, list[dict]]:
    return generate_org(n, seed)


@st.cache_resource(show_spinner="Building organisational hierarchy…", max_entries=4)
def cached_model(token: str, _raw: pd.DataFrame) -> OrgModel:
    return build_org_model(normalize(_raw))


@st.cache_data(show_spinner="Running X-Ray rules…", max_entries=16)
def cached_findings(token: str, settings_key: tuple, _model: OrgModel) -> pd.DataFrame:
    return detect_findings(_model, Settings(*settings_key))


@st.cache_data(show_spinner="Checking data quality…", max_entries=8)
def cached_quality(token: str, _model: OrgModel) -> pd.DataFrame:
    return run_checks(_model)


# ============================================================================ session
def init_state() -> None:
    ss = st.session_state
    if "settings" not in ss:
        ss.settings = Settings()
    if "dataset" not in ss:
        load_dummy(1500, 42)
    ss.setdefault("scenario", [])


def load_dummy(n: int, seed: int) -> None:
    raw, manifest = cached_dummy(n, seed)
    st.session_state.dataset = {
        "raw": raw, "token": f"dummy-{n}-{seed}", "manifest": manifest,
        "source": f"Dummy data · ~{n:,} employees · seed {seed}",
    }
    reset_filters()
    st.session_state.scenario = []


def load_uploaded(raw: pd.DataFrame, filename: str) -> None:
    st.session_state.dataset = {
        "raw": raw, "token": f"upload-{df_fingerprint(raw)}", "manifest": None,
        "source": f"Uploaded file · {filename} · {len(raw):,} rows",
    }
    reset_filters()
    st.session_state.scenario = []


def reset_filters() -> None:
    for key, _ in FILTERS:
        st.session_state[f"flt_{key}"] = []
    st.session_state["flt_manager"] = None


# ============================================================================ context
@dataclass
class Context:
    raw: pd.DataFrame
    token: str
    source: str
    manifest: list | None
    model: OrgModel
    settings: Settings
    findings_all: pd.DataFrame
    df: pd.DataFrame  # filtered enriched records
    findings: pd.DataFrame  # findings touching the filtered records
    filtered: bool
    filter_summary: str

    @property
    def active(self) -> pd.DataFrame:
        return self.df[self.df["is_active"]]


def _filter_mask(model: OrgModel) -> tuple[pd.Series, list[str]]:
    df = model.df
    mask = pd.Series(True, index=df.index)
    parts = []
    for key, label in FILTERS:
        sel = st.session_state.get(f"flt_{key}") or []
        if sel:
            col = df[key]
            if key == "job_level":
                mask &= col.isin([float(s) for s in sel])
            else:
                mask &= col.isin(sel)
            parts.append(f"{label}: {', '.join(map(str, sel[:3]))}{'…' if len(sel) > 3 else ''}")
    mgr = st.session_state.get("flt_manager")
    if mgr:
        ids = set(model.subtree_ids(mgr, include_root=True, include_vacant=True))
        mask &= df["node_id"].isin(ids)
        parts.append(f"Manager: {mgr}")
    return mask, parts


def context() -> Context:
    ss = st.session_state
    ds = ss.dataset
    settings: Settings = ss.settings
    model = cached_model(ds["token"], ds["raw"])
    findings_all = cached_findings(ds["token"], settings.cache_key(), model)
    mask, parts = _filter_mask(model)
    filtered = bool(parts)
    df = model.df[mask] if filtered else model.df
    findings = findings_all
    if filtered:
        findings = filter_findings(findings_all, set(df["node_id"]),
                                   set(df["department"]))
    return Context(raw=ds["raw"], token=ds["token"], source=ds["source"],
                   manifest=ds.get("manifest"), model=model, settings=settings,
                   findings_all=findings_all, df=df, findings=findings, filtered=filtered,
                   filter_summary=" · ".join(parts))


def filter_options(model: OrgModel) -> dict[str, list]:
    df = model.df[~model.df["is_duplicate_id"]]
    opts = {}
    for key, _ in FILTERS:
        vals = df[key].dropna()
        if key == "job_level":
            opts[key] = sorted({int(v) for v in vals})
        else:
            opts[key] = sorted(v for v in set(vals) if v != "")
    return opts


# ================================================================================================
# UI: shared components  (ui/common.py)
# ================================================================================================
CSS = """
<style>
/* Arabic prose reads right-to-left; tables and charts stay left-to-right. */
[data-testid="stMarkdownContainer"] p, [data-testid="stMarkdownContainer"] li,
[data-testid="stCaptionContainer"], .stAlert p { direction: rtl; text-align: right; }
[data-testid="stMarkdownContainer"] .ltr, [data-testid="stMarkdownContainer"] .ltr p,
[data-testid="stMarkdownContainer"] code { direction: ltr; text-align: left; unicode-bidi: plaintext; }
section[data-testid="stSidebar"] [data-testid="stMarkdownContainer"] p { direction: ltr; text-align: left; }
[data-testid="stMetric"] { padding: 0.6rem 0.8rem; }
[data-testid="stMetricValue"] { font-size: 1.6rem; }
.xr-qa { border: 1px solid rgba(137,135,129,.25); border-radius: 8px; padding: .6rem .8rem;
         margin-bottom: .5rem; direction: rtl; text-align: right; }
.xr-qa .q { font-size: .85rem; opacity: .75; }
.xr-qa .a { font-size: 1.05rem; font-weight: 600; margin: .15rem 0; }
.xr-qa .e { font-size: .8rem; opacity: .7; }
.xr-chain { direction: ltr; text-align: left; }
.xr-chain .node { border: 1px solid rgba(137,135,129,.35); border-radius: 6px; padding: .35rem .6rem;
                  display: inline-block; min-width: 320px; }
.xr-chain .node.sel { border-color: #2a78d6; border-width: 2px; }
.xr-chain .arrow { opacity: .5; padding-left: 1.2rem; }
.xr-badge { border-radius: 4px; padding: 1px 6px; font-size: .78rem; font-weight: 600; }
</style>
"""

SEVERITY_ICONS = {"Critical": "🔴", "Significant": "🟠", "Attention": "🟡", "Informational": "⚪"}


def inject_css() -> None:
    st.markdown(CSS, unsafe_allow_html=True)


def page_header(title: str, subtitle: str, ctx=None) -> None:
    st.title(title)
    st.caption(subtitle)
    if ctx is not None and ctx.filtered:
        st.info(f"الفلاتر العامة مفعّلة — {ctx.filter_summary}", icon="🔎")


def sidebar_filters(model) -> None:
    opts = filter_options(model)
    with st.sidebar:
        st.markdown("**Global Filters**")
        for key, label in FILTERS:
            fmt = (lambda v: f"{v} · {LEVEL_NAMES.get(v, '')}") if key == "job_level" else str
            st.multiselect(label, opts[key], key=f"flt_{key}", format_func=fmt,
                           placeholder="All")
        mgr = model.df[model.df["is_manager"]].sort_values("total_reports", ascending=False)
        labels = {r.node_id: f"{r.node_id} · {r.employee_name} · {r.job_title}"
                  for r in mgr.itertuples()}
        st.selectbox("Manager (organisation under)", [None] + list(labels), key="flt_manager",
                     format_func=lambda v: "All" if v is None else labels.get(v, v),
                     placeholder="All")
        st.button("Clear filters", on_click=reset_filters, width="stretch")


def kpi(col, label: str, value: str, help: str | None = None, delta: str | None = None) -> None:
    col.metric(label, value, delta=delta, help=help, border=True)


def table(df: pd.DataFrame, key: str, filename: str | None = None, height: int | None = None,
          rename: bool = True, column_config: dict | None = None, **kwargs):
    """Dataframe with a CSV download button. Returns the st.dataframe event (if selectable)."""
    shown = display_df(df) if rename else df
    cfg = default_column_config(shown)
    if column_config:
        cfg.update(column_config)
    args = dict(width="stretch", hide_index=True, column_config=cfg)
    if height:
        args["height"] = height
    args.update(kwargs)
    event = st.dataframe(shown, key=key, **args)
    st.download_button("⬇ CSV", to_csv_bytes(shown), file_name=f"{filename or key}.csv",
                       mime="text/csv", key=f"dl_{key}")
    return event


def default_column_config(df: pd.DataFrame) -> dict:
    cfg = {}
    for c in df.columns:
        lc = str(c).lower()
        if "ratio" in lc or "rate" in lc or lc in ("pct of total", "score"):
            cfg[c] = st.column_config.NumberColumn(c, format="percent")
        elif "salary" in lc:
            cfg[c] = st.column_config.NumberColumn(c, format="%,.0f")
        elif lc in ("avg span", "avg depth", "mean", "avg_span", "avg_depth", "median span"):
            cfg[c] = st.column_config.NumberColumn(c, format="%.2f")
    return cfg


def severity_label(sev: str) -> str:
    return f"{SEVERITY_ICONS.get(sev, '')} {sev}"


def empty_state(msg: str = "لا توجد بيانات مطابقة للفلاتر الحالية.") -> None:
    st.info(msg, icon="ℹ️")


def col_label(c: str) -> str:
    return DISPLAY_NAMES.get(c, c.replace("_", " ").title())


# ================================================================================================
# Page: Dashboard  (ui/dashboard.py)
# ================================================================================================
def page_dashboard() -> None:
    ctx = context()
    page_header("Organizational X-Ray", f"الملخص التنظيمي · {ctx.source}", ctx)
    df, s = ctx.df, ctx.settings
    act = ctx.active
    if act.empty:
        empty_state()
        return
    mgr = act[act["is_manager"]]
    spans = mgr["direct_reports"]
    vac = int((df["is_vacant"] & ~df["is_duplicate_id"]).sum())
    positions = len(act) + vac
    invalid = int(act["reporting_status"].isin(INVALID_REPORTING_STATUSES).sum())
    disconnected = int(act["reporting_status"].eq(RS_DISCONNECTED).sum())

    c = st.columns(5)
    kpi(c[0], "إجمالي الموظفين", f"{len(act):,}", "الموظفون النشطون (Active + On Leave) في وظائف مشغولة.")
    kpi(c[1], "إجمالي المديرين", f"{len(mgr):,}", "موظفون نشطون لديهم موظف مباشر نشط واحد على الأقل.")
    kpi(c[2], "Management Ratio", pct(len(mgr) / len(act)), "عدد المديرين ÷ إجمالي الموظفين.")
    kpi(c[3], "الطبقات التنظيمية", f"{layer_count(df)}",
        "عدد مستويات العمق المختلفة (CEO = 0) بين الموظفين المتصلين بالـCEO.")
    kpi(c[4], "Vacancy Rate", pct(vac / positions if positions else float("nan")),
        f"{vac} وظيفة شاغرة ÷ {positions:,} وظيفة.")
    c = st.columns(4)
    kpi(c[0], "Average Span of Control", fmt_num(spans.mean(), 2) if len(spans) else "—",
        "متوسط عدد الموظفين المباشرين لكل مدير.")
    kpi(c[1], "Median Span of Control", fmt_num(spans.median()) if len(spans) else "—")
    kpi(c[2], "علاقات إدارية غير صحيحة", f"{invalid:,}",
        "بدون Manager ID، Manager ID غير موجود، مدير منتهي الخدمة، أو علاقة دائرية. "
        f"يضاف إليهم {disconnected} موظف منقطعون بسبب خلل في مستوى أعلى.")
    kpi(c[3], "Organizational Findings", f"{len(ctx.findings):,}",
        " · ".join(f"{sev}: {int((ctx.findings['severity'] == sev).sum())}" for sev in SEVERITIES))

    st.subheader("إجابات الأشعة السينية")
    st.caption("كل إجابة محسوبة مباشرة من البيانات الحالية (مع الفلاتر)، والأدلة التفصيلية في الصفحات المشار إليها.")
    qa = key_questions(ctx.model, df, ctx.findings, s)
    cols = st.columns(3)
    for i, item in enumerate(qa):
        cols[i % 3].markdown(
            f"<div class='xr-qa'><div class='q'>{html.escape(item['question'])}</div>"
            f"<div class='a'>{html.escape(item['answer'])}</div>"
            f"<div class='e'>{html.escape(item['evidence'])} · <i>{item['page']}</i></div></div>",
            unsafe_allow_html=True)

    st.subheader("التوزيع والهيكل")
    left, right = st.columns(2)
    with left:
        st.markdown("**توزيع الموظفين على الطبقات (Job Level)**")
        st.plotly_chart(layer_bars(layer_metrics(df, "job_level")), width="stretch")
    with right:
        st.markdown("**توزيع Span of Control**")
        if len(spans):
            st.plotly_chart(span_histogram(spans, s.low_span_threshold, s.high_span_threshold,
                                           float(spans.median())), width="stretch")
        else:
            empty_state("لا يوجد مديرون ضمن الفلاتر الحالية.")
    left, right = st.columns(2)
    with left:
        st.markdown("**توزيع الموظفين على الإدارات**")
        st.plotly_chart(department_distribution(df), width="stretch")
    with right:
        st.markdown("**الهيكل الإداري: Business Unit → Division → Department**")
        st.plotly_chart(structure_sunburst(df), width="stretch")
    st.markdown("**Organizational Findings حسب النوع ومستوى الأهمية**")
    st.plotly_chart(findings_by_type(ctx.findings), width="stretch")


# ================================================================================================
# Page: Layers  (ui/layers.py)
# ================================================================================================
def page_layers() -> None:
    ctx = context()
    page_header("Layers", "تحليل الطبقات التنظيمية وعمق التسلسل الإداري", ctx)
    df, s = ctx.df, ctx.settings
    if ctx.active.empty:
        empty_state()
        return

    tab_level, tab_depth = st.tabs(["حسب Job Level (المعلن)", "حسب العمق الفعلي (Structural Depth)"])
    with tab_level:
        st.caption("Layer 0 → CEO، Layer 1 → C-Suite، Layer 2 → Directors، Layer 3 → Senior Managers، "
                   "Layer 4 → Managers، Layer 5 → Supervisors، Layer 6 → Individual Contributors")
        lm = layer_metrics(df, "job_level")
        st.plotly_chart(layer_bars(lm), width="stretch")
        table(lm, "layers_level", "layers_by_job_level")
    with tab_depth:
        st.caption("العمق = عدد العلاقات الإدارية بين الموظف والـCEO، محسوبًا من Manager ID وليس من المسمى.")
        lm = layer_metrics(df, "depth")
        st.plotly_chart(layer_bars(lm), width="stretch")
        table(lm, "layers_depth", "layers_by_depth")

    st.subheader("Organizational Depth")
    ds = depth_stats(ctx.model, df)
    c = st.columns(5)
    kpi(c[0], "Minimum", fmt_num(ds["min"]))
    kpi(c[1], "Maximum", fmt_num(ds["max"]))
    kpi(c[2], "Mean", fmt_num(ds["mean"], 2))
    kpi(c[3], "Median", fmt_num(ds["median"]))
    kpi(c[4], "Deep Hierarchy Threshold", str(s.deep_hierarchy_threshold),
        "الموظفون على عمق أكبر من هذا الحد يظهرون باللون البرتقالي.")
    st.caption(f"الإحصاءات محسوبة على {ds['count']:,} موظف متصل بالـCEO (بدون الـCEO نفسه).")
    left, right = st.columns(2)
    with left:
        st.markdown("**توزيع العمق**")
        st.plotly_chart(depth_histogram(df, s.deep_hierarchy_threshold), width="stretch")
    with right:
        st.markdown("**Job Level مقابل العمق الفعلي**")
        st.caption("الخلايا خارج القطر تعني أن المستوى المعلن لا يطابق موقع الموظف الفعلي في التسلسل.")
        st.plotly_chart(level_depth_heatmap(level_depth_matrix(df)), width="stretch")

    st.subheader("أطول Reporting Chains")
    top = st.slider("عدد السلاسل", 5, 50, 10, key="chains_top")
    chains = longest_chains(ctx.model, top, df)
    if chains.empty:
        empty_state()
    else:
        table(chains, "longest_chains", column_config={
            "Chain": st.column_config.TextColumn("Chain", width="large")})

    st.subheader("العمق لكل موظف")
    q = st.text_input("بحث (Employee ID / Name / Department)", key="depth_search")
    emp = df[df["is_active"]][["node_id", "employee_name", "job_title", "job_level", "department",
                              "manager_id", "depth", "reporting_status"]]
    if q:
        m = (emp["node_id"].str.contains(q, case=False, regex=False)
             | emp["employee_name"].str.contains(q, case=False, regex=False)
             | emp["department"].str.contains(q, case=False, regex=False))
        emp = emp[m]
    table(emp.sort_values("depth", ascending=False, na_position="first"), "depth_per_employee",
          height=360)


# ================================================================================================
# Page: Span of Control  (ui/span.py)
# ================================================================================================
def page_span() -> None:
    ctx = context()
    page_header("Span of Control", "عدد الموظفين المباشرين النشطين لكل مدير", ctx)
    s = ctx.settings
    mt = manager_table(ctx.df, s)
    if mt.empty:
        empty_state("لا يوجد مديرون ضمن الفلاتر الحالية.")
        return
    st_ = span_stats(mt["direct_reports"])
    c = st.columns(7)
    kpi(c[0], "Managers", f"{st_['count']:,}")
    for col, (k, label) in zip(c[1:], [("min", "Minimum"), ("max", "Maximum"), ("mean", "Mean"),
                                        ("median", "Median"), ("p25", "P25"), ("p75", "P75")]):
        kpi(col, label, fmt_num(st_[k], 2))

    st.plotly_chart(span_histogram(mt["direct_reports"], s.low_span_threshold,
                                   s.high_span_threshold, st_["median"]), width="stretch")
    st.caption(f"Low Span: ≤ {s.low_span_threshold} · High Span: > {s.high_span_threshold} "
               "(قابلة للتعديل من Settings). التصنيف يستدعي المراجعة ولا يعني حكمًا على المدير أو الدور.")

    st.subheader("المديرون")
    c1, c2, c3 = st.columns([1, 1, 2])
    cats = c1.multiselect("Span Category", ["Low", "Within Range", "High"], key="span_cat",
                          placeholder="All")
    rng = c2.slider("Direct reports", 0, int(mt["direct_reports"].max()),
                    (0, int(mt["direct_reports"].max())), key="span_rng")
    q = c3.text_input("بحث (ID / Name / Title / Department)", key="span_q")
    view = mt[mt["direct_reports"].between(*rng)]
    if cats:
        view = view[view["span_category"].isin(cats)]
    if q:
        hay = view[["node_id", "employee_name", "job_title", "department"]].fillna("").astype(str).agg(" ".join, axis=1)
        view = view[hay.str.contains(q, case=False, regex=False)]
    view = view.assign(org_median=st_["median"],
                       vs_median=(view["direct_reports"] / st_["median"]).round(2))
    table(view, "managers_span", "managers_span_of_control", height=420)

    st.subheader("Span of Control حسب الإدارة")
    table(span_by(ctx.df, "department", s), "span_by_department")
    st.subheader("Span of Control حسب Job Level")
    table(span_by(ctx.df, "job_level", s), "span_by_level")


# ================================================================================================
# Page: Departments  (ui/departments.py)
# ================================================================================================
UNIT_VIEW_COLS = ["headcount", "managers", "management_ratio", "avg_span", "avg_depth", "vacancies",
                  "avg_salary", "median_salary", "layers", "max_depth", "vacancy_rate",
                  "low_span_managers", "high_span_managers", "invalid_reporting"]


def _units_tab(ctx, by: str, key: str) -> None:
    s = ctx.settings
    units = unit_metrics(ctx.df, by, s)
    if units.empty:
        empty_state()
        return
    c1, c2, c3, c4 = st.columns([2, 1, 1, 1])
    q = c1.text_input("بحث", key=f"{key}_q", placeholder="اسم الوحدة…")
    min_hc = c2.number_input("Headcount ≥", 0, value=0, key=f"{key}_min")
    flagged = c3.selectbox("تصفية", ["الكل", "Management Ratio مرتفع", "وحدات صغيرة",
                                     "بها شواغر", "بها علاقات غير صحيحة"], key=f"{key}_flag")
    sort_col = c4.selectbox("ترتيب حسب", UNIT_VIEW_COLS, key=f"{key}_sort")
    view = units[units["headcount"] >= min_hc]
    if q:
        view = view[view[by].str.contains(q, case=False, regex=False)]
    if flagged == "Management Ratio مرتفع":
        view = view[view["management_ratio"] > s.high_management_ratio]
    elif flagged == "وحدات صغيرة":
        view = view[view["headcount"] <= s.small_department_threshold]
    elif flagged == "بها شواغر":
        view = view[view["vacancies"] > 0]
    elif flagged == "بها علاقات غير صحيحة":
        view = view[view["invalid_reporting"] > 0]
    view = view.sort_values(sort_col, ascending=False, na_position="last")
    lead = [by] + (["division", "business_unit"] if by == "department" else [])
    table(view[lead + UNIT_VIEW_COLS], f"{key}_table", f"{by}_analysis", height=420)

    left, right = st.columns(2)
    with left:
        st.markdown("**Management Ratio**")
        st.plotly_chart(management_ratio_bar(view, s.high_management_ratio, by), width="stretch",
                        key=f"{key}_management_ratio_bar")
    with right:
        st.markdown("**Headcount**")
        st.plotly_chart(headcount_bar(view, by), width="stretch",
                        key=f"{key}_headcount_bar")
    left, right = st.columns(2)
    with left:
        st.markdown("**Headcount مقابل متوسط Span of Control**")
        st.plotly_chart(span_vs_headcount(view, by), width="stretch",
                        key=f"{key}_span_vs_headcount")
    with right:
        st.markdown("**الوظائف الشاغرة**")
        st.plotly_chart(vacancy_bar(view, by), width="stretch",
                        key=f"{key}_vacancy_bar")


def _drilldown(ctx) -> None:
    units = unit_metrics(ctx.df, "department", ctx.settings)
    if units.empty:
        return
    dept = st.selectbox("اختر إدارة", units["department"].tolist(), key="dept_drill")
    u = units[units["department"].eq(dept)].iloc[0]
    c = st.columns(6)
    kpi(c[0], "Headcount", f"{int(u['headcount']):,}")
    kpi(c[1], "Managers", f"{int(u['managers'])}")
    kpi(c[2], "Management Ratio", pct(u["management_ratio"]))
    kpi(c[3], "Avg Span", fmt_num(u["avg_span"], 2))
    kpi(c[4], "Layers", f"{int(u['layers'])}")
    kpi(c[5], "Vacancies", f"{int(u['vacancies'])}")
    f = ctx.findings_all
    f = f[f["department"].eq(dept)]
    st.markdown(f"**Findings في {dept}: {len(f)}**")
    if len(f):
        v = f[["finding_id", "type", "severity", "entity_name", "description"]].copy()
        v["severity"] = v["severity"].map(severity_label)
        table(v, "dept_findings", f"findings_{dept}")


def _titles_tab(ctx) -> None:
    jt = job_title_analysis(ctx.df, ctx.settings)
    st.markdown("**المسميات المتكررة بصيغ كتابة مختلفة (Duplicate Job Titles)**")
    if jt["variants"].empty:
        st.caption("لا توجد.")
    else:
        table(jt["variants"], "title_variants")
    st.markdown(f"**المسميات المتشابهة (Similar Job Titles ≥ {ctx.settings.title_similarity_threshold:.2f})**")
    st.caption("مقارنة داخل نفس القطاع. يتم تجاهل أزواج السلم الوظيفي (مثل Accountant / Senior Accountant).")
    if jt["similar"].empty:
        st.caption("لا توجد.")
    else:
        table(jt["similar"], "title_similar")
    st.markdown("**تكرار المسميات الوظيفية**")
    table(jt["titles"].drop(columns=["normalized"]), "title_counts", height=360)


def page_departments() -> None:
    ctx = context()
    page_header("Departments", "تحليل الإدارات والقطاعات ووحدات الأعمال", ctx)
    if ctx.active.empty and not ctx.df["is_vacant"].any():
        empty_state()
        return
    tabs = st.tabs(["Departments", "Department Drill-down", "Divisions", "Business Units",
                    "Job Titles"])
    with tabs[0]:
        _units_tab(ctx, "department", "dept")
    with tabs[1]:
        _drilldown(ctx)
    with tabs[2]:
        _units_tab(ctx, "division", "div")
    with tabs[3]:
        _units_tab(ctx, "business_unit", "bu")
    with tabs[4]:
        _titles_tab(ctx)


# ================================================================================================
# Page: Organization Map  (ui/network_page.py)
# ================================================================================================
def page_network() -> None:
    ctx = context()
    page_header("Organization Map", "الشبكة التنظيمية Employee → Manager (NetworkX + Plotly)")
    st.caption("الخريطة تُرسم دائمًا على كامل البيانات حتى تبقى السلاسل الإدارية مكتملة؛ "
               "استخدم النطاق أدناه لاختيار الجزء المعروض.")
    model, s = ctx.model, ctx.settings
    df = model.df

    c1, c2 = st.columns([1, 2])
    scope = c1.radio("النطاق", SCOPES, key="net_scope")
    value = None
    with c2:
        if scope in ("Business Unit", "Division", "Department"):
            col = {"Business Unit": "business_unit", "Division": "division",
                   "Department": "department"}[scope]
            vals = sorted(v for v in df.loc[df["is_active"], col].unique() if v)
            value = st.selectbox(scope, vals, key=f"net_val_{col}")
        elif scope == "Specific Manager":
            mgr = df[df["is_manager"]].sort_values("total_reports", ascending=False)
            labels = {r.node_id: f"{r.node_id} · {r.employee_name} · {r.job_title} "
                                 f"({r.total_reports} reports)" for r in mgr.itertuples()}
            value = st.selectbox("Manager", list(labels), format_func=labels.get, key="net_mgr")
        o1, o2, o3, o4 = st.columns(4)
        managers_only = o1.toggle("Managers only", value=(scope == "Entire Organization"),
                                  key=f"net_mo_{scope}",
                                  help="عرض المديرين فقط يجعل خريطة المنظمة كاملة مقروءة.")
        include_vacant = o2.toggle("Show vacancies", value=False, key="net_vac")
        levels = o3.number_input("Max levels", 1, 15, 15 if scope != "Entire Organization" else 8,
                                 key=f"net_lv_{scope}")
        color_by = o4.selectbox("Color by", ["span", "depth"], key="net_color",
                                format_func={"span": "Span category", "depth": "Depth"}.get)

    sel = select_nodes(model, scope, value, managers_only, int(levels), s.max_network_nodes,
                       include_vacant)
    if not sel.nodes:
        empty_state("لا توجد عقد للعرض في هذا النطاق.")
        return
    msg = f"عرض {len(sel.nodes):,} من أصل {sel.total_in_scope:,} عقدة في النطاق."
    if sel.truncated:
        msg += (f" تم الاقتصار على أول {sel.levels_shown + 1} مستويات للحفاظ على وضوح الرسم "
                f"(الحد الأقصى {s.max_network_nodes} عقدة، قابل للتعديل من Settings).")
    st.caption(msg)
    st.plotly_chart(network_figure(model, sel, color_by, s.low_span_threshold,
                                   s.high_span_threshold), width="stretch",
                    config={"scrollZoom": True})
    st.caption("حجم العقدة ∝ إجمالي التابعين. مرّر المؤشر على أي عقدة لعرض التفاصيل؛ استخدم العجلة للتكبير.")
    with st.expander("العقد المعروضة (جدول)"):
        nodes = df[df["node_id"].isin(sel.nodes)][
            ["node_id", "employee_name", "job_title", "department", "manager_id", "depth",
             "direct_reports", "total_reports", "reporting_status"]]
        table(nodes, "network_nodes", height=320)


# ================================================================================================
# Page: Employee Explorer  (ui/explorer.py)
# ================================================================================================
REPORT_COLS = ["node_id", "employee_name", "job_title", "job_level", "department", "manager_id",
               "depth", "direct_reports", "total_reports", "position_status", "reporting_status"]


def page_explorer() -> None:
    ctx = context()
    page_header("Employee Explorer", "ابحث عن موظف لعرض سلسلته الإدارية والموظفين التابعين له")
    model = ctx.model
    df = model.df[~model.df["is_duplicate_id"]]
    q = st.text_input("Search (Employee ID أو Employee Name)", key="exp_q",
                      placeholder="مثال: E10492 أو Smith")
    pool = df[df["is_active"] | df["is_vacant"] | df["employment_status"].eq("Terminated")]
    if q:
        m = (pool["node_id"].str.contains(q, case=False, regex=False)
             | pool["employee_name"].str.contains(q, case=False, regex=False))
        pool = pool[m]
    else:
        pool = pool.sort_values("total_reports", ascending=False)
    if pool.empty:
        st.warning("لا توجد نتائج مطابقة.")
        return
    st.caption(f"{len(pool):,} نتيجة" + (" (أول 300 معروضة)" if len(pool) > 300 else ""))
    options = pool.head(300)
    labels = {r.node_id: f"{r.node_id} · {r.employee_name or '(Vacant)'} · {r.job_title} · {r.department}"
              for r in options.itertuples()}
    emp_id = st.selectbox("Employee", list(labels), format_func=labels.get, key="exp_sel")
    r = model.node(emp_id)
    if r is None:
        return

    st.subheader(f"{r['employee_name'] or '(Vacant position)'} — {r['job_title']}")
    c = st.columns(6)
    kpi(c[0], "Position", r["position_id"] or "—", f"Position Status: {r['position_status']}")
    kpi(c[1], "Department", r["department"] or "—", f"{r['division']} · {r['business_unit']}")
    kpi(c[2], "Salary", fmt_num(r["salary"], 0))
    lvl = r["job_level"]
    kpi(c[3], "Job Level", "—" if pd.isna(lvl) else f"{int(lvl)} · {LEVEL_NAMES.get(int(lvl), '')}",
        f"Grade: {r['grade']}")
    kpi(c[4], "Years of Service", fmt_num(r["years_of_service"]))
    kpi(c[5], "Depth", fmt_num(r["depth"]), f"Reporting status: {r['reporting_status']}")
    st.caption(f"Employment: {r['employment_status']} · {r['employment_type']} · "
               f"{r['work_arrangement']} · {r['location']} · Performance: {fmt_num(r['performance_rating'])}")

    left, right = st.columns([1, 2])
    with left:
        st.markdown("**السلسلة الإدارية للأعلى**")
        chain = model.chain_up(emp_id)
        loop = len(chain) > 1 and chain[-1] in chain[:-1]
        parts = []
        for i, nid in enumerate(chain):
            n = model.node(nid)
            if n is None:
                body = f"<b>{html.escape(nid)}</b> — <i>غير موجود في البيانات</i>"
            else:
                status = "" if n["employment_status"] in ("Active", "On Leave", "") else \
                    f" <b>[{html.escape(n['employment_status'])}]</b>"
                body = (f"<b>{html.escape(nid)}</b> · {html.escape(n['employee_name'] or '(Vacant)')}"
                        f"{status}<br><small>{html.escape(n['job_title'])} · "
                        f"{html.escape(n['department'])}</small>")
            if loop and i == len(chain) - 1:
                body = f"↻ {body} <small>(circular reporting)</small>"
            cls = "node sel" if i == 0 else "node"
            parts.append(f"<div class='{cls}'>{body}</div>")
        tail = ""
        last = model.node(chain[-1])
        if not loop and last is not None and last["manager_id"] == "" and chain[-1] != model.ceo_id:
            tail = "<div class='arrow'>↑</div><div class='node'><i>لا يوجد Manager ID</i></div>"
        st.markdown("<div class='xr-chain'>" + "<div class='arrow'>↑</div>".join(parts) + tail +
                    "</div>", unsafe_allow_html=True)
        if chain[-1] == model.ceo_id:
            st.caption(f"تصل السلسلة إلى الـCEO في {len(chain) - 1} خطوة.")
        else:
            st.warning("السلسلة لا تصل إلى الـCEO.")
    with right:
        direct = model.direct_report_ids(emp_id, include_vacant=True)
        indirect = sorted(set(model.subtree_ids(emp_id, include_root=False, include_vacant=True))
                          - set(direct))
        st.markdown(f"**Direct Reports ({len(direct)})**")
        if direct:
            table(df[df["node_id"].isin(direct)][REPORT_COLS], "exp_direct", f"direct_reports_{emp_id}",
                  height=240)
        else:
            st.caption("لا يوجد.")
        st.markdown(f"**Indirect Reports ({len(indirect)})**")
        if indirect:
            table(df[df["node_id"].isin(indirect)][REPORT_COLS].sort_values("depth"),
                  "exp_indirect", f"indirect_reports_{emp_id}", height=260)
        else:
            st.caption("لا يوجد.")

    f = ctx.findings_all
    f = f[f["related_ids"].map(lambda ids: emp_id in ids)]
    if len(f):
        st.markdown(f"**Findings المرتبطة ({len(f)})**")
        v = f[["finding_id", "type", "severity", "description"]].copy()
        v["severity"] = v["severity"].map(severity_label)
        table(v, "exp_findings", f"findings_{emp_id}")


# ================================================================================================
# Page: X-Ray Findings  (ui/findings.py)
# ================================================================================================
LIST_COLS = ["finding_id", "type", "severity", "entity_type", "entity_id", "entity_name",
             "department", "description", "evidence", "recommended_investigation"]


def _fmt_value(v, threshold=False):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return "—"
    return str(v)


def page_findings() -> None:
    ctx = context()
    page_header("X-Ray Findings",
                "ملاحظات تنظيمية مكتشفة آليًا. كل ملاحظة تستدعي المراجعة ولا تمثل حكمًا نهائيًا.", ctx)
    f = ctx.findings
    if f.empty:
        empty_state("لا توجد ملاحظات.")
        return
    c = st.columns(5)
    kpi(c[0], "Findings", f"{len(f):,}")
    for col, sev in zip(c[1:], SEVERITIES):
        kpi(col, severity_label(sev), f"{int((f['severity'] == sev).sum()):,}")

    c1, c2, c3, c4 = st.columns([1.3, 1, 1.2, 1.5])
    types = c1.multiselect("Type", [t for t in FINDING_TYPES if t in set(f["type"])],
                           key="fd_type", placeholder="All",
                           format_func=lambda t: f"{t} · {TYPE_LABELS_AR.get(t, '')}")
    sevs = c2.multiselect("Severity", SEVERITIES, key="fd_sev", placeholder="All")
    depts = c3.multiselect("Department", sorted(v for v in f["department"].unique() if v),
                           key="fd_dept", placeholder="All")
    q = c4.text_input("بحث في النص / المعرفات", key="fd_q")
    view = f
    if types:
        view = view[view["type"].isin(types)]
    if sevs:
        view = view[view["severity"].isin(sevs)]
    if depts:
        view = view[view["department"].isin(depts)]
    if q:
        hay = view[["finding_id", "entity_id", "entity_name", "description", "evidence"]] \
            .fillna("").astype(str).agg(" ".join, axis=1)
        view = view[hay.str.contains(q, case=False, regex=False)]

    with st.expander("توزيع الملاحظات حسب النوع", expanded=False):
        st.plotly_chart(findings_by_type(view), width="stretch")

    st.caption(f"{len(view):,} ملاحظة. اختر صفًا لعرض التفسير الكامل (Explainability).")
    shown = view[LIST_COLS].copy()
    shown["severity"] = shown["severity"].map(severity_label)
    event = table(shown, "findings_table", "xray_findings", height=380,
                  on_select="rerun", selection_mode="single-row")
    sel_rows = event.selection.rows if event is not None else []
    options = view["finding_id"].tolist()
    if not options:
        return
    default = options[sel_rows[0]] if sel_rows else options[0]
    fid = st.selectbox("Finding", options, index=options.index(default), key=f"fd_sel_{default}")
    _explain(ctx, fid)


def _explain(ctx, fid: str) -> None:
    e = explain_finding(ctx.findings_all, fid, ctx.model)
    st.subheader(f"{e['finding_id']} · {e['type']}")
    st.markdown(f"**{severity_label(e['severity'])}** · {e['entity_type']}: `{e['entity_id']}` — "
                f"{e['entity_name']} · Department: {e['department'] or '—'}")
    st.markdown(e["description"])
    c = st.columns(4)
    c[0].metric("Metric", e["metric"] or "—", border=True)
    c[1].metric("Value", _fmt_value(e["value"]), border=True)
    c[2].metric("Configured Threshold", _fmt_value(e["threshold"]), border=True)
    c[3].metric("Benchmark", e["benchmark"] or "—", border=True)
    st.markdown("**Evidence**")
    st.info(e["evidence"])
    st.markdown("**Rule**")
    st.code(e["rule"], language=None)
    st.markdown("**Recommended Investigation**")
    st.markdown(e["recommended_investigation"])
    recs = e["records"]
    st.markdown(f"**السجلات المصدرية ({len(recs)})**")
    cols = ["row_number", "node_id", "employee_name", "job_title", "job_level", "department",
            "manager_id", "manager_name", "employment_status", "position_id", "position_status",
            "depth", "direct_reports", "total_reports", "reporting_status"]
    table(recs[cols], "fd_records", f"evidence_{fid}", height=300)


# ================================================================================================
# Page: Data Quality  (ui/data_quality.py)
# ================================================================================================
RECORD_COLS = ["row_number", "node_id", "employee_id", "employee_name", "job_title", "job_level",
               "department", "division", "manager_id", "manager_name", "employment_status",
               "position_id", "position_status", "salary", "hire_date", "reporting_status"]


def page_data_quality() -> None:
    ctx = context()
    page_header("Data Quality", "Completeness · Uniqueness · Validity · Consistency · Referential Integrity")
    st.caption("تُحسب جودة البيانات على كامل الملف (بدون الفلاتر العامة) حتى لا تُخفى السجلات المعيبة.")
    checks = cached_quality(ctx.token, ctx.model)
    dims = dimension_scores(checks)
    c = st.columns(len(dims) + 1)
    kpi(c[0], "Overall", pct(overall_score(checks)), "نسبة الفحوص الناجحة على مستوى السجلات.")
    for col, r in zip(c[1:], dims.itertuples()):
        kpi(col, r.dimension, pct(r.score), f"{r.failed} سجل غير مطابق من {r.records:,} فحص")

    st.subheader("المقاييس")
    st.caption("اختر أي مقياس لعرض السجلات المتأثرة.")
    view = checks.drop(columns=["failing_rows", "error"])
    status_icon = {"Pass": "✅ Pass", "Warning": "🟡 Warning", "Fail": "🔴 Fail"}
    view = view.assign(status=view["status"].map(status_icon))
    event = table(view, "dq_checks", "data_quality_checks", height=420, on_select="rerun",
                  selection_mode="single-row")
    rows = event.selection.rows if event is not None else []
    failing = checks[checks["failed"] > 0]
    default_id = checks.iloc[rows[0]]["check_id"] if rows else \
        (failing.iloc[0]["check_id"] if len(failing) else checks.iloc[0]["check_id"])
    ids = checks["check_id"].tolist()
    names = dict(zip(checks["check_id"], checks["check"]))
    cid = st.selectbox("المقياس", ids, index=ids.index(default_id),
                       format_func=lambda i: f"{i} · {names[i]}", key=f"dq_sel_{default_id}")
    chk = checks[checks["check_id"].eq(cid)].iloc[0]
    st.markdown(f"**{chk['check']}** — {chk['rule']} · Score: **{pct(chk['score'])}** "
                f"({chk['failed']} من {chk['records_checked']:,})")
    recs = affected_records(ctx.model, checks, cid)
    if recs.empty:
        st.success("لا توجد سجلات متأثرة.")
    else:
        table(recs[RECORD_COLS], "dq_records", f"dq_{cid}", height=320)


# ================================================================================================
# Page: What-If Simulator  (ui/simulator.py)
# ================================================================================================
ACTION_LABELS = {
    "merge_departments": "دمج إدارات",
    "remove_layer": "إزالة طبقة إدارية",
    "move_employee": "نقل موظف",
    "remove_manager": "إزالة مدير",
}


def _describe(a: dict) -> str:
    t = a["type"]
    if t == "merge_departments":
        return (f"دمج {' + '.join(a['departments'])} → {a['new_name']}"
                + (" (توحيد القيادة)" if a.get("consolidate_heads") else ""))
    if t == "remove_layer":
        return (f"إزالة المديرين في Job Level {a['job_level']} ({LEVEL_NAMES.get(a['job_level'], '')})"
                + (f" في {a['department']}" if a.get("department") else " في كل المنظمة"))
    if t == "move_employee":
        return f"نقل {a['employee_id']} إلى المدير {a['new_manager_id']}"
    if t == "remove_manager":
        return f"إزالة المدير {a['manager_id']} → المستلم: {a['reassign_to']}"
    return json.dumps(a, ensure_ascii=False)


def _add(action: dict) -> None:
    st.session_state.scenario = st.session_state.scenario + [action]


def _labels(df: pd.DataFrame) -> dict[str, str]:
    return {r.node_id: f"{r.node_id} · {r.employee_name} · {r.job_title} · {r.department}"
            for r in df.itertuples()}


def _builder(ctx) -> None:
    model = ctx.model
    df = model.df
    act = df[df["is_active"]]
    heads = department_heads(model)
    depts = sorted(heads)
    t1, t2, t3, t4 = st.tabs(list(ACTION_LABELS.values()))

    with t1:
        sel = st.multiselect("الإدارات المراد دمجها (الأولى = الإدارة الأساسية)", depts, key="sim_mg_depts",
                             max_selections=6)
        name = st.text_input("اسم الإدارة الجديدة", value=" & ".join(sel), key=f"sim_mg_name_{'|'.join(sel)}")
        cons = st.checkbox("توحيد القيادة: نقل فرق رؤساء الإدارات الأخرى إلى رئيس الإدارة الأساسية "
                           "وحذف وظائفهم", key="sim_mg_cons")
        if sel:
            h = heads.get(sel[0])
            st.caption(f"رئيس الإدارة الأساسية: {h} · {model.node(h)['employee_name']} · "
                       f"{model.node(h)['job_title']}" if h else "")
        if st.button("➕ أضف إلى السيناريو", key="sim_mg_add", disabled=len(sel) < 2):
            _add({"type": "merge_departments", "departments": sel, "new_name": name,
                  "consolidate_heads": cons})
            st.rerun()

    with t2:
        c1, c2 = st.columns(2)
        lvl = c1.selectbox("Job Level", [1, 2, 3, 4, 5], index=3, key="sim_ly_lvl",
                           format_func=lambda v: f"{v} · {LEVEL_NAMES[v]}")
        dept = c2.selectbox("النطاق", [None] + depts, key="sim_ly_dept",
                            format_func=lambda v: "كل المنظمة" if v is None else v)
        n = act[act["is_manager"] & act["job_level"].eq(lvl)]
        if dept:
            n = n[n["department"].eq(dept)]
        st.caption(f"سيتم حذف {len(n)} وظيفة إشرافية، ونقل {int(n['direct_reports'].sum())} موظف "
                   f"و{int(n['vacant_direct_positions'].sum())} وظيفة شاغرة (تابعين مباشرين) إلى "
                   "المدير الأعلى التالي في سلسلتهم.")
        if st.button("➕ أضف إلى السيناريو", key="sim_ly_add", disabled=n.empty):
            _add({"type": "remove_layer", "job_level": int(lvl), "department": dept})
            st.rerun()

    with t3:
        labels = _labels(act.sort_values("node_id"))
        emp = st.selectbox("الموظف", list(labels), format_func=labels.get, key="sim_mv_emp")
        cur = model.node(emp)
        st.caption(f"المدير الحالي: {cur['manager_id'] or '—'} · {cur['manager_name'] or ''}")
        new = st.selectbox("المدير الجديد", [k for k in labels if k != emp], format_func=labels.get,
                           key="sim_mv_new")
        mvd = st.checkbox("نقل الموظف إلى إدارة المدير الجديد", value=True, key="sim_mv_dept")
        if st.button("➕ أضف إلى السيناريو", key="sim_mv_add"):
            _add({"type": "move_employee", "employee_id": emp, "new_manager_id": new,
                  "move_department": mvd})
            st.rerun()

    with t4:
        mgrs = _labels(act[act["is_manager"] & act["node_id"].ne(model.ceo_id)]
                       .sort_values("direct_reports"))
        mid = st.selectbox("المدير المراد إزالته", list(mgrs), format_func=mgrs.get, key="sim_rm_mgr")
        m = model.node(mid)
        head = heads.get(m["department"])
        st.caption(f"لديه {int(m['direct_reports'])} موظف مباشر · مديره: {m['manager_id']} · "
                   f"رئيس الإدارة: {head}")
        choice = st.radio("من سيستلم الموظفين؟",
                          ["manager_of_removed", "department_head", "specific"], key="sim_rm_choice",
                          format_func={"manager_of_removed": f"المدير الأعلى ({m['manager_id']})",
                                       "department_head": f"رئيس الإدارة ({head})",
                                       "specific": "مدير آخر / ترقية أحد الموظفين المباشرين"}.get)
        target = choice
        if choice == "specific":
            directs = model.direct_report_ids(mid)
            cand = act[act["node_id"].ne(mid)]
            order = directs + [n for n in cand["node_id"] if n not in set(directs)]
            lab = _labels(cand.set_index("node_id").loc[order].reset_index())
            target = st.selectbox("المستلم (الموظفون المباشرون أولًا = ترقية)", order,
                                  format_func=lab.get, key="sim_rm_target")
        if st.button("➕ أضف إلى السيناريو", key="sim_rm_add"):
            _add({"type": "remove_manager", "manager_id": mid, "reassign_to": target})
            st.rerun()


def _run(ctx, actions: list[dict]):
    key = (ctx.token, json.dumps(actions, sort_keys=True), ctx.settings.cache_key())
    cache = st.session_state.setdefault("_sim_cache", {})
    if key not in cache:
        cache.clear()
        with st.spinner("Running simulation on a copy of the data…"):
            res = apply_scenario(ctx.model.df, actions, ctx.model)
            cmp = compare(ctx.model, res.model, ctx.settings)
            f_after = detect_findings(res.model, ctx.settings)
        cache[key] = (res, cmp, f_after)
    return cache[key]


def page_simulator() -> None:
    ctx = context()
    page_header("What-If Simulator",
                "جرّب تغييرات افتراضية على نسخة من البيانات وقارن النتائج قبل التطبيق الفعلي. "
                "البيانات الأصلية لا تتغير.")
    st.caption("المحاكاة تعمل على كامل المنظمة (بدون الفلاتر العامة).")
    _builder(ctx)

    st.subheader("السيناريو")
    scen = st.session_state.scenario
    if not scen:
        st.info("أضف إجراءً واحدًا أو أكثر لبدء المحاكاة. يمكن جمع عدة إجراءات في سيناريو واحد.")
        return
    for i, a in enumerate(scen):
        c1, c2 = st.columns([10, 1])
        c1.markdown(f"{i + 1}. {_describe(a)}")
        if c2.button("✕", key=f"sim_del_{i}_{len(scen)}"):
            st.session_state.scenario = scen[:i] + scen[i + 1:]
            st.rerun()
    c1, c2 = st.columns([1, 1])
    if c1.button("مسح السيناريو", key="sim_clear"):
        st.session_state.scenario = []
        st.rerun()
    c2.download_button("⬇ Scenario JSON", json.dumps(scen, ensure_ascii=False, indent=2),
                       "scenario.json", "application/json")

    try:
        res, cmp, f_after = _run(ctx, scen)
    except SimulationError as exc:
        st.error(f"تعذر تطبيق السيناريو: {exc}")
        return
    for line in res.log:
        st.caption(line)

    st.subheader("Before / After")
    summ = cmp["summary"][["Metric", "Before (fmt)", "After (fmt)"]].rename(
        columns={"Before (fmt)": "BEFORE", "After (fmt)": "AFTER"})
    summ["Δ"] = ["" if b == a else "●" for b, a in zip(summ["BEFORE"], summ["AFTER"])]
    left, right = st.columns([1, 1])
    with left:
        table(summ, "sim_summary", "scenario_before_after", rename=False,
              height=35 * (len(summ) + 1) + 3)
    with right:
        st.markdown("**ما الذي تغير؟**")
        lines = cmp["narrative"]
        if lines:
            st.markdown("<div class='ltr'>" + "<br>".join(lines[:40]) + "</div>",
                        unsafe_allow_html=True)
            if len(lines) > 40:
                st.caption(f"… و{len(lines) - 40} تغييرات أخرى في الجداول أدناه.")
        else:
            st.caption("لا توجد تغييرات في المؤشرات.")

    last = scen[-1]
    if last["type"] == "move_employee":
        _move_focus(ctx, res, cmp, last)

    t1, t2, t3, t4 = st.tabs(["Span of Control", "Departments", "Organizational Depth", "Findings"])
    with t1:
        table(cmp["span_changes"], "sim_spans", "scenario_span_changes")
    with t2:
        table(cmp["department_changes"], "sim_depts", "scenario_department_changes", rename=False)
    with t3:
        dc = cmp["depth_changes"]
        st.caption(f"{len(dc)} موظف تغير عمقهم.")
        table(dc, "sim_depth", "scenario_depth_changes")
    with t4:
        fb = ctx.findings_all["type"].value_counts()
        fa = f_after["type"].value_counts()
        comp = pd.DataFrame({"Before": fb, "After": fa}).fillna(0).astype(int)
        comp["Change"] = comp["After"] - comp["Before"]
        table(comp.reset_index(names="Finding Type"), "sim_findings", "scenario_findings",
              rename=False)


def _move_focus(ctx, res, cmp, a: dict) -> None:
    st.subheader("أثر نقل الموظف")
    before, after = ctx.model, res.model
    emp = a["employee_id"]
    old_mgr = before.node(emp)["manager_id"]
    rows = []
    for label, nid in [("المدير الحالي", old_mgr), ("المدير الجديد", a["new_manager_id"])]:
        b, af = before.node(nid), after.node(nid)
        rows.append({"Role": label, "Manager": nid,
                     "Span Before": None if b is None else int(b["direct_reports"]),
                     "Span After": None if af is None else int(af["direct_reports"])})
    st.markdown("**Span of Control قبل وبعد**")
    table(pd.DataFrame(rows), "sim_mv_span", rename=False)
    sub = [emp] + before.subtree_ids(emp, include_root=False)
    b = before.df.set_index("node_id").loc[sub, ["employee_name", "depth", "department"]]
    af = after.df[~after.df["is_duplicate_id"]].set_index("node_id").reindex(sub)[["depth", "department"]]
    d = b.rename(columns={"depth": "Depth Before", "department": "Department Before"})
    d["Depth After"] = af["depth"]
    d["Department After"] = af["department"]
    st.markdown("**Organizational Depth قبل وبعد (الموظف وفريقه)**")
    table(d.reset_index(), "sim_mv_depth", rename=False)
    st.markdown("**Department distribution قبل وبعد**")
    dc = cmp["department_changes"]
    table(dc[["department", "headcount_before", "headcount_after"]] if len(dc) else dc,
          "sim_mv_dept", rename=False)


# ================================================================================================
# Page: Import Data  (ui/import_data.py)
# ================================================================================================
TEMPLATE_HEADERS = ["EmployeeID", "EmployeeName", "ManagerID", "JobTitle", "Department", "Division",
                    "BusinessUnit", "JobLevel", "Salary", "EmploymentStatus", "PositionID",
                    "PositionStatus", "Location", "Gender", "Age", "HireDate", "Grade",
                    "EmploymentType", "WorkArrangement", "PerformanceRating", "YearsOfService"]


@st.cache_data(show_spinner=False, max_entries=4)
def _facts(token: str, _raw: pd.DataFrame) -> dict:
    return dataset_facts(_raw)


def page_import_data() -> None:
    page_header("Import Data", "تحميل بيانات المنظمة (CSV / Excel) أو إنشاء بيانات وهمية")
    ds = st.session_state.dataset
    st.success(f"البيانات الحالية: {ds['source']}")

    t1, t2, t3 = st.tabs(["رفع ملف", "بيانات وهمية (Dummy Data)", "وصف البيانات الحالية"])
    with t1:
        st.markdown("الحقول الأساسية: " + ", ".join(f"`{h}`" for h in TEMPLATE_HEADERS[:12]))
        st.caption("يتم التعرف على أسماء الأعمدة الشائعة تلقائيًا (مثل Employee ID / employee_id / "
                   "رقم الموظف). الوظائف الشاغرة: صف بدون EmployeeID مع PositionStatus = Vacant.")
        tmpl = pd.DataFrame(columns=TEMPLATE_HEADERS)
        c1, c2 = st.columns(2)
        c1.download_button("⬇ قالب CSV", to_csv_bytes(tmpl), "org_xray_template.csv", "text/csv")
        c2.download_button("⬇ قالب Excel", to_excel_bytes({"Employees": tmpl}), "org_xray_template.xlsx")
        up = st.file_uploader("ملف CSV أو Excel", type=["csv", "xlsx", "xlsm"], key="imp_file")
        if up is not None:
            try:
                raw = read_uploaded(up.getvalue(), up.name)
            except Exception as exc:
                st.error(f"تعذر قراءة الملف: {exc}")
                return
            rep = validate(raw)
            st.markdown(f"**نتيجة التحقق (Validation)** — {rep.n_rows:,} سجل")
            mapping = pd.DataFrame({"Column in file": list(rep.column_mapping),
                                    "Mapped to": [DISPLAY_NAMES.get(v, v)
                                                  for v in rep.column_mapping.values()]})
            with st.expander("مطابقة الأعمدة", expanded=bool(rep.missing_columns)):
                st.dataframe(mapping, hide_index=True, width="stretch")
                if rep.missing_columns:
                    st.warning("أعمدة مطلوبة غير موجودة: " + ", ".join(
                        DISPLAY_NAMES.get(c, c) for c in rep.missing_columns))
            issues = rep.to_frame()
            icon = {"Error": "🔴 Error", "Warning": "🟡 Warning", "Info": "ℹ️ Info"}
            issues["Level"] = issues["Level"].map(icon)
            table(issues, "imp_issues", "validation_report", rename=False)
            st.markdown("**معاينة**")
            st.dataframe(raw.head(20), hide_index=True, width="stretch")
            if rep.ok:
                if st.button("تحميل البيانات للتحليل", type="primary", key="imp_load"):
                    load_uploaded(raw, up.name)
                    st.success("تم تحميل البيانات. انتقل إلى Dashboard.")
                    st.rerun()
            else:
                st.error("لا يمكن تحليل الملف قبل معالجة الأخطاء (Error) أعلاه. "
                         f"الأعمدة الأساسية للهيكل: {', '.join(IMPORT_REQUIRED_COLUMNS[:1])} و ManagerID.")

    with t2:
        st.markdown("يتم بناء الهيكل التنظيمي أولًا ثم توزيع موظفين (Faker) داخله، مع إدخال مشكلات "
                    "تنظيمية متعمدة يمكن للنظام اكتشافها.")
        c1, c2 = st.columns(2)
        n = c1.select_slider("عدد الموظفين التقريبي", [1500, 3000, 5000, 10000], value=1500, key="imp_n")
        seed = c2.number_input("Seed", 0, 10_000, 42, key="imp_seed")
        if st.button("إنشاء البيانات", type="primary", key="imp_gen"):
            load_dummy(int(n), int(seed))
            st.rerun()
        raw = ds["raw"]
        c1, c2 = st.columns(2)
        c1.download_button("⬇ البيانات الحالية CSV", to_csv_bytes(raw), "org_data.csv", "text/csv")
        c2.download_button("⬇ البيانات الحالية Excel", to_excel_bytes({"Employees": raw}),
                           "org_data.xlsx")

    with t3:
        f = _facts(ds["token"], ds["raw"])
        c = st.columns(4)
        c[0].metric("السجلات", f"{f['records']:,}", border=True)
        c[1].metric("الموظفون النشطون", f"{f['active_employees']:,}", border=True)
        c[2].metric("المديرون", f"{f['managers']:,}", border=True)
        c[3].metric("الطبقات", f"{f['layers']}", border=True)
        c = st.columns(4)
        c[0].metric("الإدارات", f"{f['departments']}", border=True)
        c[1].metric("القطاعات", f"{f['divisions']}", border=True)
        c[2].metric("الوظائف الشاغرة", f"{f['vacant_positions']}", border=True)
        c[3].metric("Findings", f"{f['findings_total']}", border=True)
        st.caption("الأرقام محسوبة ديناميكيًا من البيانات الحالية بالإعدادات الافتراضية.")
        if ds.get("manifest"):
            st.markdown("**المشكلات التي تم إدخالها عمدًا في البيانات الوهمية (Manifest) والنتيجة المتوقعة**")
            man = pd.DataFrame([{**m, "ids": ", ".join(m["ids"][:8]) + (" …" if len(m["ids"]) > 8 else "")}
                                for m in ds["manifest"]])
            table(man, "imp_manifest", "injected_issues", rename=False)


# ================================================================================================
# Page: Export  (ui/export_page.py)
# ================================================================================================
def page_export() -> None:
    ctx = context()
    page_header("Export", "تصدير النتائج إلى Excel أو CSV", ctx)
    use_filters = st.toggle("تطبيق الفلاتر العامة على التصدير", value=ctx.filtered,
                            disabled=not ctx.filtered, key="exp_filters")
    df = ctx.df if use_filters else ctx.model.df
    findings = ctx.findings if use_filters else ctx.findings_all
    with st.spinner("Preparing export…"):
        sheets = build_export_tables(ctx.model, findings, ctx.settings, df)
    st.markdown("**ملف Excel واحد يحتوي على الأوراق التالية:**")
    st.markdown("<div class='ltr'>" + " · ".join(f"{k} ({len(v):,})" for k, v in sheets.items())
                + "</div>", unsafe_allow_html=True)
    st.download_button("⬇ Download Excel (organizational_xray.xlsx)", to_excel_bytes(sheets),
                       "organizational_xray.xlsx", type="primary",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    st.subheader("CSV")
    name = st.selectbox("الجدول", list(sheets), key="exp_csv_sheet")
    st.dataframe(sheets[name].head(200), hide_index=True, width="stretch", height=320)
    st.download_button(f"⬇ {name}.csv", to_csv_bytes(sheets[name]),
                       f"{name.lower().replace(' ', '_')}.csv", "text/csv")
    st.caption("كل جدول في التطبيق يحتوي أيضًا على زر ⬇ CSV خاص به.")


# ================================================================================================
# Page: Settings  (ui/settings_page.py)
# ================================================================================================
def page_settings() -> None:
    page_header("Settings", "حدود التحليل (Thresholds). تُطبق فورًا على كل الصفحات والـFindings.")
    cur: Settings = st.session_state.settings
    d = Settings()
    with st.form("settings_form"):
        c1, c2 = st.columns(2)
        low = c1.number_input("Low Span Threshold", 0, 50, cur.low_span_threshold,
                              help=SETTING_DESCRIPTIONS["low_span_threshold"])
        high = c2.number_input("High Span Threshold", 1, 200, cur.high_span_threshold,
                               help=SETTING_DESCRIPTIONS["high_span_threshold"])
        deep = c1.number_input("Deep Hierarchy Threshold", 1, 30, cur.deep_hierarchy_threshold,
                               help=SETTING_DESCRIPTIONS["deep_hierarchy_threshold"])
        small = c2.number_input("Small Department Threshold", 0, 100, cur.small_department_threshold,
                                help=SETTING_DESCRIPTIONS["small_department_threshold"])
        ratio = c1.number_input("High Management Ratio (%)", 1.0, 100.0,
                                round(cur.high_management_ratio * 100, 1), step=1.0,
                                help=SETTING_DESCRIPTIONS["high_management_ratio"])
        sim = c2.slider("Title Similarity Threshold", 0.5, 1.0, float(cur.title_similarity_threshold),
                        0.01, help=SETTING_DESCRIPTIONS["title_similarity_threshold"])
        excl = c1.checkbox("استثناء CEO و C-Suite من ملاحظات Span of Control",
                           cur.exclude_executives_from_span)
        nodes = c2.number_input("Max network nodes", 50, 5000, cur.max_network_nodes, step=50,
                                help=SETTING_DESCRIPTIONS["max_network_nodes"])
        saved = st.form_submit_button("حفظ الإعدادات", type="primary")
    if saved:
        if low >= high:
            st.error("يجب أن يكون Low Span Threshold أقل من High Span Threshold.")
        else:
            st.session_state.settings = Settings(int(low), int(high), int(deep), int(small),
                                                 float(ratio) / 100, float(sim), bool(excl),
                                                 int(nodes))
            st.success("تم حفظ الإعدادات وإعادة حساب النتائج.")
    if st.button("استعادة القيم الافتراضية"):
        st.session_state.settings = Settings()
        st.rerun()

    st.subheader("القواعد الحالية")
    s = st.session_state.settings
    rules = [
        ("LOW_SPAN_OF_CONTROL", f"direct reports ≤ {s.low_span_threshold}"),
        ("HIGH_SPAN_OF_CONTROL", f"direct reports > {s.high_span_threshold}"),
        ("DEEP_HIERARCHY", f"depth to CEO > {s.deep_hierarchy_threshold}"),
        ("SMALL_DEPARTMENT", f"headcount ≤ {s.small_department_threshold}"),
        ("HIGH_MANAGEMENT_RATIO", f"managers ÷ headcount > {s.high_management_ratio:.0%} "
                                  f"(departments larger than {s.small_department_threshold})"),
        ("DUPLICATE_POSITION", f"duplicate Position ID, title variants, similarity ≥ {s.title_similarity_threshold:.2f}"),
    ]
    st.dataframe([{"Finding": r, "Rule": t} for r, t in rules], hide_index=True, width="stretch")
    st.caption(f"القيم الافتراضية: {json.dumps(d.as_dict())}")


# ================================================================================================
# Application entry point
# ================================================================================================
PAGES = [
    (page_dashboard, "Dashboard", ":material/monitor_heart:", "dashboard"),
    (page_layers, "Layers", ":material/layers:", "layers"),
    (page_span, "Span of Control", ":material/account_tree:", "span"),
    (page_departments, "Departments", ":material/apartment:", "departments"),
    (page_network, "Organization Map", ":material/hub:", "organization-map"),
    (page_explorer, "Employee Explorer", ":material/person_search:", "explorer"),
    (page_findings, "X-Ray Findings", ":material/radiology:", "findings"),
    (page_data_quality, "Data Quality", ":material/fact_check:", "data-quality"),
    (page_simulator, "What-If Simulator", ":material/science:", "simulator"),
    (page_import_data, "Import Data", ":material/upload_file:", "import"),
    (page_export, "Export", ":material/download:", "export"),
    (page_settings, "Settings", ":material/tune:", "settings"),
]


def run_app() -> None:
    st.set_page_config(page_title="Organizational X-Ray", page_icon="🩻", layout="wide")
    inject_css()
    init_state()
    pages = [st.Page(fn, title=title, icon=icon, url_path=url, default=(i == 0))
             for i, (fn, title, icon, url) in enumerate(PAGES)]
    nav = st.navigation({"Organizational X-Ray": pages})
    ds = st.session_state.dataset
    sidebar_filters(cached_model(ds["token"], ds["raw"]))
    nav.run()


def _cli() -> None:
    parser = argparse.ArgumentParser(description="Organizational X-Ray")
    parser.add_argument("--generate", action="store_true",
                        help="write the dummy dataset to data/sample and print a report")
    parser.add_argument("--employees", type=int, default=1500)
    parser.add_argument("--seed", type=int, default=42)
    args, streamlit_args = parser.parse_known_args()  # e.g. --server.port 8502 -> Streamlit
    if args.generate:
        df, manifest = generate_org(args.employees, args.seed)
        paths = save_sample(df)
        print("Saved:", ", ".join(str(p) for p in paths.values()))
        print(dataset_report(df, manifest))
        return
    # Started with plain `python app.py`: relaunch the same file through Streamlit.
    credentials = Path.home() / ".streamlit" / "credentials.toml"
    if not credentials.exists():  # skip Streamlit's first-run e-mail prompt
        credentials.parent.mkdir(exist_ok=True)
        credentials.write_text('[general]\nemail = ""\n', encoding="utf-8")
    sys.exit(subprocess.call([sys.executable, "-m", "streamlit", "run", __file__,
                              *streamlit_args]))


if __name__ == "__main__":
    if st.runtime.exists():
        run_app()
    else:
        _cli()
