"""Excel / CSV export."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from analytics.data_quality import run_checks
from analytics.departments import layer_metrics, unit_metrics
from analytics.hierarchy import OrgModel
from analytics.span_of_control import manager_table, span_by, span_stats
from utils.config import CANONICAL_COLUMNS, Settings
from utils.helpers import display_df

EMPLOYEE_EXPORT_COLUMNS = ["node_id"] + CANONICAL_COLUMNS + [
    "reporting_status", "depth", "direct_reports", "vacant_direct_positions", "total_reports",
    "is_manager",
]


def build_export_tables(model: OrgModel, findings: pd.DataFrame, settings: Settings,
                        df: pd.DataFrame | None = None) -> dict[str, pd.DataFrame]:
    """All export sheets, computed from ``df`` (filtered) or the full model."""
    data = df if df is not None else model.df
    mgr = manager_table(data, settings)
    stats = span_stats(mgr["direct_reports"])
    span_sheet = pd.concat([
        pd.DataFrame({"Statistic": [k.upper() if k.startswith("p") else k.title() for k in stats],
                      "Value": list(stats.values())}),
        pd.DataFrame({"Statistic": [""], "Value": [None]}),
        span_by(data, "department", settings).rename(columns={"department": "Statistic"}),
    ], ignore_index=True)
    dq = run_checks(model, data).drop(columns=["failing_rows"])
    layers = pd.concat([layer_metrics(data, "job_level").assign(basis="Job Level"),
                        layer_metrics(data, "depth").assign(basis="Structural Depth")],
                       ignore_index=True)
    f = findings.copy()
    if not f.empty:
        f["related_ids"] = f["related_ids"].map(lambda ids: ", ".join(ids[:50]))
    cols = [c for c in EMPLOYEE_EXPORT_COLUMNS if c in data.columns]
    return {
        "Employees": display_df(data[cols]),
        "Managers": display_df(mgr),
        "Departments": display_df(unit_metrics(data, "department", settings)),
        "Findings": display_df(f),
        "Span of Control": span_sheet,
        "Organizational Layers": display_df(layers),
        "Data Quality": display_df(dq),
    }


def _flatten(v):
    return ", ".join(map(str, v)) if isinstance(v, (list, tuple, set)) else v


def to_excel_bytes(sheets: dict[str, pd.DataFrame]) -> bytes:
    buf = io.BytesIO()
    header_fill = PatternFill("solid", fgColor="1C5CAB")
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            out = frame.copy()
            for c in out.columns:  # Excel cannot store tz-aware datetimes / lists
                if isinstance(out[c].dtype, pd.DatetimeTZDtype):
                    out[c] = out[c].dt.tz_localize(None)
                if out[c].dtype == object:
                    out[c] = out[c].map(_flatten)
            sheet = name[:31]
            out.to_excel(writer, sheet_name=sheet, index=False)
            ws = writer.sheets[sheet]
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            for i, col in enumerate(out.columns, 1):
                sample = out[col].head(200).map(lambda v: len(str(v)))
                width = min(60, max(10, len(str(col)) + 2, int(sample.max() if len(sample) else 0) + 2))
                ws.column_dimensions[get_column_letter(i)].width = width
            if len(out):
                ws.auto_filter.ref = ws.dimensions
    return buf.getvalue()


def to_csv_bytes(df: pd.DataFrame) -> bytes:
    out = df.copy()
    for c in out.columns:
        if out[c].dtype == object:
            out[c] = out[c].map(_flatten)
    return out.to_csv(index=False).encode("utf-8-sig")
